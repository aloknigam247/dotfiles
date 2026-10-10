"""Fixtures for selftest.py.

- TYPES: semantic data for the 31 diagram and chart types, from the diagram catalog's scenarios
  (tmp/diagram_catalog/spec.md and its Mermaid sources). Data only: selftest.py calls
  getattr(deck, fx["builder"])(slide, *fx["args"], **fx["kwargs"]) once that builder exists and
  lists the rest as skipped. The data model is the one in reference/diagrams.md, as plain tuples:
  nodes (id, text, kind), edges (src, dst, label, style, arrow), groups (id, label, members).
- ADVERSARIAL: layout stress cases for the graph builders (dense graph, long labels, cycles,
  nested groups), and a 40-node graph too dense for one slide that must raise "split the
  diagram" (`raises`).
- `max_crossings` of a graph fixture: the crossings of its routed connectors that bench_layout.py
  allows; for the 31 types counted on the provider rendering in the diagram catalog, for the
  adversarial cases (no rendering) the count when the builders were written.
- primitives_slides(): node kinds and link() between every supported shape, side pair and route.
- NEGATIVE: decks with one defect each, and the com_check.py check that must report it.
"""
import random
import zipfile

from pptx.oxml.ns import qn
from pptx.util import Inches

from pptlib.shapes import PRESETS
from pptlib.style import CONTENT, NODE_KINDS

MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun")


def _samples(seed, mu, sigma, n):
    """Order values ($) of one channel: deterministic log-normal samples."""
    rng = random.Random(seed)
    return [round(rng.lognormvariate(mu, sigma), 2) for _ in range(n)]


SAMPLES = {"Web": _samples(42, 3.5, 0.4, 200), "Mobile": _samples(43, 3.3, 0.35, 200),
           "Partner": _samples(44, 3.8, 0.5, 100)}


def _fx(type_name, builder, *args, graph=False, chart_xml=None, chart_type=None,
        max_crossings=None, **kwargs):
    return {"args": args, "builder": builder, "chart_type": chart_type, "chart_xml": chart_xml,
            "graph": graph, "kwargs": kwargs, "max_crossings": max_crossings, "type": type_name}


def _chart(type_name, kind, categories, series, title, chart_xml):
    return _fx(type_name, "chart", kind, *CONTENT, categories, series, chart_xml=chart_xml,
               title=title)


TYPES = [
    _fx("Flowchart", "flowchart",
        [("start", "Start", "start"), ("cart", "View cart", "process"),
         ("signed", "Signed in?", "decision"), ("signin", "Sign in", "process"),
         ("ship", "Enter shipping address", "process"), ("pay", "Pay", "process"),
         ("ok", "Payment OK?", "decision"), ("place", "Place order", "process"),
         ("email", "Send confirmation email", "process"), ("error", "Show payment error", "error"),
         ("done", "Done", "end")],
        [("start", "cart"), ("cart", "signed"), ("signed", "signin", "No"), ("signin", "signed"),
         ("signed", "ship", "Yes"), ("ship", "pay"), ("pay", "ok"), ("ok", "place", "Yes"),
         ("place", "email"), ("email", "done"), ("ok", "error", "No"), ("error", "pay")],
        graph=True, max_crossings=0),
    _fx("Sequence diagram", "sequence",
        [("customer", "Customer", "actor"), ("web", "Web App", "participant"),
         ("api", "Order API", "participant"), ("payment", "Payment Service", "participant"),
         ("inventory", "Inventory DB", "participant")],
        [("customer", "web", 'Click "Place order"'), ("web", "api", "POST /orders"),
         ("api", "inventory", "Reserve stock"), ("inventory", "api", "Reserved", "reply"),
         ("api", "payment", "Charge card"), ("payment", "api", "Approved", "reply"),
         ("api", "web", "201 Created", "reply"), ("web", "customer", "Show confirmation", "reply"),
         ("payment", "api", "Declined", "reply"), ("api", "inventory", "Release stock"),
         ("api", "web", "402 Payment Required", "reply"),
         ("web", "customer", "Show error", "reply")],
        frames=[("alt", [("payment approved", 5, 7), ("payment declined", 8, 11)])],
        notes=[("api", "Idempotency key = cart id")]),
    _fx("Class diagram", "class_diagram",
        [("Customer", ["+id: UUID", "+name: string", "+email: string"],
          ["+placeOrder(cart) Order"], ""),
         ("Order", ["+id: UUID", "+createdAt: DateTime", "+status: OrderStatus"],
          ["+total() Money"], ""),
         ("OrderItem", ["+quantity: int", "+unitPrice: Money"], ["+subtotal() Money"], ""),
         ("Product", ["+sku: string", "+title: string", "+price: Money"], [], ""),
         ("Payment", ["+amount: Money"], ["+authorize() bool"], "abstract"),
         ("CardPayment", ["+last4: string"], [], ""),
         ("WalletPayment", ["+provider: string"], [], ""),
         ("OrderStatus", ["PENDING", "PAID", "SHIPPED", "DELIVERED", "CANCELLED"], [],
          "enumeration")],
        [("Customer", "Order", "association", "places", "1", "0..*"),
         ("Order", "OrderItem", "composition", "", "1", "1..*"),
         ("OrderItem", "Product", "association", "", "0..*", "1"),
         ("Order", "Payment", "aggregation", "", "1", "1"),
         ("CardPayment", "Payment", "inheritance", "", "", ""),
         ("WalletPayment", "Payment", "inheritance", "", "", ""),
         ("Order", "OrderStatus", "dependency", "", "", "")],
        graph=True, max_crossings=0),
    _fx("State diagram", "state_diagram",
        ["Pending", "Paid", "Shipped", "In transit", "Out for delivery", "Delivered", "Cancelled",
         "Refunded"],
        [("[*]", "Pending", ""), ("Pending", "Paid", "payment captured"),
         ("Pending", "Cancelled", "customer cancels"), ("Paid", "Shipped", "label printed"),
         ("Paid", "Refunded", "refund issued"),
         ("In transit", "Out for delivery", "reached local hub"),
         ("Shipped", "Delivered", "signed for"), ("Delivered", "[*]", ""), ("Cancelled", "[*]", ""),
         ("Refunded", "[*]", "")],
        composites=[("Shipped", ["In transit", "Out for delivery"])],
        notes=[("Paid", "Stock reserved")], graph=True, max_crossings=0),
    _fx("ER diagram", "er_diagram",
        [("customer", [("id", "uuid", "PK"), ("name", "text", ""), ("email", "text", "UK")]),
         ("orders", [("id", "uuid", "PK"), ("customer_id", "uuid", "FK"),
                     ("created_at", "timestamp", ""), ("status", "text", "")]),
         ("order_item", [("order_id", "uuid", "PK, FK"), ("product_id", "uuid", "PK, FK"),
                         ("quantity", "int", ""), ("unit_price", "numeric", "")]),
         ("product", [("id", "uuid", "PK"), ("title", "text", ""), ("price", "numeric", ""),
                      ("category_id", "int", "FK")]),
         ("category", [("id", "int", "PK"), ("name", "text", "")])],
        [("customer", "orders", "1", "0..*", "customer_id"),
         ("orders", "order_item", "1", "1..*", "order_id"),
         ("product", "order_item", "1", "0..*", "product_id"),
         ("category", "product", "1", "0..*", "category_id")],
        graph=True, max_crossings=0),
    _fx("Use case diagram", "use_case",
        [("customer", "Customer"), ("support", "Support agent"), ("stripe", "Payment provider")],
        [("browse", "Browse catalogue"), ("order", "Place order"), ("pay", "Pay by card"),
         ("gift", "Apply gift card"), ("track", "Track delivery"), ("refund", "Issue refund")],
        [("customer", "browse"), ("customer", "order"), ("customer", "track"),
         ("support", "refund"),
         ("order", "pay", "include"), ("gift", "order", "extend"), ("pay", "stripe"),
         ("refund", "stripe")],
        "Inkwell storefront", graph=True, max_crossings=0),
    _fx("Software architecture", "architecture",
        [("web", "Web app", "process"), ("mobile", "Mobile app", "process"),
         ("gateway", "API gateway", "process"), ("catalog", "Catalog service", "process"),
         ("orders", "Order service", "process"), ("payments", "Payment service", "process"),
         ("postgres", "PostgreSQL", "data"), ("redis", "Redis", "data"),
         ("objects", "Object storage", "data"), ("stripe", "Stripe", "external")],
        [("web", "gateway"), ("mobile", "gateway"), ("gateway", "catalog"), ("gateway", "orders"),
         ("orders", "payments"), ("payments", "stripe"), ("catalog", "postgres"),
         ("orders", "postgres"),
         ("catalog", "redis"), ("catalog", "objects")],
        [("platform", "Inkwell platform", ["frontend", "backend", "data"]),
         ("frontend", "Frontend", ["web", "mobile"]),
         ("backend", "Backend", ["gateway", "catalog", "orders", "payments"]),
         ("data", "Data", ["postgres", "redis", "objects"])],
        graph=True, max_crossings=2),
    _fx("Agent flow", "agent_flow",
        [("email", "Customer email", "input"), ("classify", "Classify request", "task"),
         ("lookup", "orders_api", "tool"), ("policy", "Refund policy", "document"),
         ("eligible", "Eligible?", "decision"), ("refund", "Issue refund", "action"),
         ("reply", "Draft reply", "task")],
        [("email", "support"), ("classify", "lookup"), ("lookup", "eligible"),
         ("classify", "policy", "", "dotted", None), ("eligible", "refund", "yes"),
         ("eligible", "reply", "no")],
        [("support", "Refund assistant", ["classify", "lookup", "policy", "eligible"])],
        graph=True, max_crossings=0),
    _fx("Event modeling", "event_model",
        [("Interaction", ["ui", "pcr"]), ("Commands and views", ["cmd", "rmo"]),
         ("Events", ["evt"])],
        [(1, "ui", "CartPage", ""), (2, "cmd", "AddBook", "isbn: string, quantity: number"),
         (3, "evt", "BookAdded", ""), (4, "rmo", "CartSummary", ""), (5, "ui", "CheckoutPage", ""),
         (6, "cmd", "PlaceOrder", "cartId: string"), (7, "evt", "OrderPlaced", ""),
         (8, "pcr", "ReceiptSender", ""), (9, "cmd", "SendReceipt", ""),
         (10, "evt", "ReceiptSent", "")]),
    _fx("User journey", "user_journey",
        [("Discover", [("Search for a title", 4, ["Customer"]), ("Read reviews", 3, ["Customer"])]),
         ("Buy", [("Add to cart", 5, ["Customer"]), ("Checkout", 2, ["Customer", "Support"])]),
         ("After purchase",
          [("Track delivery", 3, ["Customer"]), ("Leave a review", 4, ["Customer"])])]),
    _fx("Block diagram", "block",
        [("clients", "Clients", [("web", "Web", ""), ("mobile", "Mobile", ""),
                                 ("partner", "Partner API", "")]),
         ("platform", "Platform", [("gateway", "API gateway", "full"), ("catalog", "Catalog", ""),
                                   ("orders", "Orders", ""), ("payments", "Payments", "")]),
         ("data", "Data", [("postgres", "PostgreSQL", ""), ("redis", "Redis", ""),
                           ("objects", "Object storage", "")])],
        [("clients", "gateway"), ("platform", "data")]),
    _fx("Gantt chart", "gantt",
        [("Design",
          [("UX research", "2027-01-04", "2027-01-15"), ("UI design", "2027-01-11", "2027-01-29")]),
         ("Build",
          [("Backend API", "2027-01-25", "2027-02-26"), ("Frontend", "2027-02-08", "2027-03-12")]),
         ("Ship", [("QA and fixes", "2027-03-01", "2027-03-19")])],
        [("Release", "2027-03-22")]),
    _fx("Timeline", "timeline",
        [("2019", "Founded in Pune"), ("2020", "First 1,000 customers"),
         ("2021", "Mobile app launched"),
         ("2023", "1M books sold"), ("2024", "Series B funding"),
         ("2026", "International launch")]),
    _fx("Mind map", "mind_map", "Inkwell v3",
        [("Discovery", ["Recommendations", "Reviews", "Wishlists"]),
         ("Checkout", ["One-click buy", "Gift cards"]),
         ("Mobile", ["Offline reading", "Push notifications"]),
         ("Operations", ["Observability", "Cost control"])]),
    _fx("Org chart", "org_chart", "Asha Rao \u2014 CEO",
        [("Vikram Shah \u2014 CTO", [("Neha Iyer \u2014 Engineering Manager",
                                      [("Rahul Das \u2014 Dev Lead", []),
                                       ("Meera Joshi \u2014 QA Lead", [])])]),
         ("Priya Menon \u2014 CFO", [("Arjun Nair \u2014 Finance Controller", [])]),
         ("Karan Mehta \u2014 COO", [("Sara Khan \u2014 Operations", []),
                                     ("Dev Patel \u2014 Customer Support", [])])]),
    _fx("Tree view", "tree_view",
        [(0, "inkwell/", "", False), (1, "web/", "", False), (2, "src/", "", False),
         (3, "App.tsx", "React SPA entry", False), (3, "checkout/", "", False), (1, "api/", "",
                                                                                 False),
         (2, "src/", "", False), (3, "orders.ts", "order endpoints", False),
         (3, "payments.ts", "", True), (1, "infra/", "", False), (2, "main.tf", "AWS resources",
                                                                  False),
         (1, "README.md", "", False)]),
    _fx("Quadrant chart", "quadrant",
        [("Search filters", 0.25, 0.80), ("Wishlists", 0.20, 0.45), ("Gift cards", 0.55, 0.70),
         ("Offline mode", 0.85, 0.60), ("Dark theme", 0.15, 0.20),
         ("AI recommendations", 0.80, 0.90),
         ("Loyalty points", 0.65, 0.30)],
        ("Quick wins", "Big bets", "Fill-ins", "Money pits"), ("Effort", "Impact")),
    _fx("Venn diagram", "venn",
        [("fiction", "Fiction"), ("nonfiction", "Non-fiction"), ("kids", "Kids")],
        [(("fiction", "nonfiction"), "Omnivores"), (("fiction", "kids"), "Families"),
         (("nonfiction", "kids"), "Learners"), (("fiction", "nonfiction", "kids"), "Super-fans")],
        title="Inkwell reader segments"),
    _fx("Sankey diagram", "sankey",
        [("Search", "Home", 300), ("Search", "Product page", 200), ("Social", "Home", 120),
         ("Social", "Product page", 80), ("Email", "Product page", 100),
         ("Home", "Product page", 250),
         ("Home", "Exit", 170), ("Product page", "Cart", 280), ("Product page", "Exit", 350),
         ("Cart", "Purchase", 180), ("Cart", "Exit", 100)]),
    _fx("Bar & line chart", "bar_line", MONTHS, ("Revenue", (120, 135, 150, 145, 170, 190)),
        ("Target", (130, 140, 150, 160, 170, 180)), chart_xml=("barChart", "lineChart"),
        title="Revenue vs target, H1 2027 (k$)"),
    _chart("Area chart", "area", MONTHS,
           [("Fiction", (50, 55, 60, 58, 70, 80)), ("Non-fiction", (40, 45, 50, 47, 55, 60)),
            ("Kids", (30, 35, 40, 40, 45, 50))], "Revenue by category, H1 2027 (k$)",
                ("areaChart",)),
    _chart("Pie chart", "pie", ["Search", "Social", "Email", "Direct", "Referral"],
           [("Visits (%)", (45, 25, 15, 10, 5))], "Traffic sources (% of visits)", ("pieChart",)),
    _chart("Doughnut chart", "doughnut", ["Search", "Social", "Email", "Direct", "Referral"],
           [("Visits (%)", (45, 25, 15, 10, 5))], "Traffic sources (% of visits)",
               ("doughnutChart",)),
    _fx("Scatter chart", "scatter",
        [("Orders",
          [(1, 18), (1, 25), (2, 31), (2, 44), (2, 38), (3, 52), (3, 61), (3, 47), (4, 66),
                     (4, 79), (4, 72), (5, 88), (5, 95), (6, 101), (6, 118), (7, 122), (7, 131),
                     (8, 140), (9, 162), (10, 175)])],
        chart_xml=("scatterChart",), title="Order size vs value"),
    _fx("Bubble chart", "bubble",
        [("Fiction", [(26, 540, 14040)]), ("Non-fiction", [(34, 320, 10880)]),
         ("Kids", [(18, 270, 4860)])],
        chart_xml=("bubbleChart",), title="June 2027 by category (bubble size = revenue $)"),
    _fx("Box plot", "box_plot", SAMPLES, chart_xml=("boxWhisker",), chart_type=121,
        title="Order value by channel ($)"),
    _fx("Histogram", "histogram", SAMPLES, 30, chart_xml=("barChart",),
        title="Order value by channel ($)"),
    _fx("Heatmap", "heatmap", ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
        ["00\u201304", "04\u201308", "08\u201312", "12\u201316", "16\u201320", "20\u201324"],
        [[5, 12, 48, 61, 70, 44], [4, 10, 45, 58, 66, 41], [6, 11, 50, 63, 72, 47],
         [5, 13, 52, 65, 78, 50],
         [8, 15, 49, 60, 85, 66], [12, 9, 35, 72, 90, 80], [10, 7, 30, 68, 82, 58]],
        title="Orders by weekday and time of day"),
    _fx("Sunburst chart", "sunburst",
        [("Fiction", [("Fantasy", 120), ("Mystery", 90), ("Romance", 60)]),
         ("Non-fiction", [("Business", 80), ("History", 50), ("Science", 40)]),
         ("Kids", [("Picture books", 45), ("Young adult", 70)])],
        chart_xml=("sunburst",), chart_type=120, title="Revenue by category (k$)"),
    _chart("Radar chart", "radar", ["Price", "Speed", "Reliability", "Support", "Features"],
           [("Inkwell", (4, 3, 5, 2, 4)), ("Competitor", (3, 5, 3, 4, 3))],
           "Reader-app comparison (score 0\u20135)", ("radarChart",)),
    _fx("Candlestick chart", "candlestick",
        ["2027-03-01", "2027-03-02", "2027-03-03", "2027-03-04", "2027-03-05", "2027-03-08",
         "2027-03-09",
         "2027-03-10", "2027-03-11", "2027-03-12"],
        [(100, 104, 98, 103), (103, 107, 101, 106), (106, 108, 102, 104), (104, 105, 99, 100),
         (100, 103, 97, 102), (102, 109, 101, 108), (108, 112, 106, 111), (111, 113, 107, 108),
         (108, 110, 104, 105), (105, 111, 104, 110)],
        chart_xml=("stockChart",), chart_type=89, title="INKW share price, 1\u201312 Mar 2027"),
]


# --------------------------------------------------------------------------------------------
# adversarial layout cases for the graph builders
# --------------------------------------------------------------------------------------------
def _dense(n=40, extra=30, seed=7):
    """A connected graph of n nodes: a random tree plus `extra` random edges (some backwards)."""
    edges = []
    rng = random.Random(seed)
    nodes = [(f"n{i}", f"Step {i}", "process") for i in range(n)]
    for i in range(1, n):
         edges.append((f"n{rng.randrange(i)}", f"n{i}"))
    while len(edges) < n - 1 + extra:
         a, b = rng.sample(range(n), 2)
         if (f"n{a}", f"n{b}") not in edges:
             edges.append((f"n{a}", f"n{b}"))
    return nodes, edges


ADVERSARIAL = [
    {"name": "dense graph (18 nodes, 29 edges)", "builder": "flowchart", "args": _dense(18, 12),
     "kwargs": {}, "max_crossings": 1},
    {"name": "too dense for a slide (40 nodes, 69 edges)", "builder": "flowchart",
     "args": _dense(), "kwargs": {}, "max_crossings": None, "raises": "split the diagram"},
    {"name": "long labels", "builder": "flowchart",
     "args": ([("a", "Receive the customer's order and check that every line item is in stock",
                "start"),
                ("b", "Is the delivery address inside one of the supported shipping regions?",
                 "decision"),
                ("c", "Hand the parcel over to the regional courier with a tracking number",
                 "process"),
                ("d", "Tell the customer which regions we ship to and suggest a pickup point",
                 "error"),
                ("e", "Close the order once the courier confirms the delivery signature", "end")],
               [("a", "b"), ("b", "c", "Yes, the address is supported"),
                ("b", "d", "No, the address is outside every region"), ("c", "e"), ("d", "e")]),
     "kwargs": {}, "max_crossings": 0},
    {"name": "cycles", "builder": "flowchart",
     "args": ([(c, f"State {c.upper()}", "process") for c in "abcdefgh"],
               [("a", "b"), ("b", "c"), ("c", "a", "retry"), ("c", "d"), ("d", "e"), ("e", "f"),
                ("f", "d", "loop"), ("f", "g"), ("g", "g", "poll"), ("g", "h"),
                ("h", "b", "restart")]),
     "kwargs": {}, "max_crossings": 1},
    {"name": "nested groups (3 levels)", "builder": "architecture",
     "args": ([("lb", "Load balancer", "process"), ("api1", "API node 1", "process"),
                ("api2", "API node 2", "process"), ("worker", "Worker", "process"),
                ("db", "Primary DB", "data"), ("replica", "Read replica", "data"),
                ("cache", "Cache", "data"), ("cdn", "CDN", "external")],
               [("cdn", "lb"), ("lb", "api1"), ("lb", "api2"), ("api1", "db"), ("api2", "db"),
                ("api1", "cache"), ("api2", "cache"), ("db", "replica", "replication"),
                ("worker", "db")],
               [("cloud", "Cloud account", ["region"]),
                ("region", "Region eu-west", ["app", "store"]),
                ("app", "App tier", ["lb", "api1", "api2", "worker"]),
                ("store", "Data tier", ["db", "replica", "cache"])]),
     "kwargs": {}, "max_crossings": 1},
]


# --------------------------------------------------------------------------------------------
# primitives: node kinds and link() between every supported shape, side pair and route
# --------------------------------------------------------------------------------------------
def _box(d, s, prst, x, y, w, h, name, text=""):
    return d.styled_node(s, x, y, w, h, text, "neutral", prst=prst, name=f"Node: {name}")


def _kinds_slide(d):
    prev = None
    s = d.slide()
    d.header(s, "Primitives", "Node kinds, fitted text and labelled links", len(d.prs.slides))
    for i, kind in enumerate(sorted(NODE_KINDS)):
         text = {"container": "Container frame",
                 "decision": "Payment OK?"}.get(kind, f"{kind.capitalize()} step")
         w, h, _ = d.fit_shape(text, NODE_KINDS[kind][0], 13, max_w=1.5, min_w=1.3, min_h=0.6)
         sp = d.node(s, 0.9 + (i % 4) * 3.0, 1.95 + (i // 4) * 1.25, w, h, text, kind=kind, size=13)
         if prev is not None and i % 4:
             d.link(s, prev, sp, "next" if i % 2 else "")
         prev = sp


def _presets_slides(d):
    sizes = [(1.3, 0.7), (0.8, 1.1), (1.0, 1.0), (1.5, 0.55)]
    names = sorted(PRESETS)
    for page in range(0, len(names), 15):
         grid = {}
         s = d.slide()
         d.header(s, "Primitives",
                  f"link() between presets {page + 1}-{min(page + 15, len(names))}",
                  len(d.prs.slides))
         for k, prst in enumerate(names[page:page + 15]):
             row, col = divmod(k, 5)
             w, h = sizes[k % len(sizes)]
             grid[row, col] = _box(d, s, prst, 0.9 + col * 2.45 + (1.6 - w) / 2,
                                   1.95 + row * 1.75 + (1.2 - h) / 2, w, h, prst)
         for (row, col), sp in grid.items():
             if (row, col + 1) in grid:
                 d.link(s, sp, grid[row, col + 1])
             if (row + 1, col) in grid:
                 d.link(s, sp, grid[row + 1, col])
             if (row + 1, col + 1) in grid and (row + col) % 2 == 0:
                 d.link(s, grid[row + 1, col + 1], sp, style="dashed", arrow="stealth")


def _sides_slides(d):
    pairs = [(a, b) for a in "trbl" for b in "trbl"]
    for page in range(0, 16, 8):
         s = d.slide()
         d.header(s, "Primitives", f"link() side pairs {page + 1}-{page + 8}", len(d.prs.slides))
         for k, (sa, sb) in enumerate(pairs[page:page + 8]):
             row, col = divmod(k, 4)
             x0, y0 = 0.8 + col * 3.05, 1.9 + row * 2.6
             a = _box(d, s, "roundRect", x0 + 0.35, y0 + 0.35, 0.9, 0.5, f"{sa}{sb} A",
                      f"{sa}{sb} A")
             b = _box(d, s, "rect", x0 + 1.65, y0 + 1.45, 0.9, 0.5, f"{sa}{sb} B", f"{sa}{sb} B")
             d.link(s, a, b, f"{sa} to {sb}", src_site=sa, dst_site=sb)


def _routes_slide(d):
    s = d.slide()
    d.header(s, "Primitives", "Zero-run, curved, straight and via routes; flips, ports, groups",
              len(d.prs.slides))
    a = _box(d, s, "rect", 0.9, 2.0, 1.2, 0.5, "zero A")
    b = _box(d, s, "rect", 0.9, 3.0, 1.2, 0.5, "zero B")
    d.link(s, a, b, src_site="r", dst_site="r")
    d.link(s, b, a, src_site="l", dst_site="l")
    c = _box(d, s, "ellipse", 2.6, 2.0, 0.9, 0.5, "zero C")
    e = _box(d, s, "ellipse", 3.9, 2.0, 0.9, 0.5, "zero E")
    d.link(s, c, e, src_site="b", dst_site="b")
    d.link(s, e, c, src_site="t", dst_site="t", style="dotted")
    f = _box(d, s, "roundRect", 2.6, 3.4, 1.0, 0.5, "curve F")
    g = _box(d, s, "roundRect", 4.4, 4.4, 1.0, 0.5, "curve G")
    d.link(s, f, g, route="curved")
    d.link(s, g, f, route="straight", head="oval", arrow="diamond")
    h = _box(d, s, "flowChartDecision", 6.0, 2.0, 1.2, 0.8, "via H")
    i = _box(d, s, "can", 8.2, 3.2, 1.0, 0.9, "via I")
    d.link(s, h, i, "via", via=[(7.6, 2.4), (7.6, 1.85), (9.6, 1.85), (9.6, 3.65)])
    d.link(s, i, h, src_site=0, dst_site="b")
    j = _box(d, s, "triangle", 6.0, 4.2, 1.1, 0.8, "flip J")
    j._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
    k = _box(d, s, "parallelogram", 8.0, 5.0, 1.2, 0.6, "rot K")
    k.rotation = 30
    d.link(s, j, k)
    d.link(s, k, i)
    hub = d.ports(_box(d, s, "roundRect", 10.4, 4.0, 1.4, 0.8, "ports hub"), 3, 3, 3, 3)
    hexa = d.ports(_box(d, s, "hexagon", 10.4, 5.6, 1.4, 0.7, "ports hex"), 2, 1, 2, 1)
    box = d.ports(_box(d, s, "rect", 10.6, 2.2, 1.0, 0.6, "ports rect"), 2, 2, 2, 2)
    for n, src in enumerate((i, k, g)):
         d.link(s, src, hub, dst_site=("l", n))
    d.link(s, hub, hexa, src_site=("b", 0), dst_site=("t", 0))
    d.link(s, hub, hexa, src_site=("b", 2), dst_site=("t", 1), style="dashed")
    d.link(s, hexa, hub, src_site="r", dst_site=("r", 2))
    d.link(s, box, hub, src_site=("b", 1), dst_site=("t", 1))
    m = _box(d, s, "roundRect", 0.9, 5.4, 1.2, 0.5, "group M")
    n2 = _box(d, s, "flowChartTerminator", 3.0, 6.1, 1.2, 0.5, "group N")
    grp = d.group(s, [m, n2], name="Diagram: grouped nodes", alt="Two grouped nodes and their link")
    d.link(grp, m, n2, "in group")


def primitives_slides(d):
    """Slides that exercise node(kind=), fit_shape(), label() and link() (every preset, side pair
    and route kind); com_check.py must report nothing on them."""
    _kinds_slide(d)
    _presets_slides(d)
    _sides_slides(d)
    _routes_slide(d)


# --------------------------------------------------------------------------------------------
# checker-negative decks: one defect each, and the com_check.py check that must report it
# --------------------------------------------------------------------------------------------
def _neg_overflow(d, s):
    d.node(s, 2.0, 2.5, 1.2, 0.45, "This label is far too long to fit inside its small box",
            kind="process")


def _neg_broken_word(d, s):
    d.node(s, 2.0, 2.5, 0.6, 1.4, "Internationalisation", kind="process")


def _neg_overlap(d, s):
    d.node(s, 2.0, 2.5, 2.0, 0.8, "First node", kind="process")
    d.node(s, 3.2, 2.9, 2.0, 0.8, "Second node", kind="process")


def _neg_container(d, s):
    d.node(s, 1.0, 2.0, 4.0, 2.5, "Group", kind="container", name="Container: Group")
    d.node(s, 4.2, 3.0, 2.0, 0.8, "Half outside", kind="process")


def _neg_unglued(d, s):
    d.node(s, 1.0, 2.5, 1.8, 0.8, "Source", kind="process")
    d.node(s, 5.0, 2.5, 1.8, 0.8, "Target", kind="process")
    d.connector(s, 2.8, 2.9, 5.0, 2.9).name = "Edge: Source -> Target"


def _neg_glue_drift(d, s):
    a = d.node(s, 1.0, 2.5, 1.8, 0.8, "Source", kind="process")
    b = d.node(s, 5.0, 2.5, 1.8, 0.8, "Target", kind="process")
    cn = d.link(s, a, b).connector
    # glued to the top site but drawn at the right one
    cn._element.find(f".//{qn('a:stCxn')}").set("idx", "0")


def _neg_missing_alt(d, s):
    tbl = d.table(s, 1.0, 2.0, 5.0, 1.5, ["Stage", "Owner"], [["Build", "CI"], ["Ship", "Ops"]])
    tbl._graphic_frame._element.find(f".//{qn('p:cNvPr')}").set("descr", "")


def _neg_group(d, s):
    a = d.node(s, 1.0, 2.0, 1.5, 0.6, "Member A", kind="process")
    b = d.node(s, 4.0, 4.0, 1.5, 0.6, "Member B", kind="process")
    xfrm = d.group(s, [a, b], name="Diagram: broken group", alt="A group too small for B") \
         ._element.grpSpPr.find(qn("a:xfrm"))
    for tag in ("a:ext", "a:chExt"):
         xfrm.find(qn(tag)).set("cx", str(Inches(1.5)))
         xfrm.find(qn(tag)).set("cy", str(Inches(0.6)))


def _neg_table_rows(d, s):
    tbl = d.table(s, 1.0, 2.0, 3.0, 0.6, ["Note"],
                   [["This cell holds far more text than its thin row can show"]])
    for row in tbl.rows:
         row.height = Inches(0.3)


def _neg_legend(d, s):
    d.chart(s, "column", 1.0, 2.0, 2.2, 1.6, ["Q1", "Q2"],
             [(f"Series number {i}", (i, i + 1)) for i in range(1, 13)], title="Quarterly series")


def _neg_duplicate_effect_list(d, s):
    sp = d.rect(s, 1.0, 2.0, 2.0, 1.0, fill=d.ACCENT, shadow=True)
    sp._element.spPr.append(sp._element.spPr.makeelement(qn("a:effectLst"), {}))


def _corrupt_slide_xml(path):
    """Rewrite the saved deck with a slide whose XML is not well formed."""
    with zipfile.ZipFile(path) as zin:
         items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
         for info, data in items:
             if info.filename == "ppt/slides/slide1.xml":
                 data = data.replace(b"</p:spTree>", b"</p:spTre>")
             zout.writestr(info, data)


# name: (build(deck, slide), check that must be reported, post-process of the saved file)
NEGATIVE = {
    "broken word": (_neg_broken_word, "text", None),
    "container border": (_neg_container, "overlap", None),
    "glue drift": (_neg_glue_drift, "glue", None),
    "group member outside": (_neg_group, "group", None),
    "invalid XML (second effectLst)": (_neg_duplicate_effect_list, "open", None),
    "legend overflow": (_neg_legend, "text", None),
    "malformed XML": (_neg_overflow, "manifest", _corrupt_slide_xml),
    "missing alt text": (_neg_missing_alt, "alt", None),
    "overlap": (_neg_overlap, "overlap", None),
    "table row overflow": (_neg_table_rows, "text", None),
    "text overflow": (_neg_overflow, "text", None),
    "unglued edge": (_neg_unglued, "glue", None),
}
