"""Self-test of the ppt skill's pptlib; PowerPoint is needed only for --com and --sentinel.

    python selftest.py [--strict] [--deck validation.pptx] [--com] [--sentinel] [--out DIR]

Builds every fixture in fixtures.py whose builder exists and lists the rest as SKIP (--strict fails
on any skip), and asserts: no exceptions, text at 11 pt or more, everything inside the content area,
every connector glued with its ends on the connection sites, the declared chart XML types, the
Python examples in reference/diagrams.md run, and link() refuses a route of more than five segments.

  --deck PATH  write the validation deck (one slide per diagram and chart type)
  --com        also run scripts/com_check.py: no findings on the validation, primitives and
               adversarial decks, and every checker-negative deck reports its defect
  --sentinel   COM-safety test (manual): opens a sentinel presentation in a visible PowerPoint,
               runs com_check.py on another deck, asserts PowerPoint and the sentinel are still
               open, then closes only the sentinel and reports whether PowerPoint is left running
  --out DIR    where decks, PNGs and reports go (default: selftest_out in the current directory)
"""
import json
import math
import re
import subprocess
import sys
import time
import traceback
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path[:0] = [str(SKILL / "scripts"), str(HERE)]

import fixtures  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptlib import Deck  # noqa: E402
from pptlib import deck as deck_module  # noqa: E402
from pptlib.shapes import connection_sites, connector_ends, glue_targets  # noqa: E402
from pptlib.style import CONTENT, MIN_PT  # noqa: E402

COM_CHECK = SKILL / "scripts" / "com_check.py"
DIAGRAMS_MD = SKILL / "reference" / "diagrams.md"
GLUE_TOL = 0.5 / 72  # inches
SIZE_TAGS = (qn("a:defRPr"), qn("a:endParaRPr"), qn("a:rPr"))


class Results:
    def __init__(self):
        self.rows = []

    def add(self, status, name, detail=""):
        self.rows.append((status, name, detail))
        print(f"{status:4s} {name}{': ' + detail if detail else ''}", flush=True)

    def count(self, status):
        return sum(1 for row in self.rows if row[0] == status)


def available(fx):
    """True when the fixture's builder exists (and, for chart(), supports the kind)."""
    if not hasattr(Deck, fx["builder"]):
        return False
    return fx["builder"] != "chart" or fx["args"][0] in deck_module.CHART_KINDS


# --------------------------------------------------------------------------------------------
# checks on what a builder made (no PowerPoint)
# --------------------------------------------------------------------------------------------
def _walk(shapes):
    for sp in shapes:
        yield sp
        if sp.shape_type == 6:  # group
            yield from _walk(sp.shapes)


def _ids(slide):
    return {int(el.get("id")) for el in slide.shapes._spTree.iter(qn("p:cNvPr"))}


def _new_elements(slide, before):
    """Shape-tree children not in `before`, including mc:AlternateContent (Office 2016 charts)."""
    out = []
    for el in slide.shapes._spTree:
        cnv = next(el.iter(qn("p:cNvPr")), None)
        if cnv is not None and int(cnv.get("id")) not in before:
            out.append(el)
    return out


def _chart_roots(slide, elements):
    """Chart XML roots (classic and Office 2016 charts) referenced from `elements`."""
    from lxml import etree
    out = []
    for el in elements:
        for data in el.iter(qn("a:graphicData")):
            rid = data[0].get(qn("r:id")) if len(data) else None
            if rid:
                out.append(etree.fromstring(slide.part.related_part(rid).blob))
    return out


def _xml_box(el):
    """(left, top, right, bottom) in inches of a shape-tree element, from its first xfrm."""
    off = next(el.iter(qn("a:off")), None)
    ext = next(el.iter(qn("a:ext")), None)
    if off is None or ext is None:
        return None
    x, y = int(off.get("x")) / 914400, int(off.get("y")) / 914400
    return x, y, x + int(ext.get("cx")) / 914400, y + int(ext.get("cy")) / 914400


def _name(el):
    cnv = next(el.iter(qn("p:cNvPr")), None)
    return "" if cnv is None else cnv.get("name", "")


def font_problems(slide, elements):
    out = []
    for el in elements:
        for root in [el, *_chart_roots(slide, [el])]:
            sizes = {int(r.get("sz")) / 100 for r in root.iter(*SIZE_TAGS) if r.get("sz")}
            small = sorted(v for v in sizes if v < MIN_PT)
            if small:
                out.append(f"{_name(el)!r}: text at {small[0]:g} pt (minimum {MIN_PT})")
    return out


def area_problems(elements, area=CONTENT, tol=0.02):
    x, y, w, h = area
    out = []
    for el in elements:
        box = _xml_box(el)
        outside = box and (box[0] < x - tol or box[1] < y - tol or box[2] > x + w + tol
                           or box[3] > y + h + tol)
        if outside:
            out.append(f"{_name(el)!r} ({box[0]:.2f}, {box[1]:.2f})-({box[2]:.2f}, {box[3]:.2f}) "
                       "leaves the content area")
    return out


def glue_problems(slide_shapes, shapes):
    """Connectors (except 'Deco:') must be glued at both ends with the ends on the sites."""
    out = []
    by_id = {sp.shape_id: sp for sp in _walk(slide_shapes)}
    for cn in _walk(shapes):
        if cn._element.tag != qn("p:cxnSp") or cn.name.startswith("Deco:"):
            continue
        targets = glue_targets(cn)
        if set(targets) != {"begin", "end"}:
            out.append(f"{cn.name!r} is not glued at both ends")
            continue
        for key, pt in zip(("begin", "end"), connector_ends(cn)):
            sid, idx = targets[key]
            if sid not in by_id:
                out.append(f"{cn.name!r} {key} is glued to missing shape {sid}")
                continue
            site = connection_sites(by_id[sid])[idx]
            drift = math.dist(pt, (site.x, site.y))
            if drift > GLUE_TOL:
                out.append(f"{cn.name!r} {key} is {drift * 72:.2f} pt from its site")
    return out


def chart_problems(slide, elements, expected):
    tags = set()
    for root in _chart_roots(slide, elements):
        tags |= {_local(el) for el in root.iter() if _local(el).endswith("Chart")}
        tags |= {el.get("layoutId") for el in root.iter() if el.get("layoutId")}
    return [f"no {tag} in the chart XML (found {sorted(tags)})" for tag in expected
            if tag not in tags]


def _local(el):
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


def build_checked(results, d, fx, label):
    """Build one fixture on a new slide of deck d and check what it made; returns success."""
    s = d.slide()
    d.header(s, "Self-test", label, len(d.prs.slides), "pptlib self-test")
    before = _ids(s)
    try:
        getattr(d, fx["builder"])(s, *fx["args"], **fx["kwargs"])
    except Exception as e:  # noqa: BLE001
        results.add("FAIL", label, f"{type(e).__name__}: {e}")
        traceback.print_exc(limit=3)
        return False
    elements = _new_elements(s, before)
    made = [sp for sp in s.shapes if sp.shape_id not in before]
    problems = font_problems(s, elements) + area_problems(elements) + glue_problems(s.shapes, made)
    problems += chart_problems(s, elements, fx.get("chart_xml") or ())
    if not elements:
        problems.append("the builder drew nothing")
    results.add("FAIL" if problems else "PASS", label, "; ".join(problems[:4]))
    return not problems


# --------------------------------------------------------------------------------------------
# the parts of the self-test
# --------------------------------------------------------------------------------------------
def test_types(results, d):
    """Build every type fixture whose builder exists into deck d; returns the skipped types."""
    skipped = []
    for fx in fixtures.TYPES:
        if available(fx):
            build_checked(results, d, fx, fx["type"])
        else:
            skipped.append(f"{fx['type']} ({fx['builder']}"
                           f"{' ' + repr(fx['args'][0]) if fx['builder'] == 'chart' else ''})")
    return skipped


def test_adversarial(results, d):
    skipped = []
    for fx in fixtures.ADVERSARIAL:
        if not available(fx):
            skipped.append(f"adversarial: {fx['name']} ({fx['builder']})")
        elif fx.get("raises"):
            label = f"adversarial: {fx['name']} raises {fx['raises']!r}"
            scratch = Deck()
            try:
                getattr(scratch, fx["builder"])(scratch.slide(), *fx["args"], **fx["kwargs"])
            except ValueError as e:
                results.add("PASS" if fx["raises"] in str(e) else "FAIL", label, str(e)[:90])
            else:
                results.add("FAIL", label, "no error raised")
        else:
            build_checked(results, d, fx, f"adversarial: {fx['name']}")
    return skipped


def test_primitives(results, d):
    """The primitives slides: glued connectors on their sites and text at 11 pt or more."""
    try:
        fixtures.primitives_slides(d)
    except Exception as e:  # noqa: BLE001
        results.add("FAIL", "primitives", f"{type(e).__name__}: {e}")
        traceback.print_exc(limit=3)
        return
    for i, s in enumerate(d.prs.slides, 1):
        diagram = [el for el in s.shapes._spTree
                   if _name(el).startswith(("Diagram:", "Edge:", "Label:", "Node:"))]
        problems = glue_problems(s.shapes, list(s.shapes)) + font_problems(s, diagram)
        results.add("FAIL" if problems else "PASS", f"primitives slide {i}",
                    "; ".join(problems[:4]))


def test_route_limit(results):
    """link() must refuse a route of more than five segments with a clear message."""
    d = Deck()
    s = d.slide()
    a = d.node(s, 1.0, 2.0, 1.2, 0.6, "A", kind="process")
    b = d.node(s, 4.0, 4.0, 1.2, 0.6, "B", kind="process")
    via = [(2.5, 2.3), (2.5, 1.8), (3.5, 1.8), (3.5, 3.0), (3.0, 3.0), (3.0, 4.3)]
    try:
        d.link(s, a, b, via=via, dst_site="l")
    except ValueError as e:
        ok = "at most 5" in str(e) and "split the diagram" in str(e)
        results.add("PASS" if ok else "FAIL", "link() refuses a 6+ segment route", str(e)[:90])
        return
    results.add("FAIL", "link() refuses a 6+ segment route", "no error raised")


def test_no_quit(results):
    """COM safety, statically: com_check.py never calls Quit()."""
    calls = re.findall(r"\bQuit\s*\(", COM_CHECK.read_text(encoding="utf-8"))
    results.add("FAIL" if calls else "PASS", "com_check.py contains no Quit( call",
                f"{len(calls)} found" if calls else "")


def test_examples(results):
    """Run every ```python example in reference/diagrams.md."""
    if not DIAGRAMS_MD.exists():
        results.add("SKIP", "reference/diagrams.md examples", "the file does not exist yet")
        return False
    blocks = re.findall(r"```python\n(.*?)```", DIAGRAMS_MD.read_text(encoding="utf-8"), re.S)
    for i, code in enumerate(blocks, 1):
        d = Deck()
        try:
            exec(compile(code, f"diagrams.md example {i}", "exec"),
                 {"Deck": Deck, "d": d, "s": d.slide()})
        except Exception as e:  # noqa: BLE001
            results.add("FAIL", f"diagrams.md example {i}", f"{type(e).__name__}: {e}")
            continue
        results.add("PASS", f"diagrams.md example {i}")
    return True


def build_negatives(out):
    """Save one deck per checker-negative case; returns {name: (path, expected check)}."""
    decks = {}
    for name, (build, check, post) in fixtures.NEGATIVE.items():
        d = Deck()
        path = out / f"negative_{re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')}.pptx"
        s = d.slide()
        d.header(s, "Checker negative", name, 1, "pptlib self-test")
        build(d, s)
        d.save(str(path))
        if post:
            post(path)
        decks[name] = (path, check)
    return decks


# --------------------------------------------------------------------------------------------
# PowerPoint (COM) parts
# --------------------------------------------------------------------------------------------
def run_com_check(deck, out):
    """Run scripts/com_check.py on a deck; returns its JSON report."""
    png = out / f"{deck.stem}_png"
    report = out / f"{deck.stem}_report.json"
    proc = subprocess.run(
        [sys.executable, str(COM_CHECK), str(deck), str(png), "--json", str(report)],
        capture_output=True, text=True)
    if not report.exists():
        raise RuntimeError(f"com_check.py failed on {deck.name}: {proc.stdout[-400:]} "
                           f"{proc.stderr[-400:]}")
    return json.loads(report.read_text(encoding="utf-8"))


def test_com(results, out, clean_decks, negatives):
    for deck in clean_decks:
        rep = run_com_check(deck, out)
        found = [f"slide {f['slide']} [{f['check']}] {f['shape']}: {f['message']}"
                 for f in rep["findings"]]
        results.add("FAIL" if found or rep["status"] != "checked" else "PASS",
                    f"com_check: {deck.name} has no findings ({rep['slides']} slides)",
                    "; ".join(found[:3]))
    for name, (deck, check) in negatives.items():
        rep = run_com_check(deck, out)
        checks = sorted({f["check"] for f in rep["findings"]})
        results.add("PASS" if check in checks else "FAIL",
                    f"com_check reports {name!r} as {check!r}", f"reported {checks}")


def test_sentinel(results, out, deck):
    """COM safety: com_check.py must leave a user's PowerPoint and its open decks alone."""
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    app.Visible = -1
    sentinel = app.Presentations.Add(-1)
    sentinel.Slides.Add(1, 12).Shapes.AddTextbox(1, 40, 40, 600, 60).TextFrame.TextRange.Text = \
        "pptlib self-test sentinel: com_check.py must not close this"
    name = sentinel.Name
    try:
        rep = run_com_check(deck, out)
        alive = _powerpoint_running()
        still_open = any(p.Name == name for p in app.Presentations)
        windows = sentinel.Windows.Count if still_open else 0
        ok = alive and still_open and windows >= 1 and rep["status"] == "checked"
        results.add("PASS" if ok else "FAIL",
                    "sentinel: PowerPoint and the sentinel survive com_check.py",
                    f"PowerPoint running={alive}, sentinel open={still_open}, windows={windows}")
    finally:
        sentinel.Saved = True
        sentinel.Close()
        sentinel = app = None
    time.sleep(3)
    results.add("INFO", "after closing only the sentinel",
                f"PowerPoint still running: {_powerpoint_running()}")


def _powerpoint_running():
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/NH"],
                         capture_output=True, text=True).stdout
    return "POWERPNT.EXE" in out.upper()


def main(argv):
    deck_path = Path(argv[argv.index("--deck") + 1]).resolve() if "--deck" in argv else None
    out = Path(argv[argv.index("--out") + 1] if "--out" in argv else "selftest_out").resolve()
    results = Results()
    types_deck, primitives_deck, adversarial_deck = Deck(), Deck(), Deck()
    skipped = test_types(results, types_deck) + test_adversarial(results, adversarial_deck)
    test_primitives(results, primitives_deck)
    test_route_limit(results)
    test_no_quit(results)
    test_examples(results)
    for name in skipped:
        results.add("SKIP", name, "builder not available yet")
    out.mkdir(parents=True, exist_ok=True)
    clean = []
    for deck, path in ((types_deck, deck_path or out / "validation.pptx"),
                       (primitives_deck, out / "primitives.pptx"),
                       (adversarial_deck, out / "adversarial.pptx")):
        if len(deck.prs.slides):
            deck.save(str(path))
            clean.append(path)
    if deck_path:
        print(f"validation deck: {deck_path} ({len(types_deck.prs.slides)} slides)")
    if "--com" in argv:
        test_com(results, out, clean, build_negatives(out))
    if "--sentinel" in argv:
        test_sentinel(results, out, out / "primitives.pptx")
    fails, skips = results.count("FAIL"), results.count("SKIP")
    print(f"\n{results.count('PASS')} passed, {fails} failed, {skips} skipped"
          f"{' (--strict: skips fail)' if skips and '--strict' in argv else ''}")
    if skipped:
        print("skipped types:", ", ".join(skipped))
    return 1 if fails or (skips and "--strict" in argv) else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    sys.exit(main(sys.argv))
