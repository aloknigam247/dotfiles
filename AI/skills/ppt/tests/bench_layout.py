"""Layout benchmark and crossing gate for pptlib's layout engine (pptlib/layout.py).

    python bench_layout.py [--json report.json]

Runs every case in a fresh Python process and times layout (layered()) and routing (route())
separately: 20 fixed seeds of random 40-node graphs (long edges, back edges, self-loops) and
adversarial compound graphs (nested groups, edges across them) at 20 and 40 nodes. Gate, on this
machine: 40-node p95 <= 1 s and max <= 2 s (layout + routing), peak working set <= 250 MB.

Crossings are counted on the final routed segments: for the cases above on the routes, and for the
graph fixtures in fixtures.py (whose builder exists) on the connectors of the built slide, with
shared end points not counted and filled shapes left out. Each fixture must stay within its
committed `max_crossings`.

Exit status: 0 when every gate passes, 1 otherwise.
"""
import json
import math
import random
import statistics
import subprocess
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path[:0] = [str(HERE.parent / "scripts"), str(HERE)]

GATE_MAX_S = 2.0
GATE_P95_S = 1.0
GATE_WS_MB = 250
SEEDS = range(20)


def random_graph(seed, n=40):
    """A random connected graph: every node hangs off an earlier one, plus extra edges (some long,
    some backwards) and a self-loop; node sizes vary like real labels."""
    edges = []
    rng = random.Random(seed)
    sizes = {f"n{i}": (round(0.8 + rng.random() * 1.2, 2), rng.choice((0.42, 0.42, 0.6)))
             for i in range(n)}
    for v in range(1, n):
        edges.append((f"n{rng.randrange(max(0, v - 8), v)}", f"n{v}"))
    while len(edges) < int(n * 1.5):
        a, b = rng.sample(range(n), 2)
        if (f"n{a}", f"n{b}") not in edges:
            edges.append((f"n{a}", f"n{b}"))
    edges.append((f"n{n // 2}", f"n{n // 2}"))
    return {"sizes": sizes, "edges": edges, "groups": {}}


def compound_graph(seed, n):
    """Adversarial compound graph: groups nested three levels deep (leaf groups of four in pairs,
    all in one outer group), a hub with its spokes in every leaf, hubs fanning out to the next
    leaves, spokes wired into later leaves and one long edge back, and outside nodes wired into the
    deepest groups."""
    edges, groups, leaves = [], {}, []
    rng = random.Random(1000 + seed)
    ids = [f"c{i}" for i in range(n)]
    sizes = {i: (round(0.9 + rng.random() * 0.8, 2), 0.45) for i in ids}
    outside, inner = ids[:max(2, n // 10)], ids[max(2, n // 10):]
    for j in range(0, len(inner), 4):
        leaves.append(f"leaf{j // 4}")
        groups[leaves[-1]] = inner[j:j + 4]
    for k in range(0, len(leaves), 2):
        groups[f"mid{k // 2}"] = leaves[k:k + 2]
    groups["outer"] = [g for g in groups if g.startswith("mid")]
    hubs = [groups[g][0] for g in leaves]
    spokes = [m for g in leaves for m in groups[g][1:]]
    for g in leaves:
        edges += [(groups[g][0], m) for m in groups[g][1:]]
    for k, hub in enumerate(hubs[1:], 1):
        edges.append((hubs[(k - 1) // 2], hub))
    for _ in range(len(leaves)):
        a, b = sorted(rng.sample(range(len(leaves)), 2))
        edges.append((rng.choice(groups[leaves[a]][1:]), hubs[b]))
    edges.append((groups[leaves[-1]][-1], groups[leaves[0]][-1]))
    edges += [(o, rng.choice(spokes)) for o in outside]
    edges.append((rng.choice(spokes), outside[0]))
    return {"sizes": sizes, "edges": list(dict.fromkeys(edges)), "groups": groups}


def cases():
    out = [("random", 40, s) for s in SEEDS]
    out += [("compound", n, s) for n in (20, 40) for s in range(5)]
    return out


def graph(kind, n, seed):
    return random_graph(seed, n) if kind == "random" else compound_graph(seed, n)


def peak_working_set_mb():
    """Peak working set of this process in MB (Windows), else the peak resident set size."""
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD),
                        ("PeakWorkingSetSize", ctypes.c_size_t),
                        ("WorkingSetSize", ctypes.c_size_t),
                        ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                        ("QuotaPagedPoolUsage", ctypes.c_size_t),
                        ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                        ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                        ("PagefileUsage", ctypes.c_size_t), ("PeakPagefileUsage", ctypes.c_size_t)]

        counters = Counters()
        counters.cb = ctypes.sizeof(Counters)
        current = ctypes.windll.kernel32.GetCurrentProcess
        current.restype = wintypes.HANDLE
        info = ctypes.windll.psapi.GetProcessMemoryInfo
        info.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
        if not info(current(), ctypes.byref(counters), counters.cb):
            raise OSError("GetProcessMemoryInfo failed")
        return counters.PeakWorkingSetSize / 2 ** 20
    import resource
    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024


def run_one(kind, n, seed):
    """Lay out and route one case in this process; returns its measurements."""
    from pptlib import layout
    g = graph(kind, n, seed)
    t0 = time.perf_counter()
    lay = layout.layered(g["sizes"], g["edges"], direction=("TB", "LR")[seed % 2],
                         groups=g["groups"])
    t1 = time.perf_counter()
    routes = layout.route(lay)
    t2 = time.perf_counter()
    return {"case": f"{kind}-{n}-{seed}", "crossings": layout.route_crossings(
        [r.points for r in routes]), "edges": len(g["edges"]), "kind": kind,
        "layout_s": t1 - t0, "nodes": n, "route_s": t2 - t1,
        "segments": max(len(r.points) - 1 for r in routes), "ws_mb": peak_working_set_mb()}


# --------------------------------------------------------------------------------------------
# crossings of the graph fixtures, on the connectors of the built slide
# --------------------------------------------------------------------------------------------
def connector_points(cn):
    """The polyline of a straight or bent connector on the slide, in inches."""
    from pptx.oxml.ns import qn
    from pptlib.shapes import EMU_IN, _local_to_slide
    geom = cn._element.spPr.find(qn("a:prstGeom"))
    prst = geom.get("prst")
    adj = {gd.get("name"): int(gd.get("fmla").split()[-1]) / 100000
           for gd in geom.iterfind(f"{qn('a:avLst')}/{qn('a:gd')}")}
    w, h = float(cn.width), float(cn.height)
    a1, a2, a3 = adj.get("adj1", 0.5), adj.get("adj2", 0.5), adj.get("adj3", 0.5)
    local = {"bentConnector2": [(0, 0), (w, 0), (w, h)],
             "bentConnector3": [(0, 0), (a1 * w, 0), (a1 * w, h), (w, h)],
             "bentConnector4": [(0, 0), (a1 * w, 0), (a1 * w, a2 * h), (w, a2 * h), (w, h)],
             "bentConnector5": [(0, 0), (a1 * w, 0), (a1 * w, a2 * h), (a3 * w, a2 * h),
                                (a3 * w, h), (w, h)]}.get(prst, [(0, 0), (w, h)])
    out = []
    for x, y in local:
        sx, sy, _ = _local_to_slide(cn._element, x, y)
        out.append((round(sx / EMU_IN, 4), round(sy / EMU_IN, 4)))
    return out


def fixture_crossings():
    """[(name, crossings, max_crossings)] for every graph fixture whose builder exists."""
    import fixtures
    from pptlib import Deck, layout
    out = []
    for fx in [f for f in fixtures.TYPES if f["graph"]] + fixtures.ADVERSARIAL:
        name = fx.get("type") or f"adversarial: {fx['name']}"
        if not hasattr(Deck, fx["builder"]):
            out.append((name, None, fx["max_crossings"]))
            continue
        d = Deck()
        s = d.slide()
        getattr(d, fx["builder"])(s, *fx["args"], **fx["kwargs"])
        lines = [connector_points(sp) for sp in _walk(s.shapes)
                 if sp._element.tag.endswith("cxnSp") and sp.name.startswith("Edge:")]
        out.append((name, layout.route_crossings(lines), fx["max_crossings"]))
    return out


def _walk(shapes):
    for sp in shapes:
        yield sp
        if sp.shape_type == 6:
            yield from _walk(sp.shapes)


# --------------------------------------------------------------------------------------------
def _pct(values, q):
    values = sorted(values)
    return values[min(len(values) - 1, math.ceil(q * len(values)) - 1)]


def main(argv):
    if "--one" in argv:
        kind, n, seed = argv[argv.index("--one") + 1].split(":")
        print(json.dumps(run_one(kind, int(n), int(seed))))
        return 0
    failures = []
    rows = []
    for kind, n, seed in cases():
        proc = subprocess.run([sys.executable, __file__, "--one", f"{kind}:{n}:{seed}"],
                              capture_output=True, text=True)
        if proc.returncode:
            failures.append(f"{kind}-{n}-{seed}: {proc.stderr.strip().splitlines()[-1]}")
            continue
        rows.append(json.loads(proc.stdout.strip().splitlines()[-1]))
    print(f"{'set':14s} {'runs':>4s} {'layout p50/p95/max s':>22s} {'route p50/p95/max s':>21s} "
          f"{'total p95/max s':>16s} {'peak MB':>8s} {'crossings med/max':>18s}")
    for (kind, n) in sorted({(r["kind"], r["nodes"]) for r in rows}):
        sel = [r for r in rows if r["kind"] == kind and r["nodes"] == n]
        lay = [r["layout_s"] for r in sel]
        rte = [r["route_s"] for r in sel]
        tot = [r["layout_s"] + r["route_s"] for r in sel]
        cross = [r["crossings"] for r in sel]
        print(f"{kind + '-' + str(n):14s} {len(sel):4d} "
              f"{statistics.median(lay):6.3f}/{_pct(lay, .95):6.3f}/{max(lay):6.3f} "
              f"{statistics.median(rte):6.3f}/{_pct(rte, .95):6.3f}/{max(rte):6.3f} "
              f"{_pct(tot, .95):7.3f}/{max(tot):7.3f} {max(r['ws_mb'] for r in sel):8.1f} "
              f"{statistics.median(cross):8.0f}/{max(cross):5d}")
    big = [r["layout_s"] + r["route_s"] for r in rows if r["nodes"] == 40]
    if big and _pct(big, .95) > GATE_P95_S:
        failures.append(f"40-node p95 {_pct(big, .95):.3f} s > {GATE_P95_S} s")
    if big and max(big) > GATE_MAX_S:
        failures.append(f"40-node max {max(big):.3f} s > {GATE_MAX_S} s")
    if rows and max(r["ws_mb"] for r in rows) > GATE_WS_MB:
        failures.append(f"peak working set {max(r['ws_mb'] for r in rows):.0f} MB > "
                        f"{GATE_WS_MB} MB")
    if any(r["segments"] > 5 for r in rows):
        failures.append("a route has more than 5 segments")
    print("\nfixture crossings (final connectors):")
    fixture_rows = fixture_crossings()
    for name, found, limit in fixture_rows:
        if found is None:
            print(f"  SKIP {name}: builder not available yet")
            continue
        ok = limit is not None and found <= limit
        print(f"  {'PASS' if ok else 'FAIL'} {name}: {found} crossing(s), max {limit}")
        if not ok:
            failures.append(f"{name}: {found} crossings > max_crossings {limit}")
    if "--json" in argv:
        Path(argv[argv.index("--json") + 1]).write_text(json.dumps(
            {"cases": rows, "fixtures": fixture_rows, "failures": failures}, indent=1),
            encoding="utf-8")
    for f in failures:
        print("FAIL", f)
    print("PASS: all gates met" if not failures else f"{len(failures)} gate(s) failed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    sys.exit(main(sys.argv))
