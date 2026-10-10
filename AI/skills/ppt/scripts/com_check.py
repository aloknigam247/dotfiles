"""Verify a .pptx in PowerPoint through COM (Windows): open it read-only without a window, export
every slide to PNG and report layout findings.

    python com_check.py deck.pptx png_dir [--json report.json]

Checks:
  manifest  the deck's canonical manifest (tests/manifest.py) survives a PowerPoint save-copy, which
            catches parts PowerPoint silently repairs or drops; malformed XML is reported without
            opening PowerPoint
  open      PowerPoint opens the deck
  text      text stays inside its shape and no word is broken across lines: shapes, table cells
            (rows must not grow), chart titles, axis titles, legends and data labels
  overlap   'Node:' and 'Label:' shapes do not partly overlap each other and nodes do not cross a
            'Container:' border; allowed: overlaps inside a Venn diagram, labels on containers
  group     every group contains its members
  glue      'Edge:' connectors are glued at both ends, and every glued end lies within 1 pt of the
            position PowerPoint gives its connection site (measured with a probe connector)
  alt       charts, tables, pictures and groups have alt text
  chart     PowerPoint reports the chart type the chart XML declares

COM safety: the checker works on a copy of the deck inside png_dir, opens it read-only and
windowless, closes only that presentation and never quits PowerPoint, so a PowerPoint the user is
working in, and every deck open in it, is left alone.
Exit status: 0 no findings, 1 findings, 2 the deck could not be checked.
"""
import json
import math
import os
import shutil
import sys
import time
from pathlib import Path

from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path[:0] = [str(HERE), str(HERE.parent / "tests")]
import manifest as manifest_mod  # noqa: E402
from pptlib.shapes import natural_width  # noqa: E402

# overlaps allowed inside a diagram group named "Diagram: <builder> ...", by shape-name kinds
ALLOW_OVERLAP = {"venn": {("Label", "Node"), ("Node", "Label"), ("Node", "Node")}}
CHARTEX_TYPES = {"boxWhisker": 121, "clusteredColumn": 118, "funnel": 123, "paretoLine": 122,
                 "regionMap": 140, "sunburst": 120, "treemap": 117, "waterfall": 119}
DRIFT_PT = 1.0
MSO_FALSE, MSO_TRUE = 0, -1
MSO_GROUP, MSO_PICTURE = 6, 13
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "cx": "http://schemas.microsoft.com/office/drawing/2014/chartex",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
OVERLAP_SHARE = 0.02  # of the smaller shape's area
PIE_TYPES = {5, 69, -4120, 80, 68, 71}
TOL_PT = 1.5


# --------------------------------------------------------------------------------------------
# what the deck XML declares: table row heights, chart plots and legend entries
# --------------------------------------------------------------------------------------------
def _local(el):
    return etree.QName(el).localname


def _chart_info(root):
    """Plots ([(tag, options)]), expected legend entries and deleted legend entries of a chart."""
    plots = []
    for plot in root.iterfind("c:chart/c:plotArea/*", NS):
        if _local(plot).endswith("Chart"):
            opts = {_local(el): el.get("val") for el in plot if el.get("val") is not None}
            sers = plot.findall("c:ser", NS)
            points = max((int(pc.get("val")) for pc in plot.iterfind("c:ser/c:val//c:ptCount", NS)),
                         default=0)
            plots.append((_local(plot), {**opts, "series": len(sers), "points": points}))
    for ser in root.iterfind(".//cx:plotAreaRegion/cx:series", NS):
        plots.append((ser.get("layoutId"), {"chartex": True}))
    deleted = len(root.findall("c:chart/c:legend/c:legendEntry/c:delete[@val='1']", NS))
    return {"deleted_legend": deleted, "plots": plots}


def _expected_types(plots):
    """COM ChartType values PowerPoint may report for these plots (None: not checked)."""
    out = set()
    for tag, opts in plots:
        if opts.get("chartex"):
            if tag not in CHARTEX_TYPES:
                return None
            out.add(CHARTEX_TYPES[tag])
            continue
        grouping = opts.get("grouping", "standard")
        if tag == "barChart":
            col = opts.get("barDir", "col") == "col"
            base = {"clustered": 51, "stacked": 52, "percentStacked": 53}.get(grouping, 51)
            out.add(base + (0 if col else 6))
        elif tag == "lineChart":
            out |= {"stacked": {63, 66}, "percentStacked": {64, 67}}.get(grouping, {4, 65})
        elif tag == "areaChart":
            out.add({"stacked": 76, "percentStacked": 77}.get(grouping, 1))
        elif tag in ("pieChart", "doughnutChart", "ofPieChart"):
            out |= {"pieChart": {5, 69}, "doughnutChart": {-4120, 80}, "ofPieChart": {68, 71}}[tag]
        elif tag == "scatterChart":
            out |= {-4169, 72, 73, 74, 75}
        elif tag == "bubbleChart":
            out |= {15, 87}
        elif tag == "radarChart":
            out |= {82} if opts.get("radarStyle") == "filled" else {-4151, 81}
        elif tag == "stockChart":
            out |= {88, 89, 90, 91}
        else:
            return None
    if len(plots) > 1:
        out.add(-4111)  # xlCombination
    return out or None


def deck_model(path):
    """{slide index: {shape id: info}} from the deck XML, for tables and charts."""
    from pptx import Presentation
    model = {}
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        info = model.setdefault(i, {})
        for frame in slide.shapes._spTree.iter(f"{{{NS['p']}}}graphicFrame"):
            cnv = frame.find(".//p:cNvPr", NS)
            data = frame.find(".//a:graphicData", NS)
            if cnv is None or data is None or not len(data):
                continue
            sid = int(cnv.get("id"))
            if data.get("uri", "").endswith("/table"):
                rows = data.iter(f"{{{NS['a']}}}tr")
                info[sid] = {"rows": [int(tr.get("h")) / 12700 for tr in rows]}
            elif data[0].get(f"{{{NS['r']}}}id"):
                part = slide.part.related_part(data[0].get(f"{{{NS['r']}}}id"))
                info[sid] = _chart_info(etree.fromstring(part.blob))
    return model


# --------------------------------------------------------------------------------------------
# COM helpers
# --------------------------------------------------------------------------------------------
def _prop(obj, name, *args):
    """A parameterised COM property (e.g. TextRange2.Lines(Start, Length)) through IDispatch."""
    import pythoncom
    import win32com.client
    dispid = obj._oleobj_.GetIDsOfNames(name)
    res = obj._oleobj_.Invoke(dispid, 0, pythoncom.DISPATCH_PROPERTYGET, True, *args)
    return win32com.client.Dispatch(res) if type(res).__name__ == "PyIDispatch" else res


def _lines(tr):
    out = []
    for i in range(1, 200):
        try:
            line = _prop(tr, "Lines", i, 1)
        except Exception:  # noqa: BLE001  (past the last line)
            break
        if not line.Length:
            break
        out.append(line.Text)
    return out


def _broken(text, lines):
    """Words broken across the given line split of `text`."""
    out = []
    pos = 0
    for line in lines[:-1]:
        pos += len(line)
        if line and not line[-1].isspace() and line[-1] not in "-\u2013\u2014/" \
                and pos < len(text) and not text[pos].isspace():
            out.append(line.split()[-1] if line.split() else line)
    return out


def _box(shp):
    return shp.Left, shp.Top, shp.Left + shp.Width, shp.Top + shp.Height


def _inter(a, b):
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return w * h if w > 0 and h > 0 else 0.0


def _contains(a, b, tol=1.0):
    return a[0] - tol <= b[0] and a[1] - tol <= b[1] and a[2] + tol >= b[2] and a[3] + tol >= b[3]


def _area(b):
    return max(b[2] - b[0], 0.0) * max(b[3] - b[1], 0.0)


def _ends(cn):
    """Begin and end points (pt) of a connector, honouring flips and rotation of its frame."""
    w, h = cn.Width, cn.Height
    cx, cy = cn.Left + w / 2, cn.Top + h / 2
    th = math.radians(cn.Rotation)
    out = []
    for lx, ly in ((0, 0), (w, h)):
        if cn.HorizontalFlip:
            lx = w - lx
        if cn.VerticalFlip:
            ly = h - ly
        x, y = lx - w / 2, ly - h / 2
        out.append((cx + x * math.cos(th) - y * math.sin(th),
                    cy + x * math.sin(th) + y * math.cos(th)))
    return out


class _Probe:
    """A straight connector added to the slide in memory (the deck is never saved): glued to a
    connection site, its begin point is where PowerPoint puts that site."""

    def __init__(self, slide):
        self.slide = slide
        self.shape = None

    def site(self, shp, site):
        if self.shape is None:
            self.shape = self.slide.Shapes.AddConnector(1, 0, 0, 3, 3)
        self.shape.ConnectorFormat.BeginConnect(shp, site)
        return _ends(self.shape)[0]

    def remove(self):
        if self.shape is not None:
            self.shape.Delete()
            self.shape = None


def _walk(shapes, tag=None):
    """(shape, builder tag of the enclosing 'Diagram: <builder>' group) for every shape."""
    for shp in shapes:
        yield shp, tag
        if shp.Type == MSO_GROUP:
            name = shp.Name
            inner = name[len("Diagram: "):].split()[0].lower() if name.startswith("Diagram: ") and \
                len(name) > len("Diagram: ") else tag
            yield from _walk(shp.GroupItems, inner)


# --------------------------------------------------------------------------------------------
# checks; each returns [(check, shape name, message)]
# --------------------------------------------------------------------------------------------
def text_findings(shp):
    """Text outside its shape, or a word broken across lines."""
    out = []
    if not shp.HasTextFrame or shp.Connector or shp.Rotation:
        return out
    tf = shp.TextFrame2
    if not tf.HasText:
        return out
    tr = tf.TextRange
    size = tr.Font.Size if 0 < tr.Font.Size < 400 else 18.0
    # BoundWidth runs about 0.22 x font size past the last glyph (measured), hence the tolerance
    tol_x = TOL_PT + 0.25 * size
    left, top, right, bottom = _box(shp)
    bl, bt = tr.BoundLeft, tr.BoundTop
    br, bb = bl + tr.BoundWidth, bt + tr.BoundHeight
    if bl < left - tol_x or br > right + tol_x or bt < top - TOL_PT or bb > bottom + TOL_PT:
        out.append(("text", shp.Name, f"text {br - bl:.0f} x {bb - bt:.0f} pt overflows its "
                                      f"{right - left:.0f} x {bottom - top:.0f} pt shape"))
    for word in _broken(tr.Text, _lines(tr)):
        out.append(("text", shp.Name, f"word broken across lines after {word!r}"))
    return out


def table_findings(shp, info):
    """Broken words in table cells, and rows PowerPoint had to grow because text did not fit."""
    out = []
    tbl = shp.Table
    declared = (info or {}).get("rows", [])
    for r in range(1, tbl.Rows.Count + 1):
        height = tbl.Rows(r).Height
        if r <= len(declared) and height > declared[r - 1] + TOL_PT:
            out.append(("text", shp.Name, f"row {r} grew from {declared[r - 1]:.0f} to "
                                          f"{height:.0f} pt: "
                                          "its text does not fit"))
        for c in range(1, tbl.Columns.Count + 1):
            tr = tbl.Cell(r, c).Shape.TextFrame2.TextRange
            for word in _broken(tr.Text, _lines(tr)):
                out.append(("text", shp.Name, f"cell ({r}, {c}): word broken after {word!r}"))
    return out


def _element_findings(name, what, box, chart_box, text, size, font):
    """A chart element (title, legend, data label) outside the chart or narrower than a word."""
    out = []
    if not _contains(chart_box, box, TOL_PT):
        out.append(("text", name, f"chart {what} lies outside the chart"))
    for word in (text or "").split():
        try:
            need = natural_width(word, size, font) * 0.93  # the narrowest PowerPoint/Pillow ratio
        except RuntimeError:
            break
        if need > box[2] - box[0] + TOL_PT:
            out.append(("text", name, f"chart {what}: {word!r} is wider than its "
                                      f"{box[2] - box[0]:.0f} pt box"))
    return out


def chart_findings(shp, info, font):
    """Chart title, axis titles, legend and data labels inside the chart, unbroken and not
    overlapping; and the chart type PowerPoint reports against the type the XML declares."""
    labels = []
    out = []
    ch = shp.Chart
    chart_box = (0.0, 0.0, shp.Width, shp.Height)
    expected = _expected_types((info or {}).get("plots", []))
    try:
        chart_type = ch.ChartType
    except Exception:  # noqa: BLE001
        chart_type = None
    if expected and chart_type not in expected:
        out.append(("chart", shp.Name, f"PowerPoint reports chart type {chart_type}; the XML "
                                       "declares "
                                       f"one of {sorted(expected)}"))
    elements = []
    try:
        if ch.HasTitle:
            elements.append(("title", ch.ChartTitle))
        for axis_type, what in ((1, "category axis title"), (2, "value axis title")):
            try:
                axis = ch.Axes(axis_type)
                if axis.HasTitle:
                    elements.append((what, axis.AxisTitle))
            except Exception:  # noqa: BLE001  (no such axis, e.g. pie charts)
                pass
    except Exception:  # noqa: BLE001  (Office 2016 charts expose only part of the object model)
        elements = []
    for what, el in elements:
        try:
            box = (el.Left, el.Top, el.Left + el.Width, el.Top + el.Height)
            size = el.Format.TextFrame2.TextRange.Font.Size
        except Exception:  # noqa: BLE001
            continue
        out += _element_findings(shp.Name, what, box, chart_box, el.Text, size if size > 0 else 14,
                                 font)
    try:
        has_legend = ch.HasLegend
    except Exception:  # noqa: BLE001
        has_legend = False
    if has_legend:
        out += _legend_findings(shp, ch, chart_box, info)
    try:
        series = [ch.SeriesCollection(i) for i in range(1, ch.SeriesCollection().Count + 1)]
    except Exception:  # noqa: BLE001
        series = []
    for ser in series:
        try:
            if not ser.HasDataLabels:
                continue
            points = ser.Points().Count
        except Exception:  # noqa: BLE001
            continue
        for p in range(1, points + 1):
            dl = ser.Points(p).DataLabel
            box = (dl.Left, dl.Top, dl.Left + dl.Width, dl.Top + dl.Height)
            out += _element_findings(shp.Name, f"data label {dl.Text!r}", box, chart_box, dl.Text,
                                     dl.Format.TextFrame2.TextRange.Font.Size or 12, font)
            labels.append((dl.Text, box))
    for i, (ta, a) in enumerate(labels):
        for tb, b in labels[i + 1:]:
            if _inter(a, b) > OVERLAP_SHARE * min(_area(a), _area(b)):
                out.append(("text", shp.Name, f"data labels {ta!r} and {tb!r} overlap"))
    return out


def _legend_findings(shp, ch, chart_box, info):
    """A legend inside the chart whose entries all fit (PowerPoint stacks entries that don't)."""
    out = []
    lg = ch.Legend
    legend_box = (lg.Left, lg.Top, lg.Left + lg.Width, lg.Top + lg.Height)
    if not _contains(chart_box, legend_box, TOL_PT):
        out.append(("text", shp.Name, "chart legend lies outside the chart"))
    entries = []
    for i in range(1, lg.LegendEntries().Count + 1):
        e = lg.LegendEntries(i)
        entries.append((e.Left, e.Top, e.Left + e.Width, e.Top + e.Height))
    hidden = [i for i, e in enumerate(entries, 1) if not _contains(legend_box, e, TOL_PT)]
    stacked = [i for i, e in enumerate(entries, 1)
               if any(_inter(e, f) > 0.2 * _area(e) for j, f in enumerate(entries, 1) if j != i)]
    if hidden or stacked:
        out.append(("text", shp.Name,
                    f"chart legend shows only part of its {len(entries)} entries"))
    return out


def overlap_findings(items):
    """Partial overlaps between named diagram shapes. items: [(name, box, builder tag)]."""
    out = []
    kinds = [(name.split(":", 1)[0], name, box, tag) for name, box, tag in items]
    for i, (ka, na, a, ta) in enumerate(kinds):
        for kb, nb, b, tb in kinds[i + 1:]:
            area = _inter(a, b)
            if not area or _contains(a, b) or _contains(b, a):
                continue
            pair = {ka, kb}
            if "Container" in pair:
                if pair == {"Container", "Label"}:
                    continue
                inner = nb if ka == "Container" else na
                outer = na if ka == "Container" else nb
                out.append(("overlap", inner, f"crosses the border of {outer!r}"))
                continue
            if ta is not None and ta == tb and (ka, kb) in ALLOW_OVERLAP.get(ta, ()):
                continue
            if area > OVERLAP_SHARE * min(_area(a), _area(b)):
                share = 100 * area / min(_area(a), _area(b))
                out.append(("overlap", na, f"overlaps {nb!r} ({share:.0f}%)"))
    return out


def group_findings(shp):
    """Members a group does not contain."""
    gbox = _box(shp)
    return [("group", shp.Name, f"member {it.Name!r} lies outside the group")
            for it in shp.GroupItems if not _contains(gbox, _box(it), TOL_PT)]


def glue_findings(probe, shp):
    """Unglued 'Edge:' connectors and glued ends away from their connection site."""
    out = []
    cf = shp.ConnectorFormat
    begin, end = _ends(shp)
    if shp.Name.startswith("Edge:") and not (cf.BeginConnected and cf.EndConnected):
        out.append(("glue", shp.Name, "edge is not glued at both ends"))
    ends = []
    if cf.BeginConnected:
        ends.append(("begin", begin, cf.BeginConnectedShape, cf.BeginConnectionSite))
    if cf.EndConnected:
        ends.append(("end", end, cf.EndConnectedShape, cf.EndConnectionSite))
    for which, pt, target, site in ends:
        drift = math.dist(pt, probe.site(target, site))
        if drift > DRIFT_PT:
            out.append(("glue", shp.Name,
                        f"{which} is {drift:.1f} pt away from site {site} of {target.Name!r}"))
    return out


def alt_findings(shp):
    try:
        needs = shp.Type in (MSO_GROUP, MSO_PICTURE) or shp.HasChart or shp.HasTable
    except Exception:  # noqa: BLE001
        needs = False
    if needs and not (shp.AlternativeText or "").strip():
        return [("alt", shp.Name, "no alt text")]
    return []


def check_slide(slide, model, font):
    """All checks on one slide; returns (findings, stats)."""
    findings = []
    items = []
    stats = {"charts": [], "connectors": 0, "glued": 0, "shapes": slide.Shapes.Count}
    probe = _Probe(slide)
    connectors = []
    for shp, tag in _walk(slide.Shapes):
        name = shp.Name
        if shp.Connector:
            connectors.append(shp)
            continue
        findings += alt_findings(shp)
        if shp.Type == MSO_GROUP:
            findings += group_findings(shp)
            continue
        if shp.HasTable:
            findings += table_findings(shp, model.get(shp.Id))
        elif shp.HasChart:
            findings += chart_findings(shp, model.get(shp.Id), font)
            try:
                stats["charts"].append({"name": name, "chart_type": shp.Chart.ChartType})
            except Exception as e:  # noqa: BLE001
                stats["charts"].append({"name": name, "error": str(e)[:80]})
        else:
            findings += text_findings(shp)
        if name.split(":", 1)[0] in ("Container", "Label", "Node") and ":" in name:
            items.append((name, _box(shp), tag))
    try:
        for shp in connectors:
            cf = shp.ConnectorFormat
            stats["connectors"] += 1
            stats["glued"] += bool(cf.BeginConnected and cf.EndConnected)
            findings += glue_findings(probe, shp)
    finally:
        probe.remove()
    findings += overlap_findings(items)
    return findings, stats


# --------------------------------------------------------------------------------------------
# the deck
# --------------------------------------------------------------------------------------------
def _theme_font(path):
    """The deck's theme body font (charts use it unless told otherwise)."""
    from pptx import Presentation
    master = Presentation(path).slide_masters[0].part
    theme = next(r.target_part for r in master.rels.values() if r.reltype.endswith("/theme"))
    latin = etree.fromstring(theme.blob).find(".//a:minorFont/a:latin", NS)
    return "Calibri" if latin is None else latin.get("typeface")


def check(deck, png_dir):
    """Check `deck` in PowerPoint and export its slides to `png_dir`; returns the report."""
    findings = []
    report = {"deck": str(deck), "findings": findings, "per_slide": {}, "slides": None,
              "status": "checked"}
    deck, png_dir = Path(deck).resolve(), Path(png_dir).resolve()
    start = time.perf_counter()
    work = png_dir / ".com_check"

    def add(found, slide=None):
        findings.extend({"slide": slide, "check": c, "shape": n, "message": m} for c, n, m in found)

    png_dir.mkdir(parents=True, exist_ok=True)
    try:
        original = manifest_mod.manifest(deck)
    except manifest_mod.ManifestError as e:
        add([("manifest", None, f"{e} (not opened in PowerPoint)")])
        report["status"] = "unreadable"
        return report
    font = _theme_font(deck)
    model = deck_model(deck)
    work.mkdir(exist_ok=True)
    # a unique name: PowerPoint refuses to open two presentations with the same file name
    copy = work / f"com_check_{os.getpid()}_{deck.name}"
    shutil.copyfile(deck, copy)
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(copy), MSO_TRUE, MSO_FALSE, MSO_FALSE)
    except Exception as e:  # noqa: BLE001
        add([("open", None, f"PowerPoint could not open the deck: {e}")])
        report["status"] = "unopened"
        app = None
        shutil.rmtree(work, ignore_errors=True)
        return report
    try:
        report["open_s"] = round(time.perf_counter() - start, 2)
        roundtrip = work / "roundtrip.pptx"
        pres.SaveCopyAs(str(roundtrip))
        add([("manifest", None, f"PowerPoint save-copy changed {line}")
             for line in manifest_mod.compare(original, manifest_mod.manifest(roundtrip),
                                              "the deck")])
        report["slides"] = pres.Slides.Count
        width, height = pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight
        for i in range(1, pres.Slides.Count + 1):
            slide = pres.Slides(i)
            png = png_dir / f"slide_{i:02d}.png"
            slide.Export(str(png), "PNG", 1600, round(1600 * height / width))
            found, stats = check_slide(slide, model.get(i, {}), font)
            add(found, i)
            report["per_slide"][i] = {**stats, "findings": len(found), "png": str(png)}
    finally:
        pres.Close()
        pres = app = None
        shutil.rmtree(work, ignore_errors=True)
    report["total_s"] = round(time.perf_counter() - start, 2)
    return report


def main(argv):
    args = [a for a in argv[1:] if not a.startswith("--")]
    json_path = argv[argv.index("--json") + 1] if "--json" in argv else None
    if json_path in args:
        args.remove(json_path)
    if len(args) != 2:
        print(__doc__.strip())
        return 2
    report = check(args[0], args[1])
    if json_path:
        Path(json_path).parent.mkdir(parents=True, exist_ok=True)
        Path(json_path).write_text(json.dumps(report, indent=1), encoding="utf-8")
    for i, st in report["per_slide"].items():
        print(f"slide {i:02d}: shapes={st['shapes']} connectors={st['connectors']} "
              f"glued={st['glued']} "
              f"charts={[c.get('chart_type') for c in st['charts']]} "
              f"{'OK' if not st['findings'] else str(st['findings']) + ' finding(s)'}")
    for f in report["findings"]:
        where = f"slide {f['slide']:02d}" if f["slide"] else "deck"
        print(f"  {where} [{f['check']}] {f['shape'] or ''}: {f['message']}")
    print(f"{len(report['findings'])} finding(s); slides={report['slides']} "
          f"status={report['status']} "
          f"total={report.get('total_s', 0)}s")
    if report["status"] != "checked":
        return 2
    return 1 if report["findings"] else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    sys.exit(main(sys.argv))

