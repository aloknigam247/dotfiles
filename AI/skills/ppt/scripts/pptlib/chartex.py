"""Office 2016 charts (chartEx: box & whisker, sunburst) for python-pptx, which cannot create them.

The chart part is hand-written to Microsoft's DrawingML chartEx schema. Its graphic frame sits in
an mc:AlternateContent whose fallback is a text box for readers older than PowerPoint 2016. The
chart part relates a chart-style and a chart-colour part (PowerPoint refuses the file without
them) and an embedded workbook; every data dimension and series name refers to its workbook range
through a cx:f formula, so Edit Data opens the data and changes reach the chart.
"""
import io
import uuid
from xml.sax.saxutils import escape, quoteattr

from pptx.opc.package import Part
from pptx.oxml import parse_xml
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.shapes.graphfrm import GraphicFrame

CT_CHARTEX = "application/vnd.ms-office.chartex+xml"
CT_COLORS = "application/vnd.ms-office.chartcolorstyle+xml"
CT_STYLE = "application/vnd.ms-office.chartstyle+xml"
EMU_IN = 914400
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_CS = "http://schemas.microsoft.com/office/drawing/2012/chartStyle"
NS_CX = "http://schemas.microsoft.com/office/drawing/2014/chartex"
NS_CX1 = "http://schemas.microsoft.com/office/drawing/2015/9/8/chartex"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_CHARTEX = "http://schemas.microsoft.com/office/2014/relationships/chartEx"
RT_COLORS = "http://schemas.microsoft.com/office/2011/relationships/chartColorStyle"
RT_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
RT_STYLE = "http://schemas.microsoft.com/office/2011/relationships/chartStyle"
SHEET = "Sheet1"
# every entry a chartStyle part must have, in schema order
STYLE_ENTRIES = ("axisTitle", "categoryAxis", "chartArea", "dataLabel", "dataLabelCallout",
                 "dataPoint", "dataPoint3D", "dataPointLine", "dataPointMarker",
                 "dataPointMarkerLayout", "dataPointWireframe", "dataTable", "downBar", "dropLine",
                 "errorBar", "floor", "gridlineMajor", "gridlineMinor", "hiLoLine", "leaderLine",
                 "legend", "plotArea", "plotArea3D", "seriesAxis", "seriesLine", "title",
                 "trendline", "trendlineLabel", "upBar", "valueAxis", "wall")


# --------------------------------------------------------------------------------------------
# workbook
# --------------------------------------------------------------------------------------------
def column_letter(index):
    """Spreadsheet column letters of a 0-based column index (0 -> A, 26 -> AA)."""
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def cell_ref(col, row):
    """Absolute reference to one cell of the data sheet; col and row are 0-based."""
    return f"{SHEET}!${column_letter(col)}${row + 1}"


def range_ref(col, count, first_row=1, last_col=None):
    """Absolute reference to `count` rows of column `col` (through `last_col`) from `first_row`."""
    last = column_letter(col if last_col is None else last_col)
    return (f"{SHEET}!${column_letter(col)}${first_row + 1}:${last}$"
            f"{first_row + count}")


def workbook(columns):
    """A one-sheet workbook blob: [(header, values)] as columns A, B, ... with the header in row 1
    (XlsxWriter is a python-pptx dependency)."""
    import xlsxwriter
    buf = io.BytesIO()
    book = xlsxwriter.Workbook(buf, {"in_memory": True})
    sheet = book.add_worksheet(SHEET)
    for c, (head, values) in enumerate(columns):
        sheet.write(0, c, head)
        for r, value in enumerate(values, 1):
            if value is not None and value != "":
                sheet.write(r, c, value)
        width = max([len(str(head))] + [len(str(v)) for v in values])
        sheet.set_column(c, c, min(max(width + 2, 8), 40))
    book.close()
    return buf.getvalue()


# --------------------------------------------------------------------------------------------
# chart XML
# --------------------------------------------------------------------------------------------
def _num(value):
    return repr(float(value)) if isinstance(value, float) else str(value)


def str_dim(levels, formula, dim_type="cat"):
    """A string dimension: `levels` innermost (leaf) level first, each a list of labels."""
    lvls = "".join(
        f'<cx:lvl ptCount="{len(lvl)}">'
        + "".join(f'<cx:pt idx="{i}">{escape(str(v))}</cx:pt>' for i, v in enumerate(lvl)
                  if v not in (None, ""))
        + "</cx:lvl>" for lvl in levels)
    return f'<cx:strDim type="{dim_type}"><cx:f>{formula}</cx:f>{lvls}</cx:strDim>'


def num_dim(values, formula, dim_type="val", fmt="General"):
    pts = "".join(f'<cx:pt idx="{i}">{_num(v)}</cx:pt>' for i, v in enumerate(values))
    return (f'<cx:numDim type="{dim_type}"><cx:f>{formula}</cx:f><cx:lvl ptCount="{len(values)}" '
            f'formatCode={quoteattr(fmt)}>{pts}</cx:lvl></cx:numDim>')


def solid(hexval, alpha=None):
    a = "" if alpha is None else f'<a:alpha val="{round(alpha * 100000)}"/>'
    return f'<a:solidFill><a:srgbClr val="{hexval}">{a}</a:srgbClr></a:solidFill>'


def sppr(fill=None, alpha=None, line=None, line_w=1.0, no_fill=False):
    """A cx:spPr: solid fill (optionally translucent) or none, and a solid outline or none."""
    body = "<a:noFill/>" if no_fill else (solid(fill, alpha) if fill else "")
    if line:
        body += f'<a:ln w="{round(line_w * 12700)}">{solid(line)}</a:ln>'
    elif line is not None or no_fill:
        body += "<a:ln><a:noFill/></a:ln>"
    return f"<cx:spPr>{body}</cx:spPr>"


def text_props(size, color, font, bold=False, rot=None):
    """A cx:txPr: one default run format for an element's text."""
    rot = "" if rot is None else f' rot="{rot}" vert="horz"'
    return (f'<cx:txPr><a:bodyPr{rot}/><a:lstStyle/><a:p><a:pPr>'
            f'<a:defRPr sz="{round(size * 100)}" b="{int(bold)}">{solid(color)}'
            f'<a:latin typeface={quoteattr(font)}/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="en-US"/></a:p></cx:txPr>')


def title(text, size, color, font, align="min", axis=None):
    """A chart title, or an axis title for axis 'cat' or 'val' (rotated). PowerPoint shows
    chartEx title text only as a run inside txPr, so the text is written both there and in cx:tx
    (as Excel does)."""
    run = (f'<a:r><a:rPr lang="en-US" sz="{round(size * 100)}" b="0">{solid(color)}'
           f'<a:latin typeface={quoteattr(font)}/></a:rPr><a:t>{escape(text)}</a:t></a:r>')
    algn = {"ctr": "ctr", "max": "r", "min": "l"}[align if axis is None else "ctr"]
    rot = ' rot="-5400000" vert="horz"' if axis == "val" else ""
    attrs = "" if axis else f' pos="t" align="{align}" overlay="0"'
    return (f'<cx:title{attrs}><cx:tx><cx:txData><cx:v>{escape(text)}</cx:v></cx:txData>'
            f'</cx:tx><cx:txPr><a:bodyPr{rot} spcFirstLastPara="1" vertOverflow="ellipsis" '
            f'horzOverflow="overflow" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" '
            f'anchor="ctr" anchorCtr="1"/><a:lstStyle/><a:p><a:pPr algn="{algn}" rtl="0">'
            f'<a:defRPr sz="{round(size * 100)}"/></a:pPr>{run}</a:p></cx:txPr></cx:title>')


def series(layout, data_id, name, name_ref, body="", layout_pr=""):
    """A cx:series in schema order (tx, spPr, dataPt*, dataLabels, dataId, layoutPr); `body`
    holds its spPr, dataPt and dataLabels."""
    return (f'<cx:series layoutId="{layout}" uniqueId="{{{str(uuid.uuid4()).upper()}}}">'
            f'<cx:tx><cx:txData><cx:f>{name_ref}</cx:f><cx:v>{escape(name)}</cx:v></cx:txData>'
            f'</cx:tx>{body}<cx:dataId val="{data_id}"/>{layout_pr}</cx:series>')


def chart_space(data_xml, chart_xml, text_xml):
    """The chartEx part; EXTERNAL_DATA is replaced by the workbook reference once it is related."""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<cx:chartSpace xmlns:a="{NS_A}" xmlns:r="{NS_R}" xmlns:cx="{NS_CX}">'
            f'<cx:chartData>EXTERNAL_DATA{data_xml}</cx:chartData><cx:chart>{chart_xml}</cx:chart>'
            f'{sppr(no_fill=True)}{text_xml}</cx:chartSpace>')


def style_xml(sizes, font_color, colors=None):
    """The chartStyle part: every entry PowerPoint requires, referencing the theme. Office 2016
    charts take their text sizes and colours from it ({entry: size} and {entry: colour}, e.g.
    {"dataLabel": "FFFFFF"}), not from the chart's txPr. PowerPoint swaps a data-label colour
    that contrasts too little with its point for a contrasting one."""
    colors = colors or {}

    def entry(name):
        if name == "dataPointMarkerLayout":
            return '<cs:dataPointMarkerLayout symbol="circle" size="5"/>'
        fill = ('<cs:fillRef idx="0"><cs:styleClr val="auto"/></cs:fillRef>'
                if name == "dataPoint" else '<cs:fillRef idx="0"/>')
        sp = ('<cs:spPr><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:ln w="12700">'
              '<a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln></cs:spPr>'
              if name == "dataPoint" else "")
        size = sizes.get(name)
        rpr = f'<cs:defRPr sz="{round(size * 100)}"/>' if size else ""
        return (f'<cs:{name}><cs:lnRef idx="0"/>{fill}<cs:effectRef idx="0"/><cs:fontRef '
                f'idx="minor"><a:srgbClr val="{colors.get(name, font_color)}"/></cs:fontRef>'
                f'{sp}{rpr}</cs:{name}>')

    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cs:chartStyle '
            f'xmlns:cs="{NS_CS}" xmlns:a="{NS_A}" id="410">'
            f'{"".join(entry(n) for n in STYLE_ENTRIES)}</cs:chartStyle>')


def colors_xml():
    """The chartColorStyle part: series take the theme accents (the deck palette) in turn."""
    accents = "".join(f'<a:schemeClr val="accent{i}"/>' for i in range(1, 7))
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cs:colorStyle '
            f'xmlns:cs="{NS_CS}" xmlns:a="{NS_A}" meth="cycle" id="10">{accents}'
            f'<cs:variation/><cs:variation><a:lumMod val="60000"/></cs:variation>'
            f'<cs:variation><a:lumMod val="80000"/><a:lumOff val="20000"/></cs:variation>'
            f'</cs:colorStyle>')


# --------------------------------------------------------------------------------------------
# the graphic frame
# --------------------------------------------------------------------------------------------
def add_chartex(slide, box, xml, name, descr, xlsx_blob, sizes, font_color, colors=None):
    """Insert a chartEx chart (part XML `xml`) into `slide` at box = (x, y, w, h) inches, with its
    style (text sizes and colours, see style_xml()), colour and workbook parts. Returns the
    graphic frame."""
    package = slide.part.package
    part = Part(package.next_partname("/ppt/charts/chartEx%d.xml"), CT_CHARTEX, package, b"")
    for reltype, ctype, template, body in (
            (RT_STYLE, CT_STYLE, "/ppt/charts/style%d.xml",
             style_xml(sizes, font_color, colors)),
            (RT_COLORS, CT_COLORS, "/ppt/charts/colors%d.xml", colors_xml())):
        part.relate_to(Part(package.next_partname(template), ctype, package,
                            body.encode("utf-8")), reltype)
    xrid = part.relate_to(EmbeddedXlsxPart.new(xlsx_blob, package), RT_PACKAGE)
    part._blob = xml.replace("EXTERNAL_DATA", f'<cx:externalData r:id="{xrid}" '
                                              'cx:autoUpdate="0"/>').encode("utf-8")
    rid = slide.part.relate_to(part, RT_CHARTEX)
    sid = slide.shapes._next_shape_id
    x, y, w, h = (round(v * EMU_IN) for v in box)
    xfrm = f'<a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/>'
    cnv = f'<p:cNvPr id="{sid}" name={quoteattr(name)} descr={quoteattr(descr)}/>'
    alt = parse_xml(
        f'<mc:AlternateContent xmlns:mc="{NS_MC}"><mc:Choice xmlns:cx1="{NS_CX1}" '
        f'Requires="cx1"><p:graphicFrame xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:nvGraphicFramePr>'
        f'{cnv}<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm>{xfrm}</p:xfrm>'
        f'<a:graphic><a:graphicData uri="{NS_CX}"><cx:chart xmlns:cx="{NS_CX}" '
        f'xmlns:r="{NS_R}" r:id="{rid}"/></a:graphicData></a:graphic></p:graphicFrame>'
        f'</mc:Choice><mc:Fallback><p:sp xmlns:p="{NS_P}" xmlns:a="{NS_A}"><p:nvSpPr>{cnv}'
        f'<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>{xfrm}</a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr '
        f'anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-US" '
        f'sz="1400"/><a:t>{escape(descr)} (This chart needs PowerPoint 2016 or later.)</a:t>'
        f'</a:r></a:p></p:txBody></p:sp></mc:Fallback></mc:AlternateContent>')
    tree = slide.shapes._spTree
    tree.insert_element_before(alt, "p:extLst")
    return GraphicFrame(alt[0][0], slide.shapes)
