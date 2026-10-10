"""Graph-family diagram builders, mixed into Deck: flowchart, sequence, class, state, ER, use case,
software architecture and agent flow.

Each builder takes data, lays it out automatically with pptlib.layout, draws it in the modern style
(pptlib.style) into a slide region (`box`, default: the content area under the header) and returns
the group shape that holds the diagram. Every edge is a connector glued at both ends. Text keeps
its size: when a diagram does not fit, the builder tries the other direction and tighter spacing,
then raises LayoutError ("split the diagram").
"""
import math
from dataclasses import dataclass, field

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

from . import layout
from .layout import LayoutError
from .shapes import EMU_IN, LINE_SPACING, PRESETS, _PORT_SHAPES, _rgb, bbox, connection_sites
from .style import CONTENT, MIN_PT, TYPE_SCALE, blend, kind_style, tint

_AGENT = {"action": ("hexagon", 4), "decision": ("flowChartDecision", 2),  # (geometry, slot)
          "document": ("flowChartDocument", 3), "input": ("flowChartInputOutput", 1),
          "task": ("roundRect", 0), "tool": ("flowChartPredefinedProcess", 5)}
_ANGLE = {"b": "5400000", "l": "10800000", "r": "0", "t": "16200000"}
_FIXED = {  # preset: the fixed sites the layout may use, as fractions of the box
    "can": {"b": (0.5, 1.0), "l": "spread", "r": "spread", "t": (0.5, 0.0)},
    "ellipse": "centre",
    "flowChartConnector": "centre",
    "flowChartDocument": {"b": (0.5, 20172 / 21600), "l": (0.0, 0.5), "r": (1.0, 0.5),
                          "t": (0.5, 0.0)},
    "flowChartInputOutput": {"b": (0.5, 1.0), "l": (0.1, 0.5), "r": (0.9, 0.5), "t": (0.5, 0.0)},
    "flowChartPredefinedProcess": "centre",
    "flowChartTerminator": "centre",
    "foldedCorner": "centre",
    "parallelogram": "centre",
}
_GROUP_SLOTS = (0, 2, 5, 3, 1, 4)
_MIN = {"can": (1.1, 0.75), "ellipse": (1.4, 0.6), "flowChartDecision": (1.5, 0.75),
        "flowChartDocument": (1.2, 0.6), "flowChartInputOutput": (1.4, 0.5),
        "flowChartPredefinedProcess": (1.3, 0.5), "flowChartTerminator": (1.0, 0.45),
        "foldedCorner": (1.2, 0.5), "hexagon": (1.4, 0.6)}  # minimum (w, h) by geometry
_SPREAD = {"rect", "roundRect", "flowChartAlternateProcess", "flowChartProcess"}
# (spacing scale, text wrap width in inches) tried in turn until a layout fits its region
_STEPS = ((1.25, 2.2), (1.0, 2.2), (1.0, 1.7), (0.8, 1.7), (0.8, 1.3), (0.65, 1.3), (0.55, 1.1))


@dataclass
class GNode:
    """A node to draw: text, shape and colours (hex); `draw` replaces the default box (it gets
    (deck, slide, x, y, w, h) and returns (glue target, every shape it drew))."""
    id: str
    text: str
    prst: str = "roundRect"
    fill: str = "FFFFFF"
    line: str = "64748B"
    color: str = "0F172A"
    dashed: bool = False
    bold: bool = False
    size: tuple = None
    min_w: float = 1.0
    min_h: float = 0.5
    max_w: float = None
    draw: object = None
    name: str = None
    lines: list = None


@dataclass
class GEdge:
    """An edge to draw; `decorate` adds end markers (it gets (deck, slide, points) and returns
    the shapes it made)."""
    src: str
    dst: str
    label: str = ""
    style: str = "solid"
    arrow: str = "triangle"
    head: str = None
    color: str = None
    decorate: object = None
    ends: tuple = ("", "")


@dataclass
class GGroup:
    id: str
    label: str
    members: list
    fill: str = None
    line: str = None
    color: str = None
    dashed: bool = False
    extra: dict = field(default_factory=dict)


def _node3(n):
    """(id, text, kind) from a node tuple; kind defaults to process."""
    n = tuple(n)
    return (n[0], n[1] if len(n) > 1 else n[0], n[2] if len(n) > 2 and n[2] else "process")


def _norm_edge(e):
    """(src, dst, label, style, arrow) from a tuple or Edge-like data."""
    e = tuple(e)
    return (e[0], e[1], e[2] if len(e) > 2 and e[2] else "", e[3] if len(e) > 3 and e[3] else
            "solid", (e[4] if len(e) > 4 else "triangle"))


def _overlap(a, b, pad=0.0):
    """True when boxes (x0, y0, x1, y1) overlap (grown by pad)."""
    return a[0] < b[2] + pad and b[0] < a[2] + pad and a[1] < b[3] + pad and b[1] < a[3] + pad


def _custom_sites(sp, prst, sites):
    """Replace the preset geometry of `sp` (rect, roundRect, hexagon, can or a diamond) by the
    same outline as a custom geometry whose connection sites are `sites` [(side, fx, fy)]
    (fractions of the box). Site k of the shape is sites[k]."""
    sppr = sp._element.spPr
    old = sppr.find(qn("a:prstGeom"))
    adj = {} if old is None else {gd.get("name"): gd.get("fmla").split()[-1]
                                  for gd in old.iterfind(f"{qn('a:avLst')}/{qn('a:gd')}")}
    paths = None
    if prst in ("flowChartDecision", "diamond"):
        guides, extra = PRESETS["diamond"][1], ()
        body = ('<a:moveTo><a:pt x="hc" y="t"/></a:moveTo><a:lnTo><a:pt x="r" y="vc"/></a:lnTo>'
                '<a:lnTo><a:pt x="hc" y="b"/></a:lnTo><a:lnTo><a:pt x="l" y="vc"/></a:lnTo>')
        text, defaults = PRESETS["diamond"][3], {}
    elif prst == "can":  # the preset's three paths: body, lit lid and outline
        defaults, guides, _, text = PRESETS["can"]
        extra, body = (), ""
        arc = '<a:arcTo wR="wd2" hR="y1" stAng="{}" swAng="{}"/>'
        paths = (
            '<a:path stroke="0" extrusionOk="0"><a:moveTo><a:pt x="l" y="y1"/></a:moveTo>'
            + arc.format("cd2", "-10800000") + '<a:lnTo><a:pt x="r" y="y3"/></a:lnTo>'
            + arc.format("0", "cd2") + '<a:close/></a:path>'
            '<a:path stroke="0" fill="lighten" extrusionOk="0"><a:moveTo><a:pt x="l" y="y1"/>'
            '</a:moveTo>' + arc.format("cd2", "cd2") + arc.format("0", "cd2") + '<a:close/>'
            '</a:path><a:path fill="none" extrusionOk="0"><a:moveTo><a:pt x="r" y="y1"/>'
            '</a:moveTo>' + arc.format("0", "cd2") + arc.format("cd2", "cd2")
            + '<a:lnTo><a:pt x="r" y="y3"/></a:lnTo>' + arc.format("0", "cd2")
            + '<a:lnTo><a:pt x="l" y="y1"/></a:lnTo></a:path>')
    else:
        defaults, guides, _, text = PRESETS[prst]
        extra, path = _PORT_SHAPES[prst]
        body = ""
        for cmd in path:
            if cmd[0] == "M":
                body += f'<a:moveTo><a:pt x="{cmd[1]}" y="{cmd[2]}"/></a:moveTo>'
            elif cmd[0] == "L":
                body += f'<a:lnTo><a:pt x="{cmd[1]}" y="{cmd[2]}"/></a:lnTo>'
            else:
                body += f'<a:arcTo wR="x1" hR="x1" stAng="{cmd[1]}" swAng="5400000"/>'
    paths = paths or f'<a:path>{body}<a:close/></a:path>'
    av = "".join(f'<a:gd name="{k}" fmla="val {adj.get(k, round(v))}"/>'
                 for k, v in defaults.items())
    gd = "".join(f'<a:gd name="{g.split()[0]}" fmla="{g.split(" ", 1)[1]}"/>'
                 for g in (*guides, *extra))
    cxn = ""
    for k, (side, fx, fy) in enumerate(sites):
        gd += (f'<a:gd name="s{k}x" fmla="*/ w {round(fx * 100000)} 100000"/>'
               f'<a:gd name="s{k}y" fmla="*/ h {round(fy * 100000)} 100000"/>')
        cxn += f'<a:cxn ang="{_ANGLE[side]}"><a:pos x="s{k}x" y="s{k}y"/></a:cxn>'
    rect = dict(zip("ltrb", text.split()))
    cust = etree.fromstring(
        f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:avLst>{av}</a:avLst><a:gdLst>{gd}</a:gdLst><a:ahLst/><a:cxnLst>{cxn}</a:cxnLst>'
        f'<a:rect l="{rect["l"]}" t="{rect["t"]}" r="{rect["r"]}" b="{rect["b"]}"/>'
        f'<a:pathLst>{paths}</a:pathLst></a:custGeom>')
    old.addprevious(cust)
    sppr.remove(old)
    return sp


class Diagrams:
    """Graph-family diagram builders, mixed into Deck (they use its theme, text and links)."""

    # -- core: lay out, fit, draw -------------------------------------------------------------
    def _hex(self, slot):
        return self.palette_hex[slot % len(self.palette_hex)]

    def _fit_node(self, n, font, wrap):
        if n.size:
            return n.size
        w, h, lines = self.fit_shape(n.text, n.prst if n.prst in PRESETS else "rect", font,
                                     n.bold, max_w=n.max_w or wrap, min_w=n.min_w, min_h=n.min_h)
        n.lines = lines
        return (w, h)

    def _ports_mode(self, n):
        if n.draw is not None or n.prst in _SPREAD:
            return "spread"
        if n.prst in ("flowChartDecision", "diamond"):
            return "diamond"
        return _FIXED.get(n.prst, "centre")

    def _graph(self, s, kind, title, nodes, edges, groups=(), *, box=None, direction="auto",
               prefer="TB", rank=None, pos=None, font=None, alt=None, node_gap=0.35,
               rank_gap=0.55, beside=None):
        """Lay out, fit and draw a graph; returns the diagram group shape."""
        box = box or CONTENT
        font = font or TYPE_SCALE["node"]
        if font < MIN_PT:
            raise ValueError(f"text at {font} pt is below the {MIN_PT} pt minimum")
        if direction not in ("auto", "TB", "LR"):
            raise ValueError(f"direction must be 'auto', 'TB' or 'LR', not {direction!r}")
        ids = {n.id for n in nodes} | {g.id for g in groups}
        for e in edges:
            for end in (e.src, e.dst):
                if end not in ids:
                    raise ValueError(f"{kind}: edge {e.src!r} -> {e.dst!r} names unknown "
                                     f"node {end!r}")
        dirs = [direction] if direction != "auto" else [prefer, "LR" if prefer == "TB" else "TB"]
        titles = {g.id: (self.measure(g.label, TYPE_SCALE["title"], bold=True) + 0.1,
                         TYPE_SCALE["title"] * LINE_SPACING / 72 + 0.04) for g in groups}
        pads = {g.id: (0.16, 0.24 + titles[g.id][1], 0.16, 0.16) for g in groups}
        best = None
        for d in dirs:
            for bands in (1, 2, 3):
                plan = next((p for p in (self._try_layout(
                    nodes, edges, groups, d, scale, wrap, font, box, titles, pads, rank,
                    node_gap, rank_gap, bands=bands, beside=beside) for scale, wrap in _STEPS)
                    if p), None)
                if plan is not None:
                    score = (-bands, plan["scale"], plan["wrap"], plan["fill"], d == prefer)
                    if best is None or score > best[0]:
                        best = (score, plan)
                    break
        if best is None:
            raise LayoutError(f"{kind} does not fit its {box[2]:.1f} x {box[3]:.1f} in region "
                              f"at {font} pt in either direction: split the diagram (at most "
                              "about 15 nodes per slide read well) or give it a larger box")
        plan = best[1]
        if pos:
            plan = self._pin(plan, pos, nodes, edges, groups, font, box, titles, pads,
                             rank, node_gap, rank_gap, beside)
        return self._draw_graph(s, kind, title, nodes, edges, groups, plan, font, alt)

    def _try_layout(self, nodes, edges, groups, d, scale, wrap, font, box, titles, pads,
                    rank, node_gap, rank_gap, moves=None, bands=1, beside=None):
        sizes = {n.id: self._fit_node(n, font, wrap) for n in nodes}
        texts = {i: "\n".join(self.wrap(e.label, wrap * (0.45 if d == "LR" else 0.8),
                                        TYPE_SCALE["label"]))
                 for i, e in enumerate(edges) if e.label}
        chips = {i: self._chip_size(t, TYPE_SCALE["label"]) for i, t in texts.items()}
        room = dict(chips)
        for i, e in enumerate(edges):  # end texts (multiplicities) need room along the edge too
            ends = [self._chip_size(t, TYPE_SCALE["label"]) for t in e.ends if t]
            if ends:
                w, h = room.get(i, (0.0, 0.0))
                room[i] = (w + 2 * max(c[0] for c in ends) + 0.2, h) if d == "LR" else \
                    (w, h + 2 * max(c[1] for c in ends) + 0.2)
        best_crossings = 0
        lay = routes = None
        for variant in ((0, 1, 2) if groups and moves is None else (0,)):
            cand = layout.layered(
                sizes, [(e.src, e.dst) for e in edges], direction=d,
                groups={g.id: g.members for g in groups}, rank=rank,
                ports={n.id: self._ports_mode(n) for n in nodes}, label_sizes=room,
                titles=titles, group_pad=pads, node_gap=node_gap * scale,
                rank_gap=rank_gap * scale, bands=bands, beside=beside, variant=variant)
            if moves is None and (cand.width > box[2] + 1e-6 or cand.height > box[3] + 1e-6):
                continue
            for n, (cx, cy) in (moves or {}).items():
                cand.move(n, cx, cy)
            try:
                cand_routes = layout.route(cand)
            except LayoutError:
                if moves:
                    raise
                continue
            crossings = layout.route_crossings([r.points for r in cand_routes])
            if routes is None or crossings < best_crossings:
                lay, routes, best_crossings = cand, cand_routes, crossings
            if not crossings:
                break
        if routes is None:
            return None
        labels = self._place_labels(lay, routes, chips, titles)
        if labels is None:
            return None
        ends = self._place_ends(lay, routes, edges, labels, chips)
        if ends is None:
            return None
        x0, y0, x1, y1 = self._extent(lay, routes, labels, chips, ends)
        if moves is None and (x1 - x0 > box[2] + 1e-6 or y1 - y0 > box[3] + 1e-6):
            return None
        fill = min(box[2] / max(x1 - x0, 1e-6), box[3] / max(y1 - y0, 1e-6))
        return {"bands": bands, "dir": d, "ends": ends, "extent": (x0, y0, x1, y1),
                "fill": -fill, "labels": labels, "texts": texts,
                "layout": lay, "routes": routes, "scale": scale, "sizes": sizes, "wrap": wrap,
                "box": box}

    def _place_ends(self, lay, routes, edges, labels, chips):
        """Centres of the end texts (multiplicities): beside the first or last run of their
        route, near the node, clear of nodes and labels. {(edge, end): (x, y, w, h)} or None."""
        out = {}
        taken = [(x - 0.02, y - 0.02, x + w + 0.02, y + h + 0.02) for x, y, w, h in
                 lay.nodes.values()]
        taken += [(cx - chips[i][0] / 2, cy - chips[i][1] / 2, cx + chips[i][0] / 2,
                   cy + chips[i][1] / 2) for i, (cx, cy) in labels.items()]
        for k, e in enumerate(edges):
            pts = routes[k].points
            for end, text in enumerate(e.ends):
                if not text:
                    continue
                w, h = self._chip_size(text, TYPE_SCALE["label"])
                tip, prev = (pts[0], pts[1]) if end == 0 else (pts[-1], pts[-2])
                n = math.dist(tip, prev) or 1.0
                ux, uy = (prev[0] - tip[0]) / n, (prev[1] - tip[1]) / n
                spot = None
                for along in (0.06, 0.2, 0.35, 0.5):
                    for side in (1, -1):
                        cx = tip[0] + ux * (along + (w / 2 if ux else 0)) + \
                            side * (w / 2 + 0.04) * abs(uy)
                        cy = tip[1] + uy * (along + (h / 2 if uy else 0)) + \
                            side * (h / 2 + 0.03) * abs(ux)
                        chip = (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
                        if not any(_overlap(chip, t) for t in taken):
                            spot = (cx, cy, w, h)
                            break
                    if spot:
                        break
                if spot is None:
                    return None
                out[(k, end)] = spot
                taken.append((spot[0] - w / 2 - 0.02, spot[1] - h / 2 - 0.02,
                              spot[0] + w / 2 + 0.02, spot[1] + h / 2 + 0.02))
        return out

    def _pin(self, plan, pos, nodes, edges, groups, font, box, titles, pads, rank,
             node_gap, rank_gap, beside):
        """Re-run the chosen layout with nodes pinned at slide positions (centres, inches)."""
        x0, y0, x1, y1 = plan["extent"]
        ox = box[0] + (box[2] - (x1 - x0)) / 2 - x0
        oy = box[1] + (box[3] - (y1 - y0)) / 2 - y0
        moves = {n: (x - ox, y - oy) for n, (x, y) in pos.items()}
        new = self._try_layout(nodes, edges, groups, plan["dir"], plan["scale"], plan["wrap"],
                               font, box, titles, pads, rank, node_gap, rank_gap, moves,
                               plan["bands"], beside)
        if new is None:
            raise LayoutError("pos hints leave no room for the edge labels; move the pinned nodes")
        ex0, ey0, ex1, ey1 = new["extent"]
        if ex0 + ox < box[0] - 1e-6 or ey0 + oy < box[1] - 1e-6 or \
                ex1 + ox > box[0] + box[2] + 1e-6 or ey1 + oy > box[1] + box[3] + 1e-6:
            raise LayoutError("pos hints move the diagram out of its region; choose positions "
                              "inside it")
        new["offset"] = (ox, oy)
        return new

    @staticmethod
    def _extent(lay, routes, labels, chips, ends=None):
        xs = [0.0, lay.width]
        ys = [0.0, lay.height]
        for n, (x, y, w, h) in list(lay.nodes.items()) + list(lay.groups.items()):
            xs += [x, x + w]
            ys += [y, y + h]
        for r in routes:
            xs += [p[0] for p in r.points]
            ys += [p[1] for p in r.points]
        for i, (cx, cy) in labels.items():
            w, h = chips[i]
            xs += [cx - w / 2, cx + w / 2]
            ys += [cy - h / 2, cy + h / 2]
        for cx, cy, w, h in (ends or {}).values():
            xs += [cx - w / 2, cx + w / 2]
            ys += [cy - h / 2, cy + h / 2]
        return min(xs), min(ys), max(xs), max(ys)

    @staticmethod
    def _place_labels(lay, routes, chips, titles):
        """Centres of the edge label chips: on their own route, clear of nodes, group titles and
        each other, preferring spots that hide no other route. None when one cannot be placed."""
        out = {}
        taken = [(x - 0.03, y - 0.03, x + w + 0.03, y + h + 0.03) for x, y, w, h in
                 lay.nodes.values()]
        taken += [(x + 0.14, y + 0.1, x + 0.16 + titles[g][0], y + 0.12 + titles[g][1])
                  for g, (x, y, w, h) in lay.groups.items()]
        segs = [(k, p, q) for k, r in enumerate(routes) for p, q in zip(r.points, r.points[1:])]
        for i in sorted(chips, key=lambda i: -chips[i][0]):
            w, h = chips[i]
            r = routes[i]
            spots = list(r.label_spots)
            for p, q in zip(r.points, r.points[1:]):
                spots += [(p[0] + (q[0] - p[0]) * f, p[1] + (q[1] - p[1]) * f)
                          for f in (0.5, 0.35, 0.65, 0.2, 0.8)]
            best = None
            for rank_, (cx, cy) in enumerate(spots):
                chip = (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2)
                if any(_overlap(chip, t) for t in taken):
                    continue
                hidden = sum(1 for k, p, q in segs if k != i and _overlap(
                    chip, (min(p[0], q[0]), min(p[1], q[1]), max(p[0], q[0]), max(p[1], q[1]))))
                cost = hidden * 10 + rank_ * 0.01
                if best is None or cost < best[0]:
                    best = (cost, (cx, cy), chip)
            if best is None:
                return None
            out[i] = best[1]
            taken.append(tuple(v + d for v, d in zip(best[2], (-0.03, -0.03, 0.03, 0.03))))
        return out

    def _draw_graph(self, s, kind, title, nodes, edges, groups, plan, font, alt):
        lay, routes, box = plan["layout"], plan["routes"], plan["box"]
        x0, y0, x1, y1 = plan["extent"]
        ox, oy = plan.get("offset") or (box[0] + (box[2] - (x1 - x0)) / 2 - x0,
                                         box[1] + (box[3] - (y1 - y0)) / 2 - y0)
        made, shapes = [], {}
        depth = {}
        parent = {m: g.id for g in groups for m in g.members}
        for g in groups:
            k, d = g.id, 0
            while k in parent:
                k, d = parent[k], d + 1
            depth[g.id] = d
        fills = []
        for g in sorted(groups, key=lambda g: depth[g.id]):
            x, y, w, h = lay.groups[g.id]
            sp = self._container(s, ox + x, oy + y, w, h, g, depth[g.id])
            shapes[g.id] = sp
            made.append(sp)
            fills.append(((ox + x, oy + y, ox + x + w, oy + y + h), depth[g.id],
                          sp.fill.fore_color.rgb))
        for n in nodes:
            x, y, w, h = lay.nodes[n.id]
            if n.draw is not None:
                target, extra = n.draw(self, s, ox + x, oy + y, w, h)
                made += extra  # the shapes it drew, the glue target among them
            else:
                target = self.styled_node(s, ox + x, oy + y, w, h, "\n".join(n.lines or [n.text]),
                                          "neutral", size=font, prst=n.prst, fill=_rgb(n.fill),
                                          line=_rgb(n.line), text_color=_rgb(n.color),
                                          bold=n.bold, name=n.name or f"Node: {n.text}")
                if n.dashed:
                    ln = target.line._get_or_add_ln()
                    etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
                made.append(target)
            shapes[n.id] = target
        ends = {}
        for k, r in enumerate(routes):
            for which, end, side in ((0, edges[k].src, r.src_side), (1, edges[k].dst, r.dst_side)):
                p = r.points[0 if which == 0 else -1]
                ends.setdefault(end, []).append((k, which, (ox + p[0], oy + p[1]), side))
        site = {}
        for end, lst in ends.items():
            site.update(self._glue_sites(shapes[end], end, lst, nodes, groups))
        for k, (e, r) in enumerate(zip(edges, routes)):
            pts = [(ox + p[0], oy + p[1]) for p in r.points]
            ln = self.link(s, shapes[e.src], shapes[e.dst], src_site=site[(k, 0)],
                           dst_site=site[(k, 1)], via=pts[1:-1] or None,
                           route="ortho" if len(pts) > 2 else "straight", arrow=e.arrow,
                           head=e.head, style=e.style,
                           color=_rgb(e.color) if e.color else None)
            made.append(ln.connector)
            if e.decorate is not None:
                made += e.decorate(self, s, ln.points)
        for k, (cx, cy) in plan["labels"].items():
            gx, gy = ox + cx, oy + cy
            inside = [f for f in fills if f[0][0] < gx < f[0][2] and f[0][1] < gy < f[0][3]]
            fill = max(inside, key=lambda f: f[1])[2] if inside else None
            made.append(self.label(s, gx, gy, plan["texts"][k], fill=fill,
                                   name=f"Label: {edges[k].label}"))
        for (k, end), (cx, cy, _, _) in plan["ends"].items():
            text = edges[k].ends[end]
            made.append(self.label(s, ox + cx, oy + cy, text, name=f"Label: {text}"))
        text = alt or self._alt(kind, title, nodes, edges, groups)
        return self.group(s, made, name=f"Diagram: {kind} {title or ''}".rstrip(), alt=text)

    @staticmethod
    def _alt(kind, title, nodes, edges, groups):
        names = {n.id: n.text.replace("\n", " ") for n in nodes}
        names.update({g.id: g.label for g in groups})
        links = "; ".join(f"{names[e.src]} to {names[e.dst]}" + (f" ({e.label})" if e.label else "")
                          for e in edges[:30])
        return (f"{kind.replace('_', ' ').capitalize()}{': ' + title if title else ''}. "
                f"{len(nodes)} nodes: {', '.join(names[n.id] for n in nodes[:30])}. "
                f"Connections: {links}.")

    def _container(self, s, x, y, w, h, g, depth):
        """A group frame with its title at the top left."""
        accent = g.line or (self.t["muted"] if depth == 0 and g.extra.get("neutral_root")
                            else self._hex(g.extra.get("slot", depth + 1)))
        fill = g.fill or blend("FFFFFF", accent, 0.06)
        sp = self._prst_shape(s, "roundRect", x, y, w, h)
        sp.adjustments[0] = min(0.5, 0.1 / max(min(w, h), 1e-6))
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
        sp.line.color.rgb = _rgb(accent)
        sp.line.width = Emu(12700)
        ln = sp.line._get_or_add_ln()
        if g.dashed:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        sp.shadow.inherit = False
        title = _rgb(g.color or blend(accent, "0F172A", 0.35))
        self._write(sp, g.label, TYPE_SCALE["title"], title, bold=True, align="left",
                    anchor="top", margins=(0.16, 0.12))
        sp.name = f"Container: {g.label}"
        return sp

    def _glue_sites(self, sp, node, lst, nodes, groups):
        """{(edge, end): site index}: the shape's own sites where the routes end on them, else
        a custom geometry with a site at every route end (rect, roundRect, hexagon, diamond)."""
        out = {}
        sites = connection_sites(sp)
        box = bbox(sp)
        exact = {}
        for k, which, p, side in lst:
            cands = [j for j in range(len(sites)) if sites[j].side == side] or range(len(sites))
            i = min(cands, key=lambda j: math.dist((sites[j].x, sites[j].y), p))
            exact[(k, which)] = (i, math.dist((sites[i].x, sites[i].y), p))
        if all(dist < 0.004 for _, dist in exact.values()):
            return {key: i for key, (i, _) in exact.items()}
        prst = sp._element.spPr.find(qn("a:prstGeom"))
        prst = None if prst is None else prst.get("prst")
        if prst not in ("can", "rect", "roundRect", "hexagon", "flowChartDecision", "diamond"):
            raise LayoutError(f"cannot glue to {sp.name!r}: its routes end away from its sites")
        w, h = box[2] - box[0], box[3] - box[1]
        points = []
        for k, which, p, side in lst:
            key = (side, round((p[0] - box[0]) / w, 5), round((p[1] - box[1]) / h, 5))
            if key not in points:
                points.append(key)
            out[(k, which)] = points.index(key)
        _custom_sites(sp, prst, points)
        return out

    # -- node helpers ---------------------------------------------------------------------------
    def _kind_node(self, nid, text, kind="process", *, prst=None, slot=None, **kw):
        """A GNode in the modern style of node kind `kind` (see style.NODE_KINDS), optionally
        with another geometry or palette slot."""
        kprst, fill, line, color, dashed = kind_style(kind, self.palette_hex, self.t)
        if slot is not None:
            line = self._hex(slot)
            fill = tint(line)
        prst = prst or kprst
        min_w, min_h = _MIN.get(prst, (1.0, 0.5))
        kw.setdefault("min_w", min_w)
        kw.setdefault("min_h", min_h)
        return GNode(nid, text, prst, fill, line, color, dashed, **kw)

    def _dot(self, nid, end=False):
        """A start dot or an end bullseye (state diagrams)."""
        navy = self.t["navy"]
        if not end:
            return GNode(nid, "", "ellipse", navy, navy, navy, size=(0.26, 0.26),
                         name="Node: start")

        def draw(deck, s, x, y, w, h):
            outer = deck._prst_shape(s, "ellipse", x, y, w, h)
            outer.fill.solid()
            outer.fill.fore_color.rgb = deck.WHITE
            outer.line.color.rgb = _rgb(navy)
            outer.line.width = Emu(19050)
            outer.shadow.inherit = False
            outer.name = "Node: end"
            inner = deck._prst_shape(s, "ellipse", x + w * 0.22, y + h * 0.22, w * 0.56, h * 0.56)
            inner.fill.solid()
            inner.fill.fore_color.rgb = _rgb(navy)
            inner.line.fill.background()
            inner.shadow.inherit = False
            inner.name = "Deco: end dot"
            return outer, [outer, inner]

        return GNode(nid, "", "ellipse", size=(0.3, 0.3), draw=draw)

    # -- builders: flowchart, architecture, agent flow, state diagram --------------------------
    def flowchart(self, s, nodes, edges, *, box=None, direction="auto", rank=None, pos=None,
                  font=None, title=None, alt=None):
        """A flowchart, laid out automatically. nodes: [(id, text, kind)] with kind start, end,
        process (default), decision, io, data, document, error, external or note; edges: [(src,
        dst, label, style, arrow)], label/style ('solid', 'dashed', 'dotted')/arrow optional.
        Hints: direction 'auto'|'TB'|'LR', rank {id: n}, pos {id: (x, y)} (slide inches)."""
        gn = [self._kind_node(*_node3(n)) for n in nodes]
        ge = [GEdge(*_norm_edge(e)) for e in edges]
        return self._graph(s, "flowchart", title, gn, ge, box=box, direction=direction,
                           prefer="TB", rank=rank, pos=pos, font=font, alt=alt)

    def architecture(self, s, nodes, edges, groups=(), *, box=None, direction="auto", rank=None,
                     pos=None, font=None, title=None, alt=None):
        """A software architecture with nested containers. nodes: [(id, text, kind)] (process,
        data for stores, external for third parties...); edges: [(src, dst, label, style,
        arrow)]; groups: [(id, label, members)], members being node or group ids. Nodes take
        the colour of the innermost group that holds them."""
        gg, slot_of = self._groups(groups)
        parent = {m: g.id for g in gg for m in g.members}
        gn = []
        for n in nodes:
            nid, text, kind = _node3(n)
            slot = slot_of.get(parent.get(nid)) if kind != "external" else None
            gn.append(self._kind_node(nid, text, kind, slot=slot))
        ge = [GEdge(*_norm_edge(e)) for e in edges]
        return self._graph(s, "architecture", title, gn, ge, gg, box=box, direction=direction,
                           prefer="LR", rank=rank, pos=pos, font=font, alt=alt)

    def _groups(self, groups):
        """GGroups from (id, label, members): a group of groups gets a neutral frame, the others
        a palette colour each. Returns (groups, {group id: palette slot})."""
        ids = {g[0] for g in groups}
        out, slots = [], {}
        k = 0
        for gid, label, members in (tuple(g) for g in groups):
            if any(m in ids for m in members):
                out.append(GGroup(gid, label, list(members), fill=self.t["light"],
                                  line=self.t["subtle"], color=self.t["line"]))
                continue
            slots[gid] = _GROUP_SLOTS[k % len(_GROUP_SLOTS)]
            k += 1
            out.append(GGroup(gid, label, list(members), extra={"slot": slots[gid]}))
        return out, slots

    def agent_flow(self, s, nodes, edges, groups=(), *, box=None, direction="auto", rank=None,
                   pos=None, font=None, title=None, alt=None):
        """An agent flow. nodes: [(id, text, kind)] with kind input, task, tool, document,
        decision or action (or any flowchart kind); edges: [(src, dst, label, style, arrow)],
        e.g. a dotted reference without arrow ("a", "b", "", "dotted", None); groups: [(id,
        label, members)] for the agent's boundary."""
        gn = []
        for n in nodes:
            nid, text, kind = _node3(n)
            if kind in _AGENT:
                prst, slot = _AGENT[kind]
                gn.append(self._kind_node(nid, text, "process", prst=prst, slot=slot))
            else:
                gn.append(self._kind_node(nid, text, kind))
        gg = [GGroup(g[0], g[1], list(g[2]), extra={"slot": 2}) for g in groups]
        ge = [GEdge(*_norm_edge(e)) for e in edges]
        return self._graph(s, "agent_flow", title, gn, ge, gg, box=box, direction=direction,
                           prefer="LR", rank=rank, pos=pos, font=font, alt=alt)

    def state_diagram(self, s, states, transitions, *, composites=(), notes=(), box=None,
                      direction="auto", rank=None, pos=None, font=None, title=None, alt=None):
        """A state diagram. states: [name]; transitions: [(src, dst, label)] with '[*]' as the
        start (src) or a final state (dst); composites: [(name, [member states])] drawn as a
        frame that transitions to `name` reach; notes: [(state, text)] linked by a dotted line."""
        comp = {c[0]: list(c[1]) for c in composites}
        gn = [self._kind_node(st, st, "process") for st in states if st not in comp]
        ge = []
        beside = {}
        ends = 0
        for t in transitions:
            src, dst, label = (tuple(t) + ("",))[:3]
            if src == "[*]" and not any(n.id == "[*]" for n in gn):
                gn.insert(0, self._dot("[*]"))
            if dst == "[*]":
                ends += 1
                dst = f"[*]{ends}"
                gn.append(self._dot(dst, end=True))
                beside[dst] = src
            ge.append(GEdge(src, dst, label or ""))
        for k, (st, text) in enumerate(notes):
            nid = f"note{k}"
            gn.append(self._kind_node(nid, text, "note", min_w=1.2))
            ge.append(GEdge(nid, st, "", "dotted", None))
            beside[nid] = st
        gg = [GGroup(name, name, members, extra={"slot": 5}) for name, members in comp.items()]
        return self._graph(s, "state_diagram", title, gn, ge, gg, box=box, direction=direction,
                           prefer="TB", rank=rank, pos=pos, font=font, alt=alt, beside=beside)
    # -- class and ER diagrams ------------------------------------------------------------------
    def _text_shape(self, s, x, y, w, h, anchor, name):
        """A text box without insets (like Deck.textbox) as a shape: (shape, text frame)."""
        sp = s.shapes.add_textbox(Emu(round(x * EMU_IN)), Emu(round(y * EMU_IN)),
                                  Emu(round(w * EMU_IN)), Emu(round(h * EMU_IN)))
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"bottom": 4, "middle": 3, "top": 1}[anchor]
        for m in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{m}", 0)
        sp.name = name
        return sp, tf

    def _compartments(self, title_lines, rows, columns=1, *, slot=0, header_fill=False,
                      member=None):
        """Size and draw function of a box with a title and compartments of text rows (a class or
        an entity). title_lines: [(text, size, bold, italic)]; rows: [compartment: [[cells]]]."""
        accent = self._hex(slot)
        member = member or TYPE_SCALE["title"]
        line_h = member * LINE_SPACING / 72
        pad_x, pad_y = 0.12, 0.06
        title_h = sum(sz * LINE_SPACING / 72 for _, sz, _, _ in title_lines) + 2 * pad_y + 0.04
        widths = [0.0] * columns
        for comp in rows:
            for row in comp:
                for c, cell in enumerate(row):
                    widths[c] = max(widths[c], self.measure(cell, member,
                                                            bold=columns > 1 and c == 0))
        title_w = max(self.measure(t, sz, bold=b, italic=i) for t, sz, b, i in title_lines)
        w = max(title_w, sum(widths) + 0.2 * (columns - 1)) + 2 * pad_x
        heights = [len(comp) * line_h + 2 * pad_y if comp else 0.12 for comp in rows]
        h = title_h + sum(heights)
        name = title_lines[-1][0]

        def draw(deck, s, x, y, bw, bh):
            made = []
            box = deck._prst_shape(s, "rect", x, y, bw, bh)
            box.fill.solid()
            box.fill.fore_color.rgb = _rgb(tint(accent, 0.06))
            box.line.color.rgb = _rgb(accent)
            box.line.width = Emu(12700)
            box.shadow.inherit = False
            box.name = f"Node: {name}"
            made.append(box)
            if header_fill:
                head = deck._prst_shape(s, "rect", x, y, bw, title_h)
                head.fill.solid()
                head.fill.fore_color.rgb = _rgb(tint(accent, 0.2))
                head.line.color.rgb = _rgb(accent)
                head.line.width = Emu(12700)
                head.shadow.inherit = False
                head.name = f"Deco: header {name}"
                made.append(head)
            sp, tf = deck._text_shape(s, x, y + pad_y, bw, title_h - 2 * pad_y, "middle",
                                      f"Deco: title {name}")
            for k, (text, sz, bold, italic) in enumerate(title_lines):
                deck.run(deck.para(tf, first=k == 0, align="center"), text, size=sz, bold=bold,
                         italic=italic, color=deck.TEXT)
            made.append(sp)
            top = y + title_h
            col_x = [x + pad_x]
            for wcol in widths[:-1]:
                col_x.append(col_x[-1] + wcol + 0.2)
            for comp, ch in zip(rows, heights):
                sep = deck.connector(s, x, top, x + bw, top, color=_rgb(accent), width=1.0)
                sep.name = "Deco: separator"
                made.append(sep)
                for c in range(columns):
                    cw = (x + bw - pad_x - col_x[c]) if c == columns - 1 else widths[c]
                    sp, tf = deck._text_shape(s, col_x[c], top + pad_y, cw + 0.02,
                                              ch - 2 * pad_y, "top", f"Deco: cells {name}")
                    for k, row in enumerate(comp):
                        cell = row[c] if c < len(row) else ""
                        deck.run(deck.para(tf, first=k == 0), cell, size=member,
                                 bold=columns > 1 and c == 0 and bool(cell),
                                 color=deck.LINE if columns > 1 and c == columns - 1 else
                                 deck.TEXT)
                    made.append(sp)
                top += ch
            return box, made

        return (w, h), draw

    def _marker(self, s, tip, prev, kind, color, filled=False):
        """A UML marker ('triangle' or 'diamond', hollow or filled) at a route end, pointing at
        the node."""
        d = (tip[0] - prev[0], tip[1] - prev[1])
        n = math.hypot(*d) or 1.0
        ux, uy = round(d[0] / n), round(d[1] / n)
        length, width = (0.17, 0.15) if kind == "triangle" else (0.26, 0.15)
        cx, cy = tip[0] - ux * length / 2, tip[1] - uy * length / 2
        if kind == "triangle":
            sp = self._prst_shape(s, "triangle", cx - width / 2, cy - length / 2, width, length)
            sp.rotation = {(0, -1): 0, (1, 0): 90, (0, 1): 180, (-1, 0): 270}[(ux, uy)]
        else:
            w, h = (width, length) if uy else (length, width)
            sp = self._prst_shape(s, "diamond", cx - w / 2, cy - h / 2, w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = color if filled else self.WHITE
        sp.line.color.rgb = color
        sp.line.width = Emu(12700)
        sp.shadow.inherit = False
        sp.name = f"Deco: {kind} marker"
        return sp

    def class_diagram(self, s, classes, relations, *, box=None, direction="auto", rank=None,
                      pos=None, font=None, title=None, alt=None):
        """A UML class diagram. classes: [(name, attributes, methods, stereotype)]; relations:
        [(a, b, kind, label, multiplicity at a, multiplicity at b)] with kind association,
        navigation, composition or aggregation (a is the whole), inheritance or realization (b
        is the parent) or dependency. Parents and wholes are laid out first."""
        edges, nodes = [], []
        for k, c in enumerate(classes):
            name, attrs, meths, stereo = (tuple(c) + ([], [], ""))[:4]
            head = [(f"\u00ab{stereo}\u00bb", TYPE_SCALE["label"], False, True)] if stereo else []
            head.append((name, font or TYPE_SCALE["node"], True, stereo == "abstract"))
            size, draw = self._compartments(head, [[[a] for a in attrs], [[m] for m in meths]],
                                            slot=k, member=TYPE_SCALE["label"])
            nodes.append(GNode(name, name, "rect", size=size, draw=draw))
        color = _rgb(self.t["line"])
        for r in relations:
            a, b, kind, label, ma, mb = (tuple(r) + ("association", "", "", ""))[:6]
            if kind in ("inheritance", "realization"):
                edges.append(GEdge(b, a, label, "dashed" if kind == "realization" else "solid",
                                   None, decorate=_deco("triangle", 0, color), ends=(mb, ma)))
            elif kind == "composition":
                edges.append(GEdge(a, b, label, "solid", None, ends=(ma, mb),
                                   decorate=_deco("diamond", 0, color, filled=True)))
            elif kind == "aggregation":
                edges.append(GEdge(a, b, label, "solid", None, ends=(ma, mb),
                                   decorate=_deco("diamond", 0, color)))
            else:
                arrow = {"dependency": "arrow", "navigation": "arrow"}.get(kind)
                edges.append(GEdge(a, b, label, "dashed" if kind == "dependency" else "solid",
                                   arrow, ends=(ma, mb)))
        return self._graph(s, "class_diagram", title, nodes, edges, box=box, direction=direction,
                           prefer="TB", rank=rank, pos=pos, font=font, alt=alt, node_gap=0.5,
                           rank_gap=0.7)

    def er_diagram(self, s, entities, relations, *, box=None, direction="auto", rank=None,
                   pos=None, font=None, title=None, alt=None):
        """An entity-relationship diagram. entities: [(name, [(column, type, key)])] with key
        'PK', 'FK', 'UK', 'PK, FK' or ''; relations: [(a, b, cardinality at a, cardinality at b,
        label)] with cardinalities '1', '0..1', '1..*' or '0..*' drawn as crow's feet."""
        edges, nodes = [], []
        for k, (name, cols) in enumerate(entities):
            rows = [[key or "", col, typ] for col, typ, key in cols]
            size, draw = self._compartments([(name, font or TYPE_SCALE["node"], True, False)],
                                            [rows], 3, slot=k, header_fill=True)
            nodes.append(GNode(name, name, "rect", size=size, draw=draw))
        color = _rgb(self.t["line"])
        for r in relations:
            a, b, ca, cb, label = (tuple(r) + ("1", "0..*", ""))[:5]
            edges.append(GEdge(a, b, label, "solid", None, decorate=_crows(ca, cb, color)))
        return self._graph(s, "er_diagram", title, nodes, edges, box=box, direction=direction,
                           prefer="LR", rank=rank, pos=pos, font=font, alt=alt, node_gap=0.5,
                           rank_gap=1.1)

    def _crow(self, s, tip, prev, card, color):
        """Crow's-foot marks for `card` at a route end (tip on the entity, prev before it)."""
        d = (tip[0] - prev[0], tip[1] - prev[1])
        n = math.hypot(*d) or 1.0
        ux, uy = round(d[0] / n), round(d[1] / n)
        made = []

        def at(a, b):  # a back along the route from the tip, b across it
            return tip[0] - ux * a - uy * b, tip[1] - uy * a + ux * b

        def line(p, q):
            cn = self.connector(s, *p, *q, color=color, width=1.25)
            cn.name = "Deco: crow's foot"
            made.append(cn)

        many, zero = card.endswith("*"), card.startswith("0")
        if many:
            for b in (-0.08, 0.0, 0.08):
                line(at(0.16, 0.0), at(0.0, b))
            if not zero:
                line(at(0.22, -0.07), at(0.22, 0.07))
        else:
            line(at(0.08, -0.07), at(0.08, 0.07))
            if not zero:
                line(at(0.14, -0.07), at(0.14, 0.07))
        if zero:
            cx, cy = at(0.26 if many else 0.18, 0.0)
            o = self._prst_shape(s, "ellipse", cx - 0.05, cy - 0.05, 0.1, 0.1)
            o.fill.solid()
            o.fill.fore_color.rgb = self.LIGHT
            o.line.color.rgb = color
            o.line.width = Emu(15875)
            o.shadow.inherit = False
            o.name = "Deco: crow's foot circle"
            made.append(o)
        return made

    # -- use case and sequence diagrams -----------------------------------------------------------
    def _actor(self, name, font):
        """A stick-figure actor above its name; the glue target is an invisible frame around
        both."""
        tw, th, lines = self.text_size(name, font, max_w=1.4)
        w, h = max(tw + 0.1, 0.8), 0.5 + th
        accent = self._hex(5)

        def draw(deck, s, x, y, bw, bh):
            made = []
            frame = deck._prst_shape(s, "rect", x, y, bw, bh)
            frame.fill.background()
            frame.line.fill.background()
            frame.shadow.inherit = False
            deck._write(frame, "\n".join(lines), font, deck.TEXT, anchor="bottom",
                        margins=(0.02, 0.0))
            frame.name = f"Node: {name}"
            made.append(frame)
            cx, top = x + bw / 2, y + 0.03
            head = deck._prst_shape(s, "ellipse", cx - 0.075, top, 0.15, 0.15)
            head.fill.solid()
            head.fill.fore_color.rgb = _rgb(tint(accent, 0.2))
            head.line.color.rgb = _rgb(accent)
            head.line.width = Emu(15875)
            head.shadow.inherit = False
            head.name = f"Deco: actor head {name}"
            made.append(head)
            for x1, y1, x2, y2 in ((cx, top + 0.15, cx, top + 0.31),
                                   (cx - 0.13, top + 0.21, cx + 0.13, top + 0.21),
                                   (cx, top + 0.31, cx - 0.1, top + 0.44),
                                   (cx, top + 0.31, cx + 0.1, top + 0.44)):
                ln = deck.connector(s, x1, y1, x2, y2, color=_rgb(accent), width=1.25)
                ln.name = f"Deco: actor {name}"
                made.append(ln)
            return frame, made

        return GNode(name, name, "rect", size=(w, h), draw=draw, name=f"Node: {name}")

    def use_case(self, s, actors, cases, links, system="", *, box=None, direction="auto",
                 rank=None, pos=None, font=None, title=None, alt=None):
        """A use case diagram. actors: [(id, name)] drawn as stick figures; cases: [(id, text)]
        drawn as ovals inside the `system` boundary; links: [(a, b)] associations, or (a, b,
        'include' | 'extend') for dashed dependencies."""
        font = font or TYPE_SCALE["node"]
        nodes = [self._actor(name, font) for _, name in actors]
        ids = {aid: name for aid, name in actors}
        for n, (aid, _) in zip(nodes, actors):
            n.id = aid
        for cid, text in cases:
            nodes.append(self._kind_node(cid, text, "process", prst="ellipse", slot=3))
            ids[cid] = text
        edges = []
        for ln in links:
            a, b, kind = (tuple(ln) + ("",))[:3]
            if kind in ("include", "extend"):
                edges.append(GEdge(a, b, f"\u00ab{kind}\u00bb", "dashed", "arrow"))
            else:
                edges.append(GEdge(a, b, "", "solid", None))
        groups = [GGroup("system", system, [c for c, _ in cases], extra={"slot": 1})] \
            if system else []
        return self._graph(s, "use_case", title, nodes, edges, groups, box=box,
                           direction=direction, prefer="LR", rank=rank, pos=pos, font=font,
                           alt=alt, node_gap=0.3, rank_gap=0.8)

    def _flat_label(self, s, cx, bottom, text, color=None, fill=None, align="center"):
        """A one-line text chip on the slide background (no rounding, tight insets) whose bottom
        edge is at `bottom`, centred on cx (or starting at cx when align is 'left')."""
        size = TYPE_SCALE["label"]
        w = self.measure(text, size) + 0.12
        h = size * LINE_SPACING / 72 + 0.03
        x = cx - w / 2 if align == "center" else cx
        sp = self._prst_shape(s, "rect", x, bottom - h, w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill or self.LIGHT
        sp.line.fill.background()
        sp.shadow.inherit = False
        self._write(sp, text, size, color or self.TEXT, margins=(0.06, 0.0), wrap=False)
        sp.name = f"Label: {text}"
        return sp

    def sequence(self, s, participants, messages, *, frames=(), notes=(), box=None, font=None,
                 title=None, alt=None):
        """A sequence diagram. participants: [(id, name, 'participant' | 'actor')]; messages:
        [(src, dst, text)] or (src, dst, text, 'reply') for dashed replies, in time order;
        frames: [(kind, [(guard, first, last)])] with message indexes from 0, e.g. an 'alt' frame
        and its branches; notes: [(participant, text)] shown on the participant's lifeline after
        the first message that reaches it. A request opens an activation bar on its receiver and
        the receiver's next reply closes it; the branches of a frame are alternatives."""
        box = box or CONTENT
        font = font or TYPE_SCALE["node"]
        label = TYPE_SCALE["label"]
        made = []
        x0, y0, w, h = box
        heads = []
        for pid, name, *kind in participants:
            tw, th, lines = self.text_size(name, font, max_w=1.6)
            heads.append((pid, name, kind[0] if kind else "participant", tw, th, lines))
        head_h = max(0.5, max(hd[4] for hd in heads) + 0.24)
        if any(hd[2] == "actor" for hd in heads):
            head_h = max(head_h, 0.5 + max(hd[4] for hd in heads))
        col = (w - 0.2) / len(heads)
        xs = {pid: x0 + 0.1 + col * (i + 0.5) for i, (pid, *_) in enumerate(heads)}
        for src, dst, text, *_ in messages:
            if self.measure(text, label) + 0.4 > abs(xs[src] - xs[dst]):
                raise LayoutError(f"sequence: message {text!r} is wider than the gap between "
                                  f"{src!r} and {dst!r}; shorten it or split the diagram")
        line_h = label * LINE_SPACING / 72 + 0.03
        pitch, guard_h = line_h + 0.06, line_h + 0.04
        branch_of = {g[1]: (k, i) for k, (_, branches) in enumerate(frames)
                     for i, g in enumerate(branches)}
        note_at = {}
        for pid, text in notes:
            first = next((i for i, m in enumerate(messages) if m[1] == pid), None)
            if first is not None:
                note_at.setdefault(first, []).append((pid, text))
        y = y0 + head_h + 0.08
        rows, note_rows, guards = [], [], {}
        for i, (src, dst, text, *kind) in enumerate(messages):
            if i in branch_of:
                guards[i] = y
                y += guard_h
            y += pitch
            rows.append(y - 0.03)
            for pid, text in note_at.get(i, ()):
                tw, th, lines = self.text_size(text, label, max_w=col * 1.6)
                note_rows.append((pid, y + 0.04, tw + 0.3, th + 0.1, lines, text))
                y += th + 0.18
        bottom = y + 0.12
        if bottom > y0 + h + 1e-6:
            raise LayoutError(f"sequence needs {bottom - y0:.1f} in of height but its region has "
                              f"{h:.1f} in: split the diagram")
        bars = self._activations(messages, frames, rows, pitch)
        colors = {pid: self._hex(i) for i, (pid, *_) in enumerate(heads)}
        lifelines = {}
        for pid, name, kind, tw, th, lines in heads:
            cx = xs[pid]
            if kind == "actor":
                node = self._actor(name, font)
                _, extra = node.draw(self, s, cx - node.size[0] / 2, y0, node.size[0], head_h)
                made += extra
            else:
                bw = max(tw + 0.36, 1.2)
                made.append(self.styled_node(s, cx - bw / 2, y0, bw, head_h, "\n".join(lines),
                                             "neutral", size=font, fill=_rgb(tint(colors[pid])),
                                             line=_rgb(colors[pid]), name=f"Node: {name}"))
            dash = self.connector(s, cx, y0 + head_h, cx, bottom, color=_rgb(self.t["subtle"]),
                                  width=1.0)
            etree.SubElement(dash.line._get_or_add_ln(), qn("a:prstDash")).set("val", "dash")
            dash.name = f"Deco: lifeline {name}"
            made.append(dash)
            life = self._prst_shape(s, "rect", cx - 0.06, y0 + head_h, 0.12,
                                    bottom - y0 - head_h)
            life.fill.background()
            life.line.fill.background()
            life.shadow.inherit = False
            life.name = f"Deco: lifeline target {name}"
            lifelines[pid] = life
        for kind, branches in frames:
            top = guards[branches[0][1]]
            last = rows[branches[-1][2]] + 0.1
            fx = x0 + 0.02
            frame = self._prst_shape(s, "rect", fx, top, w - 0.04, last - top)
            frame.fill.background()
            frame.line.color.rgb = _rgb(self.t["muted"])
            frame.line.width = Emu(12700)
            frame.shadow.inherit = False
            frame.name = f"Deco: {kind} frame"
            made.append(frame)
            tag_w = self.measure(kind, label, bold=True) + 0.3
            tag = self._prst_shape(s, "snip1Rect", fx, top, tag_w, guard_h)
            tag.adjustments[0] = 0.3
            tag.fill.solid()
            tag.fill.fore_color.rgb = _rgb(self.t["panel"])
            tag.line.color.rgb = _rgb(self.t["muted"])
            tag.line.width = Emu(12700)
            tag.shadow.inherit = False
            self._write(tag, kind, label, self.TEXT, bold=True, margins=(0.06, 0.0), wrap=False)
            tag.name = f"Label: {kind} frame"
            made.append(tag)
            for i, (guard, first, _) in enumerate(branches):
                gy = guards[first]
                if i:
                    sep = self.connector(s, fx, gy, fx + w - 0.04, gy,
                                         color=_rgb(self.t["muted"]), width=1.0)
                    etree.SubElement(sep.line._get_or_add_ln(), qn("a:prstDash")).set("val",
                                                                                        "dash")
                    sep.name = "Deco: frame separator"
                    made.append(sep)
                made.append(self._flat_label(s, fx + tag_w + 0.1, gy + guard_h - 0.01,
                                             f"[{guard}]", color=self.LINE, fill=self.LIGHT,
                                             align="left"))
        for pid, a, b in bars:
            bar = self._prst_shape(s, "rect", xs[pid] - 0.06, a, 0.12, b - a)
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(tint(colors[pid], 0.25))
            bar.line.color.rgb = _rgb(colors[pid])
            bar.line.width = Emu(9525)
            bar.shadow.inherit = False
            bar.name = f"Deco: activation {pid}"
            made.append(bar)
        sites = {pid: sorted({yy for m, yy in zip(messages, rows) if pid in m[:2]})
                 for pid in xs}
        for pid, life in lifelines.items():
            hh = bottom - y0 - head_h
            spots = []
            for yy in sites[pid]:
                on_bar = any(p == pid and a - 1e-6 <= yy <= b + 1e-6 for p, a, b in bars)
                for side in ("l", "r"):
                    fx = 0.5 if not on_bar else (0.0 if side == "l" else 1.0)
                    spots.append((side, fx, (yy - y0 - head_h) / hh))
            _custom_sites(life, "rect", spots)
            made.append(life)
        for (src, dst, text, *kind), yy in zip(messages, rows):
            reply = bool(kind) and kind[0] == "reply"
            right = xs[dst] > xs[src]
            ia = sites[src].index(yy) * 2 + (1 if right else 0)
            ib = sites[dst].index(yy) * 2 + (0 if right else 1)
            ln = self.link(s, lifelines[src], lifelines[dst], src_site=ia, dst_site=ib,
                           route="straight", arrow="arrow" if reply else "triangle",
                           style="dashed" if reply else "solid",
                           name=f"Edge: {src} -> {dst}: {text}")
            made.append(ln.connector)
            made.append(self._flat_label(s, (xs[src] + xs[dst]) / 2, yy - 0.025, text))
        for pid, ny, nw, nh, lines, text in note_rows:
            made.append(self.styled_node(s, xs[pid] - nw / 2, ny, nw, nh, "\n".join(lines),
                                         "note", size=label, name=f"Label: note {text}"))
        names = ", ".join(hd[1] for hd in heads)
        steps = "; ".join(f"{m[0]} to {m[1]}: {m[2]}" for m in messages[:30])
        return self.group(s, made, name=f"Diagram: sequence {title or ''}".rstrip(),
                          alt=alt or f"Sequence diagram{': ' + title if title else ''}. "
                                     f"Participants: {names}. Messages: {steps}.")

    @staticmethod
    def _activations(messages, frames, rows, pitch):
        """[(participant, top, bottom)] activation bars: a request opens one on its receiver,
        the receiver's next reply closes it, a bar without a reply is short; every branch of a
        frame starts from the state before the frame, and bars with one start merge."""
        bars = {}
        open_ = {}
        starts = {g[1]: (k, i) for k, (_, branches) in enumerate(frames)
                  for i, g in enumerate(branches)}
        saved = {}
        for i, (src, dst, _, *kind) in enumerate(messages):
            if i in starts:
                k, branch = starts[i]
                if branch == 0:
                    saved[k] = {p: list(v) for p, v in open_.items()}
                else:
                    open_ = {p: list(v) for p, v in saved[k].items()}
            if kind and kind[0] == "reply":
                if open_.get(src):
                    a = open_[src].pop()
                    bars[(src, a)] = max(bars.get((src, a), a), rows[i])
            else:
                open_.setdefault(dst, []).append(rows[i])
                bars.setdefault((dst, rows[i]), rows[i])
        return [(p, a - 0.05, max(b, a + pitch * 0.6) + 0.05) for (p, a), b in bars.items()]

def _deco(kind, end, color, filled=False):
    """A GEdge.decorate callback: a `kind` marker at the start (end=0) or the end."""
    def decorate(deck, s, pts):
        tip, prev = (pts[0], pts[1]) if end == 0 else (pts[-1], pts[-2])
        return [deck._marker(s, tip, prev, kind, color, filled)]
    return decorate


def _crows(ca, cb, color):
    """A GEdge.decorate callback: crow's feet for both ends of an ER relation."""
    def decorate(deck, s, pts):
        return deck._crow(s, pts[0], pts[1], ca, color) + deck._crow(s, pts[-1], pts[-2], cb, color)
    return decorate


# @@BUILDERS3@@
