"""Native, editable charts mixed into Deck, in the modern style (see reference/diagrams.md).

chart() draws bar, column, line, pie, area, stacked area, doughnut and radar charts; bar_line(),
scatter(), bubble(), histogram() and candlestick() draw combo, XY, bubble, pre-binned and stock
charts; box_plot() and sunburst() draw Office 2016 charts (chartex.py); heatmap() draws a table
with colour-scaled cells and a gradient legend. Charts are python-pptx charts adjusted in their
XML, so every one keeps its data in an embedded workbook and Edit Data works in PowerPoint.

The style: theme palette colours, light gridlines, no chart border or background, direct labels
where they read better than a legend, deck font, text at 11 pt or more, alt text on every chart.
"""
import datetime as dt
import math
import re
from xml.sax.saxutils import quoteattr

from lxml import etree
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.chart.datalabel import DataLabels
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.chart import XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import chartex
from .shapes import alt_text, check_size
from .style import CONTENT, blend

CHART_KINDS = {
    "area": XL_CHART_TYPE.AREA, "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE,
    "radar": XL_CHART_TYPE.RADAR_FILLED,
}
CHART_PT = {"label": 12, "legend": 12, "tick": 12, "title": 16}  # the chart type scale (pt)
EMU_IN = 914400
GRID_PT = 0.75
HOLE = 50  # doughnut hole, % of the diameter
LEGEND_POS = {
    "bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP,
    "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT,
}
LINE_PT = 2.25
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
PIE_START = 0  # degrees clockwise from 12 o'clock where the first slice starts
SHORT_TITLE = 0.6  # inches: a vertical axis title this short stays upright
TITLE_BAND = 0.5  # inches the chart title takes at the top of a chart
TITLE_INSET = 0.06  # inches from the chart's top-left corner to its title
_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}$")
# series children that follow c:dLbls (schema order), for inserting a series' data labels
_AFTER_DLBLS = ("c:trendline", "c:errBars", "c:cat", "c:val", "c:xVal", "c:yVal", "c:smooth",
                "c:bubbleSize", "c:bubble3D", "c:shape", "c:extLst")


# --------------------------------------------------------------------------------------------
# colours and XML snippets
# --------------------------------------------------------------------------------------------
def _hex(color):
    return str(color) if isinstance(color, RGBColor) else color.lstrip("#").upper()


def _luminance(hexval):
    def channel(v):
        v /= 255
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (channel(int(hexval[i:i + 2], 16)) for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def on_color(fill, dark):
    """White or `dark`, whichever contrasts more with `fill` (WCAG contrast ratio)."""
    lum = _luminance(fill)
    return "FFFFFF" if 1.05 / (lum + 0.05) >= (lum + 0.05) / (_luminance(dark) + 0.05) else dark


def _solid_xml(hexval, alpha=None):
    a = "" if alpha is None else f'<a:alpha val="{round(alpha * 100000)}"/>'
    return f'<a:solidFill><a:srgbClr val="{hexval}">{a}</a:srgbClr></a:solidFill>'


def _sppr(fill=None, alpha=None, line=None, width=GRID_PT, dash=None):
    """A c:spPr: a solid (optionally translucent) fill or none, and a solid line or none."""
    body = _solid_xml(fill, alpha) if fill else "<a:noFill/>"
    if line:
        dash = f'<a:prstDash val="{dash}"/>' if dash else ""
        body += f'<a:ln w="{round(width * 12700)}" cap="rnd">{_solid_xml(line)}{dash}</a:ln>'
    else:
        body += "<a:ln><a:noFill/></a:ln>"
    return parse_xml(f'<c:spPr xmlns:c="{NS_C}" xmlns:a="{NS_A}">{body}</c:spPr>')


def _txpr(size, color, font, bold=False):
    check_size(size)
    return parse_xml(
        f'<c:txPr xmlns:c="{NS_C}" xmlns:a="{NS_A}"><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
        f'<a:defRPr sz="{round(size * 100)}" b="{int(bold)}">{_solid_xml(color)}'
        f'<a:latin typeface="{font}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p>'
        f'</c:txPr>')


def _replace(parent, new, *successors):
    """Put element `new` into `parent` in place of any child with its tag, before the first of
    `successors` (schema order) or at the end."""
    for old in parent.findall(new.tag):
        parent.remove(old)
    for tag in successors:
        nxt = parent.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(new)
            return new
    parent.append(new)
    return new


def _set_val(parent, tag, val, *successors):
    el = parent.find(qn(tag))
    if el is None:
        el = _replace(parent, etree.Element(qn(tag)), *successors)
    el.set("val", str(val))
    return el


def _series_sppr(ser, fill=None, alpha=None, line=None, width=LINE_PT, dash=None):
    ser._element.get_or_add_spPr()
    old = ser._element.spPr
    old.addprevious(_sppr(fill, alpha, line, width, dash))
    ser._element.remove(old)


def _manual_layout(parent, x, y, w=None, h=None, target=None):
    """Give a chart element (plot area, title, legend, data label) a manual layout in fractions
    of the chart: an edge position for plot areas, titles and legends, else an offset."""
    layout = parent.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        first = parent.find(qn("c:idx"))
        if first is not None:
            first.addnext(layout)
        elif parent.tag == qn("c:title"):
            tx = parent.find(qn("c:tx"))
            (tx.addnext(layout) if tx is not None else parent.insert(0, layout))
        elif parent.tag == qn("c:legend"):
            _replace(parent, layout, "c:overlay", "c:spPr", "c:txPr", "c:extLst")
        else:
            parent.insert(0, layout)
    for child in list(layout):
        layout.remove(child)
    mode = "" if w is None and parent.tag == qn("c:dLbl") else (
        '<c:xMode val="edge"/><c:yMode val="edge"/>')
    size = "" if w is None else f'<c:w val="{w:.5f}"/><c:h val="{h:.5f}"/>'
    target = "" if target is None else f'<c:layoutTarget val="{target}"/>'
    layout.append(parse_xml(f'<c:manualLayout xmlns:c="{NS_C}">{target}{mode}<c:x val="{x:.5f}"/>'
                            f'<c:y val="{y:.5f}"/>{size}</c:manualLayout>'))


def _nice(lo, hi, ticks=6):
    """(minimum, maximum, step) of a nice axis covering lo..hi with about `ticks` intervals."""
    if hi <= lo:
        hi = lo + (abs(lo) or 1)
    raw = (hi - lo) / ticks
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return math.floor(lo / step + 1e-9) * step, math.ceil(hi / step - 1e-9) * step, step


def _bin_width(lo, hi, bins):
    """The nice bin width (1, 2, 2.5 or 5 times a power of ten) nearest (hi - lo) / bins."""
    raw = max(hi - lo, 1e-9) / max(bins, 1)
    mag = 10 ** math.floor(math.log10(raw))
    return min((m * mag for m in (1, 2, 2.5, 5, 10)), key=lambda c: abs(math.log(c / raw)))


def _fmt(value):
    return f"{value:,.0f}" if float(value).is_integer() else f"{value:,.4g}"


def _summary(categories, series, limit=8):
    """'name: cat value, ...' for each series, shortened for alt text."""
    parts = []
    for name, values in series:
        pairs = [f"{c} {_fmt(v)}" for c, v in zip(categories, values)]
        if len(pairs) > limit:
            pairs = pairs[:limit] + [f"and {len(values) - limit} more"]
        parts.append(f"{name}: {', '.join(pairs)}")
    return "; ".join(parts)


def _pairs(data):
    """[(name, values)] from a {name: values} dict or a list of (name, values) pairs."""
    items = list(data.items()) if isinstance(data, dict) else list(data)
    for item in items:
        if not (isinstance(item, (tuple, list)) and len(item) == 2 and isinstance(item[0], str)):
            raise ValueError(f"expected (name, values) pairs, got {item!r}")
    return items


def _date_label(value, date_format):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime(date_format)
    if isinstance(value, str) and _ISO_DATE.match(value):
        return dt.date.fromisoformat(value).strftime(date_format)
    return str(value)


# --------------------------------------------------------------------------------------------
# the mixin
# --------------------------------------------------------------------------------------------
class Charts:
    """Chart builders, mixed into Deck (they use its theme palette, font and text helpers)."""

    # -- shared style -------------------------------------------------------------------------
    @property
    def _ink(self):
        """Colour of chart text other than titles: the text colour softened toward the lines."""
        return blend(self._hx("text"), self._hx("line"), 0.45)

    def _hx(self, key):
        """A theme colour as 'RRGGBB'."""
        return _hex(self.t[key])

    def _color(self, i):
        return self.palette_hex[i % len(self.palette_hex)]

    def _frame(self, s, xl_type, box, data, name, alt):
        x, y, w, h = box
        gf = s.shapes.add_chart(xl_type, Emu(round(x * EMU_IN)), Emu(round(y * EMU_IN)),
                                Emu(round(w * EMU_IN)), Emu(round(h * EMU_IN)), data)
        gf.name = name
        alt_text(gf, alt)
        return gf

    def _style(self, ch, title, box):
        """The modern look shared by every classic chart: no border or background, deck font."""
        cs = ch._chartSpace
        _set_val(cs, "c:roundedCorners", 0, "c:style", "c:clrMapOvr", "c:pivotSource",
                 "c:protection", "c:chart")
        _replace(cs, _sppr(), "c:txPr", "c:externalData", "c:printSettings", "c:userShapes",
                 "c:extLst")
        _replace(cs, _txpr(CHART_PT["tick"], self._ink, self.FONT), "c:externalData",
                 "c:printSettings", "c:userShapes", "c:extLst")
        plot_area = cs.find(qn("c:chart")).find(qn("c:plotArea"))
        _replace(plot_area, _sppr(), "c:extLst")
        self._title(ch, title, box)

    def _title(self, ch, title, box):
        """A left-aligned chart title at the top of the chart (the plot area stays below it)."""
        if not title:
            ch.has_title = False
            return
        ch.has_title = True
        tf = ch.chart_title.text_frame
        tf.text = title
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        run = para.runs[0]
        run.font.size = Pt(CHART_PT["title"])
        run.font.bold = False
        run.font.name = self.FONT
        run.font.color.rgb = RGBColor.from_string(self._hx("text"))
        # a hair inside the chart's corner: at the very edge COM places the title outside it
        _manual_layout(ch._chartSpace.find(qn("c:chart")).find(qn("c:title")),
                       TITLE_INSET / box[2], TITLE_INSET / box[3])

    def _axis(self, axis, grid=False, line=False, fmt=None, title=None, minimum=None,
              maximum=None, step=None, vertical=False):
        """Tick labels in the chart font, light gridlines or none, a light axis line or none.
        A vertical axis's title reads upward, unless it is short enough to stand upright."""
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.has_major_gridlines = grid
        if grid:
            axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(
                self._hx("card_border"))
            axis.major_gridlines.format.line.width = Pt(GRID_PT)
        if line:
            axis.format.line.color.rgb = RGBColor.from_string(self._hx("subtle"))
            axis.format.line.width = Pt(GRID_PT)
        else:
            axis.format.line.fill.background()
        font = axis.tick_labels.font
        font.size, font.name = Pt(CHART_PT["tick"]), self.FONT
        font.color.rgb = RGBColor.from_string(self._ink)
        if fmt:
            axis.tick_labels.number_format = fmt
            axis.tick_labels.number_format_is_linked = False
        if minimum is not None:
            axis.minimum_scale = minimum
        if maximum is not None:
            axis.maximum_scale = maximum
        if step is not None:
            axis.major_unit = step
        if title:
            axis.has_title = True
            tf = axis.axis_title.text_frame
            tf.text = title
            run = tf.paragraphs[0].runs[0]
            run.font.size, run.font.bold, run.font.name = Pt(CHART_PT["tick"]), False, self.FONT
            run.font.color.rgb = RGBColor.from_string(self._ink)
            if vertical and self.measure(title, CHART_PT["tick"]) <= SHORT_TITLE:
                body = tf._txBody.find(qn("a:bodyPr"))
                body.set("rot", "0")
                body.set("vert", "horz")

    def _legend(self, ch, show, pos):
        ch.has_legend = bool(show)
        if not show:
            return
        if pos not in LEGEND_POS:
            raise ValueError(f"unknown legend position {pos!r}; expected one of "
                             f"{sorted(LEGEND_POS)}")
        ch.legend.position = LEGEND_POS[pos]
        ch.legend.include_in_layout = False
        font = ch.legend.font
        font.size, font.name = Pt(CHART_PT["legend"]), self.FONT
        font.color.rgb = RGBColor.from_string(self._ink)

    def _labels(self, owner, color, fmt=None, position=None, value=True, category=False,
                percent=False, series=False, bold=False):
        """Data labels of a plot or series in the chart font, without a legend key."""
        dls = owner.data_labels if hasattr(owner, "data_labels") else \
            DataLabels(owner._element.get_or_add_dLbls())
        font = dls.font
        font.size, font.name, font.bold = Pt(CHART_PT["label"]), self.FONT, bold
        font.color.rgb = RGBColor.from_string(color)
        if fmt:
            dls.number_format = fmt
            dls.number_format_is_linked = False
        if position is not None:
            dls.position = position
        dls.show_legend_key = False
        dls.show_value, dls.show_category_name = value, category
        dls.show_percentage, dls.show_series_name = percent, series
        return dls

    def _chart_alt(self, alt, kind, title, detail):
        return alt or f"{kind}{f' {title}' if title else ''}: {detail}."

    def _note(self, s, box, lines, name, align="left", anchor="middle"):
        """A text box of (text, size, colour, bold) lines at box = (x, y, w, h); returns it."""
        x, y, w, h = box
        tb = s.shapes.add_textbox(Emu(round(x * EMU_IN)), Emu(round(y * EMU_IN)),
                                  Emu(round(w * EMU_IN)), Emu(round(h * EMU_IN)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"bottom": MSO_ANCHOR.BOTTOM, "middle": MSO_ANCHOR.MIDDLE,
                              "top": MSO_ANCHOR.TOP}[anchor]
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (text, size, color, bold) in enumerate(lines):
            self.run(self.para(tf, first=i == 0, align=align), text, size=size,
                     color=RGBColor.from_string(color), bold=bold)
        tb.name = name
        return tb

    # -- chart() --------------------------------------------------------------------------------
    def chart(self, s, kind, x, y, w, h, categories, series, title=None, has_legend=True,
              legend_pos="bottom", alt=None, *, labels=None, number_format=None, filled=True,
              center_label=None, x_title=None, y_title=None, name=None):
        """A native, editable chart in the modern style; returns the python-pptx chart.

        kind: 'bar', 'column', 'line', 'pie', 'area', 'area_stacked', 'doughnut' or 'radar'.
        series: [(name, values)], one value per category; pie and doughnut take one series.
        labels: data labels; None picks them where they read better (pie and doughnut slices,
        single-series bars and columns). Pie and doughnut labels show the percentage, and the
        category too when there is no legend. number_format: value format, e.g. '0.0' or
        '"$"#,##0'. filled: radar areas filled and translucent (False: lines with markers).
        center_label: text in a doughnut's hole; its first line is large ('100%\\nof visits').
        x_title / y_title: axis titles. alt: alternative text (default: a data summary).
        """
        if kind not in CHART_KINDS:
            raise ValueError(f"unknown chart kind {kind!r}; expected one of {sorted(CHART_KINDS)}")
        if kind in ("pie", "doughnut") and len(series) != 1:
            raise ValueError(f"{kind} charts require exactly one series")
        if center_label and kind != "doughnut":
            raise ValueError("center_label is for doughnut charts")
        for sname, values in series:
            if len(values) != len(categories):
                raise ValueError(f"series {sname!r} has {len(values)} values, "
                                 f"expected {len(categories)}")
        data = CategoryChartData(number_format=number_format or "General")
        data.categories = categories
        for sname, values in series:
            data.add_series(sname, values)
        xl_type = XL_CHART_TYPE.RADAR_MARKERS if kind == "radar" and not filled else \
            CHART_KINDS[kind]
        words = {"area_stacked": "Stacked area", "column": "Column"}
        label = f"{words.get(kind, kind.capitalize())} chart"
        gf = self._frame(s, xl_type, (x, y, w, h), data, name or f"Chart: {title or label}",
                         self._chart_alt(alt, label, title, _summary(categories, series)))
        ch = gf.chart
        self._style(ch, title, (x, y, w, h))
        if kind in ("pie", "doughnut"):
            self._round(s, ch, kind, (x, y, w, h), categories, series[0][1], title, has_legend,
                        legend_pos, labels, number_format, center_label)
            return ch
        self._legend(ch, has_legend, legend_pos)
        plot = ch.plots[0]
        if kind == "radar":
            self._radar(ch, filled)
        else:
            vertical = kind != "bar"
            value_axis, cat_axis = ch.value_axis, ch.category_axis
            self._axis(value_axis, grid=True, fmt=number_format, title=y_title if vertical
                       else x_title, vertical=vertical)
            self._axis(cat_axis, line=True, title=x_title if vertical else y_title,
                       vertical=not vertical)
        for i, ser in enumerate(plot.series):
            color = self._color(i)
            if kind in ("bar", "column"):
                _series_sppr(ser, color)
            elif kind == "line":
                _series_sppr(ser, line=color, width=LINE_PT)
                self._series_marker(ser, color)
            elif kind == "area":
                _series_sppr(ser, color, 0.35, color, LINE_PT)
            elif kind == "area_stacked":
                _series_sppr(ser, color, 0.8, "FFFFFF", 1.0)
        if kind in ("bar", "column"):
            plot.gap_width = 60
            plot.overlap = -10 if len(series) > 1 else 0
            if labels or (labels is None and len(series) == 1):
                for i, ser in enumerate(plot.series):
                    self._labels(ser, on_color(self._color(i), self._hx("text")), number_format,
                                 XL_LABEL_POSITION.INSIDE_END)
        elif labels:
            self._labels(plot, self._ink, number_format, None)
        if kind.startswith("area"):
            _set_val(ch._chartSpace.find(f".//{qn('c:valAx')}"), "c:crossBetween", "midCat",
                     "c:majorUnit", "c:minorUnit", "c:dispUnits", "c:extLst")
        return ch

    def _series_marker(self, ser, color, size=7):
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = size
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
        ser.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        ser.marker.format.line.width = Pt(1.0)

    def _radar(self, ch, filled):
        """Radar: translucent filled areas with a solid outline, light rings and spokes."""
        for i, ser in enumerate(ch.plots[0].series):
            color = self._color(i)
            if filled:
                _series_sppr(ser, color, 0.25, color, LINE_PT)
            else:
                _series_sppr(ser, line=color, width=LINE_PT)
                self._series_marker(ser, color)
        values = [v for ser in ch.plots[0].series for v in ser.values if v is not None]
        lo, hi, step = _nice(min(0, min(values)), max(values), 5)
        self._axis(ch.value_axis, grid=True, minimum=lo, maximum=hi, step=step)
        cat_axis = ch.category_axis
        self._axis(cat_axis, grid=True)
        font = cat_axis.tick_labels.font
        font.size = Pt(CHART_PT["tick"] + 1)
        font.color.rgb = RGBColor.from_string(self._hx("text"))

    def _round(self, s, ch, kind, box, categories, values, title, has_legend, legend_pos, labels,
               number_format, center_label):
        """Pie and doughnut: palette slices with white separators; the largest circle that fits
        with its outside labels, in a square plot area whose position is known (so the labels
        and the doughnut's centre label can be placed); labels inside the slices that hold
        them, outside the rest."""
        x, y, w, h = box
        hole = HOLE / 100 if kind == "doughnut" else 0.0
        top = TITLE_BAND if title else 0.1
        total = float(sum(values)) or 1.0
        show_category = labels is not False and not has_legend
        texts = [("\n".join(([str(c)] if show_category else []) + [f"{v / total:.0%}"]))
                 for c, v in zip(categories, values)]
        sizes = [self.text_size(t, CHART_PT["label"], bold=True)[:2] for t in texts]
        legend_w = (max(self.measure(str(c), CHART_PT["legend"]) for c in categories) + 0.55
                    if has_legend and legend_pos in ("left", "right") else 0.0)
        legend_h = 0.45 if has_legend and legend_pos in ("top", "bottom") else 0.0
        area = (x + 0.1 + (legend_w if legend_pos == "left" else 0.0),
                y + top + (legend_h if legend_pos == "top" else 0.0),
                w - legend_w - 0.2, h - top - legend_h - 0.1)
        slices = []
        start = 0.0
        for v in values:
            slices.append((math.radians((PIE_START + 360 * (start + v / total / 2)) % 360),
                           v / total))
            start += v / total

        def inside(i, radius):
            share, (lw, lh) = slices[i][1], sizes[i]
            r_label = radius * ((1 + hole) / 2 if hole else 0.6)
            depth = radius * (1 - hole) if hole else radius * 0.8
            return lw <= 0.95 * 2 * r_label * math.sin(min(math.pi * share, math.pi / 2)) and \
                lh <= 0.95 * depth

        def extent(radius):
            """Bounding box of the circle and its outside labels, around the centre."""
            box_ = [-radius, -radius, radius, radius]
            for i, (mid, _) in enumerate(slices):
                if labels is False or inside(i, radius):
                    continue
                lw, lh = sizes[i]
                cx_, cy_ = (math.sin(mid) * self._reach(radius, mid, lw, lh),
                            -math.cos(mid) * self._reach(radius, mid, lw, lh))
                box_ = [min(box_[0], cx_ - lw / 2), min(box_[1], cy_ - lh / 2),
                        max(box_[2], cx_ + lw / 2), max(box_[3], cy_ + lh / 2)]
            return box_

        lo, hi = 0.25, min(area[2], area[3]) / 2
        for _ in range(40):
            mid_r = (lo + hi) / 2
            left, upper, right, lower = extent(mid_r)
            lo, hi = (mid_r, hi) if right - left <= area[2] and lower - upper <= area[3] else \
                (lo, mid_r)
        radius = lo
        left, upper, right, lower = extent(radius)
        cx = area[0] + (area[2] - (right - left)) / 2 - left
        cy = area[1] + (area[3] - (lower - upper)) / 2 - upper
        plot_area = ch._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
        _manual_layout(plot_area, (cx - radius - x) / w, (cy - radius - y) / h, 2 * radius / w,
                       2 * radius / h, "inner")
        plot = ch.plots[0]
        _set_val(plot._element, "c:firstSliceAng", PIE_START, "c:holeSize", "c:extLst")
        if hole:
            _set_val(plot._element, "c:holeSize", HOLE, "c:extLst")
        ser = plot.series[0]
        for i, point in enumerate(ser.points):
            fmt = point.format
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = RGBColor.from_string(self._color(i))
            fmt.line.color.rgb = RGBColor.from_string("FFFFFF")
            fmt.line.width = Pt(1.5)
        self._legend(ch, has_legend, legend_pos)
        if labels is not False:
            within = [inside(i, radius) for i in range(len(slices))]
            self._slice_labels(ch, ser, slices, sizes, within, (x, y, w, h), radius, hole,
                               show_category, number_format)
        if center_label:
            self._center_label(s, center_label, cx, cy, radius * hole)

    @staticmethod
    def _reach(radius, mid, lw, lh):
        """Distance from the centre to an outside label's centre at angle `mid`."""
        return radius + 0.1 + 0.5 * (abs(math.sin(mid)) * lw + abs(math.cos(mid)) * lh)

    def _slice_labels(self, ch, ser, slices, sizes, inside, box, radius, hole, show_category,
                      number_format):
        """Per-slice data labels: inside ones in a colour that contrasts with the slice, outside
        ones in the chart ink (a doughnut's are moved out of the ring by an offset)."""
        points = []
        x, y, w, h = box
        fmt = number_format or "0%"
        for i, ((mid, _), (lw, lh), within) in enumerate(zip(slices, sizes, inside)):
            color = on_color(self._color(i), self._hx("text")) if within else self._ink
            offset = ""
            pos = "" if hole else f'<c:dLblPos val="{"ctr" if within else "outEnd"}"/>'
            if hole and not within:
                move = self._reach(radius, mid, lw, lh) - radius * (1 + hole) / 2
                offset = (f'<c:layout><c:manualLayout><c:x val="{move * math.sin(mid) / w:.5f}"/>'
                          f'<c:y val="{-move * math.cos(mid) / h:.5f}"/></c:manualLayout>'
                          f'</c:layout>')
            points.append(
                f'<c:dLbl><c:idx val="{i}"/>{offset}<c:numFmt formatCode={quoteattr(fmt)} '
                f'sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
                f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr '
                f'sz="{CHART_PT["label"] * 100}" b="1">{_solid_xml(color)}<a:latin '
                f'typeface={quoteattr(self.FONT)}/></a:defRPr></a:pPr><a:endParaRPr '
                f'lang="en-US"/></a:p></c:txPr>{pos}<c:showLegendKey val="0"/>'
                f'<c:showVal val="0"/><c:showCatName val="{int(show_category)}"/>'
                f'<c:showSerName val="0"/><c:showPercent val="1"/><c:showBubbleSize val="0"/>'
                f'<c:separator>\n</c:separator></c:dLbl>')
        dlbls = parse_xml(
            f'<c:dLbls xmlns:c="{NS_C}" xmlns:a="{NS_A}">{"".join(points)}'
            f'<c:numFmt formatCode={quoteattr(fmt)} sourceLinked="0"/>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:showLegendKey val="0"/>'
            f'<c:showVal val="0"/><c:showCatName val="{int(show_category)}"/>'
            f'<c:showSerName val="0"/><c:showPercent val="1"/><c:showBubbleSize val="0"/>'
            f'<c:separator>\n</c:separator><c:showLeaderLines val="0"/></c:dLbls>')
        _replace(ser._element, dlbls, *_AFTER_DLBLS)
        plot_dlbls = ch.plots[0]._element.find(qn("c:dLbls"))
        if plot_dlbls is not None:
            ch.plots[0]._element.remove(plot_dlbls)

    def _center_label(self, s, text, cx, cy, hole_radius):
        """Text centred in a doughnut's hole: the first line large and bold, the rest smaller."""
        big, small = 22, 14
        lines = text.split("\n")
        need_w = max([self.measure(lines[0], big, bold=True)] +
                     [self.measure(t, small) for t in lines[1:]])
        tw = hole_radius * 1.6
        if need_w > tw:
            raise ValueError(f"center_label {text!r} is too wide for the doughnut's hole; "
                             "shorten it")
        th = (big + small * (len(lines) - 1)) * 1.25 / 72 + 0.1
        self._note(s, (cx - tw / 2, cy - th / 2, tw, th),
                   [(t, big if i == 0 else small, self._hx("text") if i == 0 else self._ink, i == 0)
                    for i, t in enumerate(lines)], f"Chart label: {lines[0]}", align="center")

    # -- bar & line -----------------------------------------------------------------------------
    def bar_line(self, s, categories, bars, lines, title=None, *, box=CONTENT, labels=True,
                 dashed=False, number_format=None, x_title=None, y_title=None, legend_pos="top",
                 alt=None, name=None):
        """A combo chart: columns and lines on the same axes; returns the graphic frame.

        bars / lines: a (name, values) series or a list of them, one value per category.
        labels: value labels on the columns. dashed: dashed lines (e.g. a target).
        """
        bars = [bars] if isinstance(bars[0], str) else list(bars)
        lines = [lines] if isinstance(lines[0], str) else list(lines)
        series = bars + lines
        for sname, values in series:
            if len(values) != len(categories):
                raise ValueError(f"series {sname!r} has {len(values)} values, "
                                 f"expected {len(categories)}")
        data = CategoryChartData(number_format=number_format or "General")
        data.categories = categories
        for sname, values in series:
            data.add_series(sname, values)
        detail = (f"{', '.join(n for n, _ in bars)} as columns and "
                  f"{', '.join(n for n, _ in lines)} as lines; {_summary(categories, series)}")
        gf = self._frame(s, XL_CHART_TYPE.COLUMN_CLUSTERED, box, data,
                         name or f"Chart: {title or 'bar and line chart'}",
                         self._chart_alt(alt, "Bar and line chart", title, detail))
        ch = gf.chart
        self._style(ch, title, box)
        bar_el = ch.plots[0]._element
        line_el = parse_xml(f'<c:lineChart xmlns:c="{NS_C}"><c:grouping val="standard"/>'
                            f'<c:varyColors val="0"/></c:lineChart>')
        for ser_el in bar_el.findall(qn("c:ser"))[len(bars):]:
            bar_el.remove(ser_el)
            for tag in ("c:invertIfNegative", "c:spPr"):
                for el in ser_el.findall(qn(tag)):
                    ser_el.remove(el)
            etree.SubElement(ser_el, qn("c:smooth")).set("val", "0")
            line_el.append(ser_el)
        etree.SubElement(line_el, qn("c:marker")).set("val", "1")
        for ax in bar_el.findall(qn("c:axId")):
            etree.SubElement(line_el, qn("c:axId")).set("val", ax.get("val"))
        bar_el.addnext(line_el)
        bar_plot, line_plot = ch.plots[0], ch.plots[1]
        bar_plot.gap_width = 60
        bar_plot.overlap = -10 if len(bars) > 1 else 0
        for i, ser in enumerate(bar_plot.series):
            _series_sppr(ser, self._color(i))
            if labels:
                self._labels(ser, on_color(self._color(i), self._hx("text")), number_format,
                             XL_LABEL_POSITION.CENTER)
        for i, ser in enumerate(line_plot.series):
            color = self._color(max(len(bars), 2) + i)
            _series_sppr(ser, line=color, width=2.75, dash="dash" if dashed else None)
            self._series_marker(ser, color, 8)
        values = [v for _, vals in series for v in vals]
        lo, hi, step = _nice(min(0, min(values)), max(values) * 1.04)
        self._axis(ch.value_axis, grid=True, fmt=number_format, title=y_title, minimum=lo,
                   maximum=hi, step=step, vertical=True)
        self._axis(ch.category_axis, line=True, title=x_title)
        self._legend(ch, True, legend_pos)
        return gf

    # -- scatter and bubble -----------------------------------------------------------------------
    def _xy_axes(self, ch, xs, ys, pad, x_title, y_title, x_format, y_format):
        """Nice x and y scales around the points, padded by `pad` (shares of each range), and
        not below 0 for values that are not."""
        scales = []
        for values, share in ((xs, pad[0]), (ys, pad[1])):
            lo, hi = min(values), max(values)
            margin = (hi - lo) * share or 1
            scales.append(_nice(lo - margin if lo < 0 else max(0, lo - margin), hi + margin))
        (x_min, x_max, x_step), (y_min, y_max, y_step) = scales
        self._axis(ch.category_axis, grid=True, line=True, fmt=x_format, title=x_title,
                   minimum=x_min, maximum=x_max, step=x_step)
        self._axis(ch.value_axis, grid=True, fmt=y_format, title=y_title, minimum=y_min,
                   maximum=y_max, step=y_step, vertical=True)

    def scatter(self, s, series, title=None, *, box=CONTENT, x_title=None, y_title=None,
                x_format=None, y_format=None, marker_size=10, legend_pos="top", alt=None,
                name=None):
        """An XY scatter chart of [(name, [(x, y), ...])]; returns the graphic frame.
        The legend shows only for several series."""
        series = _pairs(series)
        data = XyChartData()
        for sname, points in series:
            ser = data.add_series(sname)
            for px, py in points:
                ser.add_data_point(px, py)
        detail = "; ".join(f"{n}: {len(p)} points, x {_fmt(min(q[0] for q in p))} to "
                           f"{_fmt(max(q[0] for q in p))}, y {_fmt(min(q[1] for q in p))} to "
                           f"{_fmt(max(q[1] for q in p))}" for n, p in series)
        gf = self._frame(s, XL_CHART_TYPE.XY_SCATTER, box, data,
                         name or f"Chart: {title or 'scatter chart'}",
                         self._chart_alt(alt, "Scatter chart", title, detail))
        ch = gf.chart
        self._style(ch, title, box)
        for i, ser in enumerate(ch.plots[0].series):
            color = self._color(i)
            _series_sppr(ser)
            self._series_marker(ser, color, marker_size)
            ser.marker.format.line.width = Pt(1.25)
        self._xy_axes(ch, [p[0] for _, pts in series for p in pts],
                      [p[1] for _, pts in series for p in pts], (0.06, 0.08), x_title, y_title,
                      x_format, y_format)
        self._legend(ch, len(series) > 1, legend_pos)
        return gf

    def bubble(self, s, series, title=None, *, box=CONTENT, x_title=None, y_title=None,
               x_format=None, y_format=None, size_format=None, labels=True, scale=150,
               alt=None, name=None):
        """A bubble chart of [(name, [(x, y, size), ...])]; bubble areas follow `size`.
        labels: each bubble's series name and size above it (size_format, e.g. '"$"#,##0');
        with labels there is no legend. Returns the graphic frame."""
        series = _pairs(series)
        data = BubbleChartData()
        for sname, points in series:
            ser = data.add_series(sname)
            for px, py, size in points:
                ser.add_data_point(px, py, size)
        detail = "; ".join(f"{n}: " + ", ".join(f"x {_fmt(p[0])}, y {_fmt(p[1])}, size "
                                                f"{_fmt(p[2])}" for p in pts)
                           for n, pts in series)
        gf = self._frame(s, XL_CHART_TYPE.BUBBLE, box, data,
                         name or f"Chart: {title or 'bubble chart'}",
                         self._chart_alt(alt, "Bubble chart", title, detail))
        ch = gf.chart
        self._style(ch, title, box)
        plot = ch.plots[0]
        plot.bubble_scale = scale
        for i, ser in enumerate(plot.series):
            _series_sppr(ser, self._color(i), 0.75, "FFFFFF", 1.0)
            if labels:
                dls = self._labels(ser, self._hx("text"), size_format, XL_LABEL_POSITION.ABOVE,
                                   value=False, series=True)
                dls_el = dls._element
                _set_val(dls_el, "c:showBubbleSize", 1, "c:separator", "c:showLeaderLines",
                         "c:extLst")
                sep = etree.Element(qn("c:separator"))
                sep.text = "\n"
                _replace(dls_el, sep, "c:showLeaderLines", "c:leaderLines", "c:extLst")
        # pad the axes by the largest bubble's radius (25% of the plot height at scale 100)
        reach = 0.125 * scale / 100
        self._xy_axes(ch, [p[0] for _, pts in series for p in pts],
                      [p[1] for _, pts in series for p in pts],
                      (reach * 1.3 * box[3] / box[2] + 0.05, reach * 1.6 + 0.1), x_title, y_title,
                      x_format, y_format)
        self._legend(ch, not labels and len(series) > 1, "top")
        return gf

    # -- histogram -------------------------------------------------------------------------------
    def histogram(self, s, samples, bins=30, title=None, *, box=CONTENT, x_title=None,
                  y_title="Count", x_format=None, opacity=0.6, legend_pos="top", alt=None,
                  name=None):
        """Overlaid histograms of {name: values}, pre-binned into about `bins` bins of a nice
        width shared by every series (the counts stay editable); returns the graphic frame.
        Columns touch and overlap translucently; bins are labelled by their lower edge."""
        series = _pairs(samples)
        values = [v for _, vals in series for v in vals]
        if not values:
            raise ValueError("histogram() needs at least one value")
        lo, hi = min(values), max(values)
        width = _bin_width(lo, hi, bins)
        label_step = _nice(lo, hi, 8)[2]
        skip = max(1, round(label_step / width))
        start = math.floor(lo / (width * skip) + 1e-9) * width * skip
        count = max(1, math.floor((hi - start) / width + 1e-9) + 1)
        edges = [round(start + i * width, 10) for i in range(count)]
        data = CategoryChartData(number_format="0")
        data.categories = edges
        counts = []
        for sname, vals in series:
            bucket = [0] * count
            for v in vals:
                bucket[min(int((v - start) / width + 1e-9), count - 1)] += 1
            counts.append(bucket)
            data.add_series(sname, bucket)
        detail = (f"{', '.join(f'{n} ({len(v)} values)' for n, v in series)} in {count} bins of "
                  f"{_fmt(width)} from {_fmt(start)} to {_fmt(start + count * width)}")
        gf = self._frame(s, XL_CHART_TYPE.COLUMN_CLUSTERED, box, data,
                         name or f"Chart: {title or 'histogram'}",
                         self._chart_alt(alt, "Histogram", title, detail))
        ch = gf.chart
        self._style(ch, title, box)
        plot = ch.plots[0]
        plot.gap_width, plot.overlap = 0, 100
        for i, ser in enumerate(plot.series):
            _series_sppr(ser, self._color(i), opacity)
        _, y_hi, y_step = _nice(0, max(max(c) for c in counts))
        self._axis(ch.value_axis, grid=True, title=y_title, minimum=0, maximum=y_hi, step=y_step,
                   vertical=True)
        cat_axis = ch.category_axis
        self._axis(cat_axis, line=True, fmt=x_format or "General", title=x_title)
        cat_el = cat_axis._element
        for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
            _set_val(cat_el, tag, skip, "c:noMultiLvlLbl", "c:extLst")
        self._legend(ch, len(series) > 1, legend_pos)
        return gf

    # -- candlestick -----------------------------------------------------------------------------
    def candlestick(self, s, dates, ohlc, title=None, *, box=CONTENT, y_title=None,
                    y_format=None, date_format="%a %b %d", up=None, down=None, alt=None,
                    name=None):
        """A native stock chart (open-high-low-close) with high-low lines and up/down bars;
        returns the graphic frame. dates: labels, datetime.date values or ISO 'YYYY-MM-DD'
        strings (shown with date_format); ohlc: [(open, high, low, close)] per date.
        up / down: bar colours (default: palette slots 4 and 5, green and rose)."""
        if len(dates) != len(ohlc):
            raise ValueError(f"{len(dates)} dates but {len(ohlc)} open-high-low-close rows")
        for i, (o, hi, lo, c) in enumerate(ohlc):
            if not lo <= min(o, c) <= max(o, c) <= hi:
                raise ValueError(f"row {i}: low <= open, close <= high does not hold")
        up, down = _hex(up or self._color(3)), _hex(down or self._color(4))
        labels = [_date_label(d, date_format) for d in dates]
        data = CategoryChartData(number_format=y_format or "General")
        data.categories = labels
        for i, sname in enumerate(("Open", "High", "Low", "Close")):
            data.add_series(sname, [row[i] for row in ohlc])
        rises = sum(1 for o, _, _, c in ohlc if c >= o)
        detail = (f"{len(ohlc)} trading days from {labels[0]} to {labels[-1]}; low "
                  f"{_fmt(min(r[2] for r in ohlc))}, high {_fmt(max(r[1] for r in ohlc))}, first "
                  f"open {_fmt(ohlc[0][0])}, last close {_fmt(ohlc[-1][3])}; {rises} up days, "
                  f"{len(ohlc) - rises} down days")
        gf = self._frame(s, XL_CHART_TYPE.LINE, box, data,
                         name or f"Chart: {title or 'candlestick chart'}",
                         self._chart_alt(alt, "Candlestick chart", title, detail))
        ch = gf.chart
        self._style(ch, title, box)
        plot_area = ch._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
        line_el = plot_area.find(qn("c:lineChart"))
        stock = etree.Element(qn("c:stockChart"))
        for ser_el in line_el.findall(qn("c:ser")):
            for tag in ("c:spPr", "c:marker", "c:smooth"):
                for el in ser_el.findall(qn(tag)):
                    ser_el.remove(el)
            ser_el.find(qn("c:tx")).addnext(parse_xml(
                f'<c:spPr xmlns:c="{NS_C}" xmlns:a="{NS_A}"><a:ln w="19050"><a:noFill/></a:ln>'
                f'</c:spPr>'))
            ser_el.find(qn("c:spPr")).addnext(parse_xml(
                f'<c:marker xmlns:c="{NS_C}"><c:symbol val="none"/></c:marker>'))
            stock.append(ser_el)

        def bars(tag, color):
            return (f'<c:{tag}><c:spPr>{_solid_xml(color, 0.55)}<a:ln w="19050">'
                    f'{_solid_xml(color)}</a:ln></c:spPr></c:{tag}>')

        stock.append(parse_xml(
            f'<c:hiLowLines xmlns:c="{NS_C}" xmlns:a="{NS_A}"><c:spPr><a:ln w="19050">'
            f'{_solid_xml(self._hx("line"))}</a:ln></c:spPr></c:hiLowLines>'))
        stock.append(parse_xml(
            f'<c:upDownBars xmlns:c="{NS_C}" xmlns:a="{NS_A}"><c:gapWidth val="80"/>'
            f'{bars("upBars", up)}{bars("downBars", down)}</c:upDownBars>'))
        for ax in line_el.findall(qn("c:axId")):
            stock.append(ax)
        line_el.addprevious(stock)
        plot_area.remove(line_el)
        lo, hi, step = _nice(min(r[2] for r in ohlc), max(r[1] for r in ohlc), 8)
        self._axis(ch.value_axis, grid=True, fmt=y_format, title=y_title, minimum=lo, maximum=hi,
                   step=step, vertical=True)
        self._axis(ch.category_axis, line=True)
        self._legend(ch, False, "top")
        return gf

    # -- Office 2016 charts -----------------------------------------------------------------------
    def _chartex_sizes(self):
        return {"axisTitle": CHART_PT["tick"], "categoryAxis": CHART_PT["tick"],
                "dataLabel": CHART_PT["label"], "legend": CHART_PT["legend"],
                "title": CHART_PT["title"], "valueAxis": CHART_PT["tick"]}

    def _chartex_axis_title(self, text, axis):
        return chartex.title(text, CHART_PT["tick"], self._ink, self.FONT, axis=axis) if text \
            else ""

    def box_plot(self, s, samples, title=None, *, box=CONTENT, x_title=None, y_title=None,
                 y_format=None, mean=True, outliers=True, quartiles="exclusive", alt=None,
                 name=None):
        """An Office 2016 box & whisker chart of {group: values}: one box per group (median,
        quartiles, whiskers, outliers and a mean marker); returns the graphic frame. The values
        stay editable through Edit Data. quartiles: 'exclusive' or 'inclusive' median.
        Needs PowerPoint 2016 or later (verified on Microsoft 365)."""
        series = _pairs(samples)
        if quartiles not in ("exclusive", "inclusive"):
            raise ValueError("quartiles must be 'exclusive' or 'inclusive'")
        groups = [g for g, vals in series for _ in vals]
        values = [v for _, vals in series for v in vals]
        if not values:
            raise ValueError("box_plot() needs at least one value")
        group_head, value_head = x_title or "Group", y_title or "Value"
        color = self._color(0)
        data = (f'<cx:data id="0">{chartex.str_dim([groups], chartex.range_ref(0, len(groups)))}'
                f'{chartex.num_dim(values, chartex.range_ref(1, len(values)))}</cx:data>')
        layout = (f'<cx:layoutPr><cx:visibility meanLine="0" meanMarker="{int(mean)}" '
                  f'nonoutliers="0" outliers="{int(outliers)}"/><cx:statistics '
                  f'quartileMethod="{quartiles}"/></cx:layoutPr>')
        ser = chartex.series("boxWhisker", 0, value_head, chartex.cell_ref(1, 0),
                             chartex.sppr(color, 0.3, color, 1.5), layout)
        num_fmt = (f'<cx:numFmt formatCode={quoteattr(y_format)} sourceLinked="0"/>'
                   if y_format else "")
        grid = chartex.sppr(line=self._hx("card_border"), line_w=GRID_PT)
        tick = chartex.text_props(CHART_PT["tick"], self._ink, self.FONT)
        head = chartex.title(title, CHART_PT["title"], self._hx("text"), self.FONT) if title else ""
        chart = (f'{head}<cx:plotArea><cx:plotAreaRegion>{ser}</cx:plotAreaRegion>'
                 f'<cx:axis id="0"><cx:catScaling gapWidth="1"/>'
                 f'{self._chartex_axis_title(x_title, "cat")}<cx:tickLabels/>'
                 f'{chartex.sppr(line=self._hx("subtle"), line_w=GRID_PT)}{tick}</cx:axis>'
                 f'<cx:axis id="1"><cx:valScaling/>{self._chartex_axis_title(y_title, "val")}'
                 f'<cx:majorGridlines>{grid}</cx:majorGridlines><cx:tickLabels/>{num_fmt}'
                 f'{chartex.sppr(no_fill=True)}{tick}</cx:axis></cx:plotArea>')
        stats = []
        for group, vals in series:
            ordered = sorted(vals)
            mid = len(ordered) // 2
            median = ordered[mid] if len(ordered) % 2 else (ordered[mid - 1] + ordered[mid]) / 2
            stats.append(f"{group} ({len(vals)} values, median {_fmt(median)}, range "
                         f"{_fmt(ordered[0])} to {_fmt(ordered[-1])})")
        return chartex.add_chartex(
            s, box, chartex.chart_space(data, chart, tick),
            name or f"Chart: {title or 'box plot'}",
            self._chart_alt(alt, "Box plot", title, "; ".join(stats)),
            chartex.workbook([(group_head, groups), (value_head, values)]),
            self._chartex_sizes(), self._ink, {})

    def sunburst(self, s, tree, title=None, *, box=CONTENT, levels=None, value_label="Value",
                 number_format="General", alt=None, name=None):
        """An Office 2016 sunburst chart of a hierarchy; returns the graphic frame.

        tree: [(name, children)] where children is again such a list or, at the leaves, a
        number: [("Fiction", [("Fantasy", 120), ("Mystery", 90)]), ...]. Every leaf sits at the
        same depth. levels: workbook column headers per level (default 'Level 1', ...).
        Inner rings take the palette darkened, outer rings the palette, with white labels
        (name and value). Needs PowerPoint 2016 or later (verified on Microsoft 365)."""
        paths = []

        def walk(items, prefix):
            for item in items:
                if not (isinstance(item, (tuple, list)) and len(item) == 2):
                    raise ValueError(f"sunburst nodes are (name, children or value), got {item!r}")
                label, rest = item
                if isinstance(rest, (int, float)):
                    paths.append((prefix + [str(label)], rest))
                else:
                    walk(rest, prefix + [str(label)])

        walk(tree, [])
        if not paths:
            raise ValueError("sunburst() needs at least one leaf")
        depth = len(paths[0][0])
        if any(len(p) != depth for p, _ in paths):
            raise ValueError("every sunburst leaf must sit at the same depth")
        heads = list(levels or [f"Level {i}" for i in range(1, depth + 1)])
        if len(heads) != depth:
            raise ValueError(f"levels names {len(heads)} columns but the tree is {depth} deep")
        columns = [(heads[d], [p[d] for p, _ in paths]) for d in range(depth)]
        values = [v for _, v in paths]
        cats = chartex.str_dim([col for _, col in reversed(columns)],
                               chartex.range_ref(0, len(paths), last_col=depth - 1))
        data = (f'<cx:data id="0">{cats}'
                f'{chartex.num_dim(values, chartex.range_ref(depth, len(paths)), "size")}'
                f'</cx:data>')
        # data points are the tree's nodes in pre-order (each parent before its children)
        points = []
        roots = list(dict.fromkeys(p[0] for p, _ in paths))
        seen = {}
        for path, _ in paths:
            for d in range(depth):
                key = tuple(path[:d + 1])
                if key in seen:
                    continue
                seen[key] = len(seen)
                shade = blend(self._color(roots.index(path[0])), self._hx("text"),
                              0.22 * (depth - 1 - d) / max(depth - 1, 1))
                points.append(f'<cx:dataPt idx="{seen[key]}">'
                              f'{chartex.sppr(shade, None, "FFFFFF", 1.5)}</cx:dataPt>')
        labels = (f'<cx:dataLabels pos="ctr"><cx:numFmt formatCode={quoteattr(number_format)} '
                  f'sourceLinked="0"/><cx:visibility seriesName="0" categoryName="1" '
                  f'value="1"/><cx:separator>\n</cx:separator></cx:dataLabels>')
        ser = chartex.series("sunburst", 0, value_label, chartex.cell_ref(depth, 0),
                             chartex.sppr(None, None, "FFFFFF", 1.5) + "".join(points) + labels)
        head = chartex.title(title, CHART_PT["title"], self._hx("text"), self.FONT) if title else ""
        chart = f'{head}<cx:plotArea><cx:plotAreaRegion>{ser}</cx:plotAreaRegion></cx:plotArea>'
        totals = {}
        for path, v in paths:
            totals[path[0]] = totals.get(path[0], 0) + v
        detail = ", ".join(f"{k} {_fmt(v)}" for k, v in totals.items())
        return chartex.add_chartex(
            s, box, chartex.chart_space(data, chart, chartex.text_props(
                CHART_PT["tick"], self._ink, self.FONT)),
            name or f"Chart: {title or 'sunburst chart'}",
            self._chart_alt(alt, "Sunburst chart", title, f"{len(paths)} leaves; {detail}"),
            chartex.workbook(columns + [(value_label, values)]), self._chartex_sizes(),
            self._ink, {"dataLabel": "FFFFFF"})

    # -- heatmap ---------------------------------------------------------------------------------
    def heatmap(self, s, rows, cols, values, title=None, *, box=CONTENT, value_format="{:g}",
                scale_label=None, corner="", alt=None, name=None):
        """A heatmap: a native table whose cells are coloured on a scale from the palette's first
        colour, with a gradient legend; returns the table. rows / cols: labels; values: one row
        of numbers per row label. Cell text is the value (value_format), light or dark to
        contrast with its fill; scale_label titles the legend (e.g. 'Orders'); corner is the
        text of the top-left cell. The fills are computed once: editing a value does not
        recolour its cell."""
        if not rows or not cols or len(values) != len(rows) or \
                any(len(r) != len(cols) for r in values):
            raise ValueError(f"values must be {len(rows)} rows of {len(cols)} numbers")
        cell_pt, head_pt, title_pt = 14, CHART_PT["tick"], CHART_PT["title"]
        x, y, w, h = box
        flat = [v for r in values for v in r]
        lo, hi = min(flat), max(flat)
        base = self._color(0)
        stops = [(0.0, blend("FFFFFF", base, 0.08)), (0.5, base),
                 (1.0, blend(base, self._hx("text"), 0.6))]

        def scale(v):
            t = 0.0 if hi == lo else (v - lo) / (hi - lo)
            for (t0, c0), (t1, c1) in zip(stops, stops[1:]):
                if t <= t1:
                    return blend(c0, c1, (t - t0) / (t1 - t0))
            return stops[-1][1]

        mid = (lo + hi) / 2
        ticks = [(hi, value_format.format(hi)),
                 (mid, value_format.format(round(mid) if all(float(v).is_integer() for v in flat)
                                           else mid)),
                 (lo, value_format.format(lo))]
        texts = [[value_format.format(v) for v in r] for r in values]
        tick_w = max(self.measure(t, head_pt) for _, t in ticks)
        legend_w = max(0.3 + 0.12 + tick_w + 0.08,
                       self.measure(scale_label, head_pt) if scale_label else 0.0) + 0.1
        title_lines = self.wrap(title, w - legend_w, title_pt) if title else []
        top = len(title_lines) * title_pt * 1.2 / 72 + 0.2 if title else 0.0
        label_w = max(self.measure(str(r), head_pt) for r in [*rows, corner]) + 0.2
        table_w = w - legend_w - 0.25
        cell_w = (table_w - label_w) / len(cols)
        need = max([self.measure(str(c), head_pt) for c in cols] +
                   [self.measure(t, cell_pt) for r in texts for t in r]) + 0.16
        if cell_w < need:
            raise ValueError(f"heatmap columns are {cell_w:.2f} in wide but need {need:.2f} in; "
                             "split the heatmap or shorten the labels")
        head_h = 0.4
        row_h = (h - top - head_h) / len(rows)
        if row_h < cell_pt * 1.2 / 72 + 0.12:
            raise ValueError(f"heatmap rows are too short for {cell_pt} pt values; split the "
                             "heatmap")
        if title:
            self._note(s, (x, y, w - legend_w, top - 0.1),
                       [("\n".join(title_lines), title_pt, self._hx("text"), False)],
                       f"Title: {title}", anchor="top")
        peak = max(((r, c) for r in range(len(rows)) for c in range(len(cols))),
                   key=lambda rc: values[rc[0]][rc[1]])
        detail = (f"{len(rows)} rows ({rows[0]} to {rows[-1]}) by {len(cols)} columns "
                  f"({cols[0]} to {cols[-1]}); values {ticks[2][1]} to {ticks[0][1]}, highest "
                  f"at {rows[peak[0]]}, {cols[peak[1]]}; darker cells hold higher values")
        tbl = self.table(s, x, y + top, table_w, h - top, [corner] + [str(c) for c in cols],
                         [[str(r)] + t for r, t in zip(rows, texts)],
                         col_widths=[label_w] + [cell_w] * len(cols),
                         col_align=["right"] + ["center"] * len(cols), font_size=cell_pt,
                         alt=self._chart_alt(alt, "Heatmap", title, detail))
        tbl._graphic_frame.name = name or f"Chart: {title or 'heatmap'}"
        tbl.rows[0].height = Emu(round(head_h * EMU_IN))
        for r in range(1, len(rows) + 1):
            tbl.rows[r].height = Emu(round(row_h * EMU_IN))
        gap = _hex(self.LIGHT)
        for r in range(len(rows) + 1):
            for c in range(len(cols) + 1):
                cell = tbl.cell(r, c)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = cell.margin_right = Emu(round(0.06 * EMU_IN))
                font = cell.text_frame.paragraphs[0].runs[0].font
                if r == 0 or c == 0:
                    cell.fill.background()
                    font.size, font.bold = Pt(head_pt), False
                    font.color.rgb = RGBColor.from_string(self._ink)
                    continue
                fill = scale(values[r - 1][c - 1])
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(fill)
                font.color.rgb = RGBColor.from_string(on_color(fill, self._hx("text")))
                tc_pr = cell._tc.get_or_add_tcPr()
                for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                    for old in tc_pr.findall(qn(tag)):
                        tc_pr.remove(old)
                    tc_pr.insert(i, parse_xml(f'<{tag} xmlns:a="{NS_A}" w="28575">'
                                              f'{_solid_xml(gap)}</{tag}>'))
        self._scale_legend(s, (x + w - legend_w + 0.1, y + top + head_h, legend_w - 0.1,
                               h - top - head_h), stops, ticks, lo, hi, tick_w, scale_label)
        return tbl

    def _scale_legend(self, s, box, stops, ticks, lo, hi, tick_w, scale_label):
        """A vertical gradient bar (high values at the top) with tick labels and a title."""
        x, y, w, h = box
        if scale_label:
            self._note(s, (x, y, w, 0.3), [(scale_label, CHART_PT["tick"], self._ink, False)],
                       f"Legend: {scale_label}", anchor="bottom")
            y, h = y + 0.4, h - 0.4
        # inset so the end tick labels, centred on the bar's ends, stay inside the box
        y, h = y + 0.13, h - 0.26
        bar = self.rect(s, x, y, 0.3, h, fill=self.WHITE)
        bar.name = "Legend: colour scale"
        sppr = bar._element.spPr
        sppr.find(qn("a:solidFill")).addprevious(parse_xml(
            f'<a:gradFill xmlns:a="{NS_A}" rotWithShape="1"><a:gsLst>'
            + "".join(f'<a:gs pos="{round((1 - t) * 100000)}"><a:srgbClr val="{c}"/></a:gs>'
                      for t, c in reversed(stops))
            + '</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>'))
        sppr.remove(sppr.find(qn("a:solidFill")))
        for v, text in ticks:
            ty = y + h * (1 - (0.5 if hi == lo else (v - lo) / (hi - lo)))
            self._note(s, (x + 0.42, ty - 0.13, tick_w + 0.08, 0.26),
                       [(text, CHART_PT["tick"], self._ink, False)], f"Legend: {text}")
