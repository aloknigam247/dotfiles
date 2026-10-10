"""Diagram primitives mixed into Deck: connection sites, glued orthogonal links with arrowheads and
label chips, custom-site geometry, theme-font text measuring, alt text and groups.

Shape names follow one convention that scripts/com_check.py relies on: 'Node: ...' shapes and
'Label: ...' chips must not overlap each other, 'Container: ...' frames contain nodes, 'Edge: ...'
connectors are glued at both ends, and 'Deco: ...' shapes are not checked.
"""
import math
import os
from dataclasses import dataclass

from lxml import etree
from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .style import EDGE_WIDTH, MIN_PT, NODE_LINE_WIDTH, NODE_RADIUS, TYPE_SCALE, kind_style

ADJ_MAX = 2_000_000_000  # adjust values are 32-bit integers
CHIP_PAD = (0.06, 0.02)  # inches: text insets of a label chip
EMU_IN = 914400
EXT_MIN = 127  # EMU: connector frame extent along an axis the route has no net run on
LINE_SPACING = 1.2  # PowerPoint's single line height / font size (measured through COM)
MAX_SEGMENTS = 5  # bentConnector5 is the longest preset connector
MEASURE_PX = 1000  # Pillow measures at this size and scales, for fractional advance widths
MIN_FIRST, MIN_LAST = 0.03, 0.06  # inches: shortest first segment and arrowhead segment
# Pillow width x TEXT_MARGIN + TEXT_PAD_PT >= the narrowest width PowerPoint keeps on one line.
# Measured through COM (Calibri Light, Calibri; regular and bold; 11-18 pt; 256 strings): the
# widest ratio was 1.0163, so no measured line needs more than this margin.
TEXT_MARGIN = 1.03
TEXT_PAD_PT = 0.5

_ANGLES = {"cd2": 10800000, "cd4": 5400000, "cd8": 2700000, "3cd4": 16200000, "3cd8": 8100000,
           "5cd8": 13500000, "7cd8": 18900000}
_OPS = {
    "*/": lambda x, y, z: x * y / z,
    "+-": lambda x, y, z: x + y - z,
    "+/": lambda x, y, z: (x + y) / z,
    "?:": lambda x, y, z: y if x > 0 else z,
    "abs": abs,
    "at2": lambda x, y: math.degrees(math.atan2(y, x)) * 60000,
    "cat2": lambda x, y, z: x * math.cos(math.atan2(z, y)),
    "cos": lambda x, y: x * math.cos(math.radians(y / 60000)),
    "max": max,
    "min": min,
    "mod": lambda x, y, z: math.sqrt(x * x + y * y + z * z),
    "pin": lambda x, y, z: x if y < x else (z if y > z else y),
    "sat2": lambda x, y, z: x * math.sin(math.atan2(z, y)),
    "sin": lambda x, y: x * math.sin(math.radians(y / 60000)),
    "sqrt": math.sqrt,
    "tan": lambda x, y: x * math.tan(math.radians(y / 60000)),
    "val": lambda x: x,
}

_OVAL_GD = ("idx cos wd2 2700000", "idy sin hd2 2700000", "il +- hc 0 idx", "ir +- hc idx 0",
            "it +- vc 0 idy", "ib +- vc idy 0")
_OVAL = ("3cd4 hc t", "3cd4 il it", "cd2 l vc", "cd4 il ib", "cd4 hc b", "cd4 ir ib", "0 r vc",
         "3cd4 ir it")
_RECT = ("3cd4 hc t", "cd2 l vc", "cd4 hc b", "0 r vc")
_RECT_RBLT = ("0 r vc", "cd4 hc b", "cd2 l vc", "3cd4 hc t")

# preset: (adjust defaults, guides "name op args", connection sites "angle x y", text rectangle
# "l t r b"), from the ECMA-376 preset shape definitions; positions checked against PowerPoint 16
# through COM at several sizes and adjust values.
PRESETS = {
    "can": ({"adj": 25000}, ("maxAdj */ 50000 h ss", "a pin 0 adj maxAdj", "y1 */ ss a 200000",
                             "y2 +- y1 y1 0", "y3 +- b 0 y1"),
            ("3cd4 hc y2", "3cd4 hc t", "cd2 l vc", "cd4 hc b", "0 r vc"), "l y2 r y3"),
    "chevron": ({"adj": 50000}, ("maxAdj */ 100000 w ss", "a pin 0 adj maxAdj", "x1 */ ss a 100000",
                                 "x2 +- r 0 x1", "x3 */ x2 1 2", "dx +- x2 0 x1", "il ?: dx x1 l",
                                 "ir ?: dx x2 r"),
                ("3cd4 x3 t", "cd2 x1 vc", "cd4 x3 b", "0 r vc"), "il t ir b"),
    "diamond": ({}, ("ir */ w 3 4", "ib */ h 3 4"), _RECT, "wd4 hd4 ir ib"),
    "ellipse": ({}, _OVAL_GD, _OVAL, "il it ir ib"),
    "flowChartAlternateProcess": ({}, ("il */ ssd6 29289 100000", "ir +- r 0 il", "ib +- b 0 il"),
                                  _RECT, "il il ir ib"),
    "flowChartConnector": ({}, _OVAL_GD, _OVAL, "il it ir ib"),
    "flowChartDecision": ({}, ("ir */ w 3 4", "ib */ h 3 4"), _RECT, "wd4 hd4 ir ib"),
    "flowChartDelay": ({}, ("idx cos wd2 2700000", "idy sin hd2 2700000", "ir +- hc idx 0",
                            "it +- vc 0 idy", "ib +- vc idy 0"), _RECT, "l it ir ib"),
    "flowChartDocument": ({}, ("y1 */ h 17322 21600", "y2 */ h 20172 21600"),
                          ("3cd4 hc t", "cd2 l vc", "cd4 hc y2", "0 r vc"), "l t r y1"),
    "flowChartInputOutput": ({}, ("x3 */ w 2 5", "x4 */ w 3 5", "x5 */ w 4 5", "x6 */ w 9 10"),
                             ("3cd4 x4 t", "3cd4 hc t", "cd2 wd10 vc", "cd4 x3 b", "cd4 hc b",
                              "0 x6 vc"), "wd5 t x5 b"),
    "flowChartMagneticDisk": ({}, ("y3 */ h 1 3", "y2 */ h 5 6"),
                              ("3cd4 hc y3", "3cd4 hc t", "cd2 l vc", "cd4 hc b", "0 r vc"),
                              "l y3 r y2"),
    "flowChartManualInput": ({}, (), ("3cd4 hc hd10", "cd2 l vc", "cd4 hc b", "0 r vc"),
                             "l hd5 r b"),
    "flowChartManualOperation": ({}, ("x3 */ w 4 5", "x4 */ w 9 10"),
                                 ("3cd4 hc t", "cd2 wd10 vc", "cd4 hc b", "0 x4 vc"), "wd5 t x3 b"),
    "flowChartOffpageConnector": ({}, ("y1 */ h 4 5",), _RECT, "l t r y1"),
    "flowChartPredefinedProcess": ({}, ("x2 */ w 7 8",), _RECT, "wd8 t x2 b"),
    "flowChartPreparation": ({}, ("x2 */ w 4 5",), _RECT, "wd5 t x2 b"),
    "flowChartProcess": ({}, (), _RECT, "l t r b"),
    "flowChartTerminator": ({}, ("il */ w 1018 21600", "ir */ w 20582 21600", "it */ h 3163 21600",
                                 "ib */ h 18437 21600"), _RECT, "il it ir ib"),
    "foldedCorner": ({"adj": 16667}, ("a pin 0 adj 50000", "dy2 */ ss a 100000", "y2 +- b 0 dy2"),
                     _RECT, "l t r y2"),
    "hexagon": ({"adj": 25000, "vf": 115470},
                ("maxAdj */ 50000 w ss", "a pin 0 adj maxAdj", "shd2 */ hd2 vf 100000",
                 "x1 */ ss a 100000", "x2 +- r 0 x1", "dy1 sin shd2 3600000", "y1 +- vc 0 dy1",
                 "y2 +- vc dy1 0", "q1 */ maxAdj -1 2", "q2 +- a q1 0", "q3 ?: q2 4 2",
                 "q4 ?: q2 3 2", "q5 ?: q2 q1 0", "q6 +/ a q5 q1", "q7 */ q6 q4 -1",
                 "q8 +- q3 q7 0", "il */ w q8 24", "it */ h q8 24", "ir +- r 0 il", "ib +- b 0 it"),
                ("0 r vc", "cd4 x2 y2", "cd4 x1 y2", "cd2 l vc", "3cd4 x1 y1", "3cd4 x2 y1"),
                "il it ir ib"),
    "homePlate": ({"adj": 50000}, ("maxAdj */ 100000 w ss", "a pin 0 adj maxAdj",
                                   "dx1 */ ss a 100000", "x1 +- r 0 dx1", "ir +/ x1 r 2",
                                   "x2 */ x1 1 2"),
                  ("3cd4 x2 t", "cd2 l vc", "cd4 x2 b", "0 r vc"), "l t ir b"),
    "octagon": ({"adj": 29289}, ("a pin 0 adj 50000", "x1 */ ss a 100000", "x2 +- r 0 x1",
                                 "y2 +- b 0 x1", "il */ x1 1 2", "ir +- r 0 il", "ib +- b 0 il"),
                ("0 r x1", "0 r y2", "cd4 x2 b", "cd4 x1 b", "cd2 l y2", "cd2 l x1", "3cd4 x1 t",
                 "3cd4 x2 t"), "il il ir ib"),
    "parallelogram": ({"adj": 25000},
                      ("maxAdj */ 100000 w ss", "a pin 0 adj maxAdj", "x1 */ ss a 200000",
                       "x2 */ ss a 100000", "x6 +- r 0 x1", "x5 +- r 0 x2", "x3 */ x5 1 2",
                       "x4 +- r 0 x3", "q1 */ 5 a maxAdj", "q2 +/ 1 q1 12", "il */ q2 w 1",
                       "it */ q2 h 1", "ir +- r 0 il", "ib +- b 0 it", "q3 */ h hc x2",
                       "y1 pin 0 q3 h", "y2 +- b 0 y1"),
                      ("3cd4 hc y2", "3cd4 x4 t", "0 x6 vc", "cd4 x3 b", "cd4 hc y1", "cd2 x1 vc"),
                      "il it ir ib"),
    "plaque": ({"adj": 16667}, ("a pin 0 adj 50000", "x1 */ ss a 100000", "il */ x1 70711 100000",
                                "ir +- r 0 il", "ib +- b 0 il"), _RECT, "il il ir ib"),
    "rect": ({}, (), _RECT, "l t r b"),
    "round2SameRect": ({"adj1": 16667, "adj2": 0},
                       ("a1 pin 0 adj1 50000", "a2 pin 0 adj2 50000", "tx1 */ ss a1 100000",
                        "bx1 */ ss a2 100000", "d +- tx1 0 bx1", "tdx */ tx1 29289 100000",
                        "bdx */ bx1 29289 100000", "il ?: d tdx bdx", "ir +- r 0 il",
                        "ib +- b 0 bdx"), _RECT_RBLT, "il tdx ir ib"),
    "roundRect": ({"adj": 16667}, ("a pin 0 adj 50000", "x1 */ ss a 100000",
                                   "il */ x1 29289 100000", "ir +- r 0 il", "ib +- b 0 il"),
                  _RECT, "il il ir ib"),
    "snip1Rect": ({"adj": 16667}, ("a pin 0 adj 50000", "dx1 */ ss a 100000", "x1 +- r 0 dx1",
                                   "it */ dx1 1 2", "ir +/ x1 r 2"), _RECT_RBLT, "l it ir b"),
    "trapezoid": ({"adj": 25000}, ("maxAdj */ 50000 w ss", "a pin 0 adj maxAdj",
                                   "x1 */ ss a 200000", "x4 +- r 0 x1", "il */ wd3 a maxAdj",
                                   "it */ hd3 a maxAdj", "ir +- r 0 il"),
                  ("3cd4 hc t", "cd2 x1 vc", "cd4 hc b", "0 x4 vc"), "il it ir b"),
    "triangle": ({"adj": 50000}, ("a pin 0 adj 100000", "x1 */ w a 200000", "x2 */ w a 100000",
                                  "x3 +- x1 wd2 0"),
                 ("3cd4 x2 t", "cd2 x1 vc", "cd4 l b", "cd4 x2 b", "cd4 r b", "0 x3 vc"),
                 "x1 vc x3 b"),
}
SIDES = {"r": 0, "b": 90, "l": 180, "t": 270}  # outward direction of each side, in degrees


def _builtins(w, h):
    """The DrawingML built-in guides of a w x h shape."""
    ss = min(w, h)
    env = {"b": h, "h": h, "hc": w / 2, "l": 0.0, "ls": max(w, h), "r": w, "ss": ss, "t": 0.0,
           "vc": h / 2, "w": w, **_ANGLES}
    for n in (2, 3, 4, 5, 6, 8, 10, 12, 16, 32):
        env.update({f"hd{n}": h / n, f"ssd{n}": ss / n, f"wd{n}": w / n})
    return env


def _value(token, env):
    return env[token] if token in env else float(token)


def _guides(w, h, adj, guides):
    """Evaluate a geometry's guides for a w x h shape with adjust values `adj`."""
    env = _builtins(w, h)
    env.update(adj)
    for guide in guides:
        name, op, *args = guide.split()
        env[name] = _OPS[op](*(_value(a, env) for a in args))
    return env


def _adjusts(el):
    """{name: value} of the a:avLst under a prstGeom or custGeom element."""
    av = el.find(qn("a:avLst"))
    return {} if av is None else {gd.get("name"): float(gd.get("fmla").split()[-1])
                                  for gd in av.iterfind(qn("a:gd"))}


def geometry(sp):
    """(adjust values, guides, connection sites, text rectangle) of a shape's preset or custom
    geometry, in the PRESETS notation."""
    sppr = sp._element.spPr
    prst = sppr.find(qn("a:prstGeom"))
    if prst is not None:
        name = prst.get("prst")
        if name not in PRESETS:
            raise ValueError(f"shape {sp.name!r}: preset {name!r} has no connection-site table; "
                             f"use one of {sorted(PRESETS)} or a custom geometry")
        adj, guides, sites, text = PRESETS[name]
        return {**adj, **_adjusts(prst)}, guides, sites, text
    cust = sppr.find(qn("a:custGeom"))
    if cust is None:
        raise ValueError(f"shape {sp.name!r} has no geometry")
    guides = [f"{gd.get('name')} {gd.get('fmla')}"
              for gd in cust.iterfind(f"{qn('a:gdLst')}/{qn('a:gd')}")]
    rect = cust.find(qn("a:rect"))
    sites = [f"{c.get('ang')} {c[0].get('x')} {c[0].get('y')}"
             for c in cust.iterfind(f"{qn('a:cxnLst')}/{qn('a:cxn')}")]
    text = "l t r b" if rect is None else " ".join(rect.get(k) for k in ("l", "t", "r", "b"))
    return _adjusts(cust), guides, sites, text


def text_rect(prst, w, h, adj=None):
    """(l, t, r, b) of a preset's text rectangle in a w x h shape (same units as w and h)."""
    defaults, guides, _, text = PRESETS[prst]
    env = _guides(float(w), float(h), {**defaults, **(adj or {})}, guides)
    return tuple(_value(k, env) for k in text.split())


def size_for_text(prst, need_w, need_h, adj=None, min_w=0.0, min_h=0.0):
    """The smallest (w, h) of a preset whose text rectangle is at least need_w x need_h."""
    h = max(need_h, min_h, 1e-6)
    w = max(need_w, min_w, 1e-6)
    for _ in range(60):
        left, top, right, bottom = text_rect(prst, w, h, adj)
        if right - left >= need_w - 1e-9 and bottom - top >= need_h - 1e-9:
            return w, h
        if right - left < need_w - 1e-9:
            w *= min(need_w / max(right - left, 1e-6), 1.5) * 1.0001 + 1e-6
        if bottom - top < need_h - 1e-9:
            h *= min(need_h / max(bottom - top, 1e-6), 1.5) * 1.0001 + 1e-6
    raise ValueError(f"no {prst} shape holds {need_w:.2f} x {need_h:.2f}")


# --------------------------------------------------------------------------------------------
# frames: shape-local coordinates <-> slide coordinates (EMU)
# --------------------------------------------------------------------------------------------
def _xfrm(el):
    for tag in ("p:spPr", "p:grpSpPr"):
        pr = el.find(qn(tag))
        if pr is not None:
            return pr.find(qn("a:xfrm"))
    return el.find(qn("p:xfrm"))  # graphic frames (tables, charts)


def _flag(xfrm, name):
    return xfrm.get(name) in ("1", "true")


def _groups(node):
    """Maps (off, ext, chOff, chExt) of `node` and its enclosing groups, innermost first; empty for
    the slide's shape tree."""
    out = []
    while node is not None and etree.QName(node).localname == "grpSp":
        xfrm = _xfrm(node)
        if xfrm is not None:
            if xfrm.get("rot", "0") != "0" or _flag(xfrm, "flipH") or _flag(xfrm, "flipV"):
                raise ValueError("rotated or flipped groups are not supported")
            out.append([float(xfrm.find(qn(f"a:{tag}")).get(k)) for tag, k in
                        (("off", "x"), ("off", "y"), ("ext", "cx"), ("ext", "cy"), ("chOff", "x"),
                         ("chOff", "y"), ("chExt", "cx"), ("chExt", "cy"))])
        node = node.getparent()
    return out


def _local_to_slide(el, x, y, angle=0.0):
    """Map a point and a direction local to shape element `el` onto the slide (EMU, degrees),
    honouring the shape's flips and rotation and its enclosing groups."""
    xfrm = _xfrm(el)
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    rot = float(xfrm.get("rot", 0)) / 60000
    w, h = float(ext.get("cx")), float(ext.get("cy"))
    if _flag(xfrm, "flipH"):
        x, angle = w - x, 180 - angle
    if _flag(xfrm, "flipV"):
        y, angle = h - y, -angle
    if rot:
        th = math.radians(rot)
        dx, dy = x - w / 2, y - h / 2
        x = w / 2 + dx * math.cos(th) - dy * math.sin(th)
        y = h / 2 + dx * math.sin(th) + dy * math.cos(th)
        angle += rot
    x, y = x + float(off.get("x")), y + float(off.get("y"))
    for ox, oy, cx, cy, chx, chy, chcx, chcy in _groups(el.getparent()):
        x = ox + (x - chx) * (cx / chcx if chcx else 1.0)
        y = oy + (y - chy) * (cy / chcy if chcy else 1.0)
    return x, y, angle % 360


def _slide_to_container(container, x, y):
    """Map a slide point (EMU) into the child coordinates of `container` (spTree or grpSp)."""
    for ox, oy, cx, cy, chx, chy, chcx, chcy in reversed(_groups(container)):
        x = chx + (x - ox) * (chcx / cx if cx else 1.0)
        y = chy + (y - oy) * (chcy / cy if cy else 1.0)
    return x, y


@dataclass(frozen=True)
class Site:
    """A connection site on the slide: position in inches, outward direction in degrees
    (0 right, 90 down, 180 left, 270 up)."""
    x: float
    y: float
    angle: float

    @property
    def side(self):
        return min(SIDES, key=lambda s: abs((self.angle - SIDES[s] + 180) % 360 - 180))


def connection_sites(sp):
    """The connection sites of shape `sp`, by DrawingML site index (what a:stCxn idx refers to)."""
    out = []
    adj, guides, sites, _ = geometry(sp)
    env = _guides(float(sp.width), float(sp.height), adj, guides)
    for site in sites:
        ang, x, y = site.split()
        px, py, a = _local_to_slide(sp._element, _value(x, env), _value(y, env),
                                    _value(ang, env) / 60000)
        out.append(Site(px / EMU_IN, py / EMU_IN, a))
    return out


def bbox(sp):
    """(left, top, right, bottom) of a shape on the slide, in inches."""
    w, h = float(sp.width), float(sp.height)
    pts = [_local_to_slide(sp._element, x, y) for x, y in ((0, 0), (w, 0), (0, h), (w, h))]
    return (min(p[0] for p in pts) / EMU_IN, min(p[1] for p in pts) / EMU_IN,
            max(p[0] for p in pts) / EMU_IN, max(p[1] for p in pts) / EMU_IN)


def side_sites(sites, side):
    """Indexes of the sites facing `side` ('t', 'r', 'b', 'l'), in reading order along it."""
    idx = [i for i, st in enumerate(sites) if st.side == side]
    return sorted(idx, key=(lambda i: sites[i].x) if side in "tb" else (lambda i: sites[i].y))


def connector_ends(cn):
    """Begin and end points of a connector on the slide, in inches."""
    w, h = float(cn.width), float(cn.height)
    (bx, by, _), (ex, ey, _) = (_local_to_slide(cn._element, x, y) for x, y in ((0, 0), (w, h)))
    return (bx / EMU_IN, by / EMU_IN), (ex / EMU_IN, ey / EMU_IN)


def glue_targets(cn):
    """{'begin'|'end': (shape id, site index)} of a connector's a:stCxn / a:endCxn."""
    out = {}
    nv = cn._element.find(f"{qn('p:nvCxnSpPr')}/{qn('p:cNvCxnSpPr')}")
    for tag, key in (("a:stCxn", "begin"), ("a:endCxn", "end")):
        el = None if nv is None else nv.find(qn(tag))
        if el is not None:
            out[key] = (int(el.get("id")), int(el.get("idx")))
    return out


# --------------------------------------------------------------------------------------------
# orthogonal routes and connector geometry
# --------------------------------------------------------------------------------------------
_DASH = {"dash_dot": "dashDot", "dashed": "dash", "dotted": "sysDot", "solid": None}
_DIRS = {"b": (0, 1), "l": (-1, 0), "r": (1, 0), "t": (0, -1)}


def _same(p, q, eps=1e-6):
    return abs(p[0] - q[0]) < eps and abs(p[1] - q[1]) < eps


def simplify(pts, eps=1e-6):
    """Drop repeated points and merge consecutive segments that lie on one line."""
    out = []
    for p in pts:
        p = (float(p[0]), float(p[1]))
        if out and _same(p, out[-1], eps):
            continue
        if len(out) >= 2:
            (ax, ay), (bx, by) = out[-2], out[-1]
            vertical = abs(ax - bx) < eps and abs(bx - p[0]) < eps
            if vertical or (abs(ay - by) < eps and abs(by - p[1]) < eps):
                out.pop()
                if _same(p, out[-1], eps):
                    continue
        out.append(p)
    return out


def _direction(p, q):
    """Unit axis vector from p to q, or None when the segment is not horizontal or vertical."""
    dx, dy = q[0] - p[0], q[1] - p[1]
    if abs(dy) < 1e-6 and abs(dx) >= 1e-6:
        return (1 if dx > 0 else -1, 0)
    if abs(dx) < 1e-6 and abs(dy) >= 1e-6:
        return (0, 1 if dy > 0 else -1)
    return None


def _hits(p, q, box, clear):
    """True when segment p-q enters the interior of `box` grown by `clear`."""
    left, top, right, bottom = box[0] - clear, box[1] - clear, box[2] + clear, box[3] + clear
    if abs(p[1] - q[1]) < 1e-9:
        return top < p[1] < bottom and min(p[0], q[0]) < right and max(p[0], q[0]) > left
    return left < p[0] < right and min(p[1], q[1]) < bottom and max(p[1], q[1]) > top


def _score(pts, d0, d1, boxes, clear):
    """(segments, length, balance) of a candidate route, or None when it is not acceptable."""
    n = len(pts) - 1
    if n < 1 or n > MAX_SEGMENTS:
        return None
    dirs = [_direction(p, q) for p, q in zip(pts, pts[1:])]
    if None in dirs or dirs[0] != d0 or dirs[-1] != (-d1[0], -d1[1]):
        return None
    lengths = [abs(q[0] - p[0]) + abs(q[1] - p[1]) for p, q in zip(pts, pts[1:])]
    if n > 1 and (lengths[0] < MIN_FIRST - 1e-9 or lengths[-1] < MIN_LAST - 1e-9):
        return None
    for i, (p, q) in enumerate(zip(pts, pts[1:])):
        for box in boxes:
            own = (i == 0 and box == boxes[0]) or (i == n - 1 and box == boxes[1])
            if not own and _hits(p, q, box, clear):
                return None
    mid = ((pts[0][0] + pts[-1][0]) / 2, (pts[0][1] + pts[-1][1]) / 2)
    seg = n // 2
    centre = ((pts[seg][0] + pts[seg + 1][0]) / 2, (pts[seg][1] + pts[seg + 1][1]) / 2)
    return n, round(sum(lengths), 4), round(abs(centre[0] - mid[0]) + abs(centre[1] - mid[1]), 4)


def orthogonal_route(p0, d0, p1, d1, boxes, stub=0.15, clear=0.04):
    """The orthogonal route with the fewest segments (at most 5), then the shortest, from p0 leaving
    along d0 to p1 entering against d1 (d1 points out of the target). boxes[0] and boxes[1] are the
    source and target boxes, the rest are obstacles; units are inches. None when nothing fits."""
    if d0[0] == 0:
        def flip(p):
            return (p[1], p[0])
        res = orthogonal_route(flip(p0), flip(d0), flip(p1), flip(d1),
                               [(b[1], b[0], b[3], b[2]) for b in boxes], stub, clear)
        return None if res is None else [flip(p) for p in res]
    best = None
    cands = []
    (x0, y0), (x1, y1) = p0, p1
    a, b = boxes[0], boxes[1]
    xs = {x0 + d0[0] * stub, (x0 + x1) / 2, a[0] - stub, a[2] + stub, b[0] - stub, b[2] + stub}
    ys = {y0, y1, (y0 + y1) / 2, a[1] - stub, a[3] + stub, b[1] - stub, b[3] + stub}
    for lo, hi, coords in ((a[2], b[0], xs), (b[2], a[0], xs), (a[3], b[1], ys), (b[3], a[1], ys)):
        if lo < hi:
            coords.add((lo + hi) / 2)
    if d1[0]:
        xs.add(x1 + d1[0] * stub)
        cands.append([p0, p1])
        cands += [[p0, (x, y0), (x, y1), p1] for x in xs]
        cands += [[p0, (xa, y0), (xa, ym), (xb, ym), (xb, y1), p1]
                  for xa in xs for xb in xs for ym in ys]
    else:
        ys.add(y1 + d1[1] * stub)
        cands.append([p0, (x1, y0), p1])
        cands += [[p0, (xa, y0), (xa, ym), (x1, ym), p1] for xa in xs for ym in ys]
    for pts in cands:
        pts = simplify(pts)
        score = _score(pts, d0, d1, boxes, clear)
        if score is not None and (best is None or score < best[0]):
            best = (score, pts)
    return None if best is None else best[1]


def _rot(deg, x, y):
    """Rotate (x, y) clockwise on screen (y down) by 0, 90 or 270 degrees."""
    if deg == 90:
        return -y, x
    if deg == 270:
        return y, -x
    return x, y


def _frame(pts):
    """Connector frame for a polyline in EMU: (segments, adjust values, (off x, off y), (cx, cy),
    rotation, flipH, flipV). Bent connectors start horizontally in their own frame, so a route that
    starts vertically is drawn in a frame rotated by 90 or 270 degrees."""
    n = len(pts) - 1
    (x0, y0), (xn, yn) = pts[0], pts[-1]
    dx, dy = xn - x0, yn - y0
    if n == 1:
        return 1, [], (min(x0, xn), min(y0, yn)), (abs(dx), abs(dy)), 0, dx < 0, dy < 0
    if abs(pts[1][1] - y0) < abs(pts[1][0] - x0):
        du, dv = abs(dx), abs(dy)
        rot, fh, fv = 0, -1 if dx < 0 else 1, -1 if dy < 0 else 1
    else:
        down = dy > 0 or (dy == 0 and pts[1][1] > y0)
        du, dv = abs(dy), abs(dx)
        rot, fh, fv = 90 if down else 270, 1, -1 if (dx > 0 if down else dx < 0) else 1

    def local(p):
        u, v = _rot((360 - rot) % 360, p[0] - x0, p[1] - y0)
        return u * fh, v * fv

    u1 = local(pts[1])[0] if n >= 3 else 0.0
    v2 = local(pts[2])[1] if n >= 4 else 0.0
    u3 = local(pts[3])[0] if n == 5 else 0.0
    w = max(du, EXT_MIN, max(abs(u1), abs(u3)) * 100000 / ADJ_MAX)
    h = max(dv, EXT_MIN, abs(v2) * 100000 / ADJ_MAX)
    adjs = [round(u1 / w * 100000), round(v2 / h * 100000), round(u3 / w * 100000)][:n - 2]
    cx, cy = _rot(rot, fh * w / 2, fv * h / 2)
    return n, adjs, (x0 + cx - w / 2, y0 + cy - h / 2), (w, h), rot, fh < 0, fv < 0


def _shape_connector(cn, pts, curved=False):
    """Give connector `cn` the geometry of polyline `pts` (container EMU)."""
    n, adjs, (ox, oy), (w, h), rot, flip_h, flip_v = _frame(pts)
    sppr = cn._element.spPr
    geom = sppr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    xfrm = sppr.find(qn("a:xfrm"))
    geom.set("prst", "straightConnector1" if n == 1 else
             f"{'curved' if curved else 'bent'}Connector{n}")
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for gd in list(av):
        av.remove(gd)
    for i, val in enumerate(adjs, 1):
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", f"adj{i}")
        gd.set("fmla", f"val {val}")
    for key in ("rot", "flipH", "flipV"):
        xfrm.attrib.pop(key, None)
    if rot:
        xfrm.set("rot", str(rot * 60000))
    if flip_h:
        xfrm.set("flipH", "1")
    if flip_v:
        xfrm.set("flipV", "1")
    xfrm.find(qn("a:off")).set("x", str(round(ox)))
    xfrm.find(qn("a:off")).set("y", str(round(oy)))
    xfrm.find(qn("a:ext")).set("cx", str(round(w)))
    xfrm.find(qn("a:ext")).set("cy", str(round(h)))


def _style_line(shape, color, width, style="solid", head=None, tail=None, size="med"):
    """Solid colour, width (pt), dash style and line ends of a line or connector."""
    if style not in _DASH:
        raise ValueError(f"unknown line style {style!r}; expected one of {sorted(_DASH)}")
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    ln = shape.line._get_or_add_ln()
    if _DASH[style]:
        etree.SubElement(ln, qn("a:prstDash")).set("val", _DASH[style])
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind and kind != "none":
            el = etree.SubElement(ln, qn(tag))
            el.set("type", kind)
            el.set("w", size)
            el.set("len", size)


# --------------------------------------------------------------------------------------------
# custom-site geometry: a rect, roundRect or hexagon with evenly spaced sites on every side
# --------------------------------------------------------------------------------------------
_PORT_SHAPES = {
    "hexagon": (("fw +- x2 0 x1",), (("M", "l", "vc"), ("L", "x1", "y1"), ("L", "x2", "y1"),
                                     ("L", "r", "vc"), ("L", "x2", "y2"), ("L", "x1", "y2"))),
    "rect": ((), (("M", "l", "t"), ("L", "r", "t"), ("L", "r", "b"), ("L", "l", "b"))),
    "roundRect": (("x2 +- r 0 x1", "y2 +- b 0 x1"),
                  (("M", "x1", "t"), ("L", "x2", "t"), ("A", 16200000), ("L", "r", "y2"), ("A", 0),
                   ("L", "x1", "b"), ("A", 5400000), ("L", "l", "x1"), ("A", 10800000))),
}


def _port_sites(prst, counts):
    """(guides, sites) for `counts` = (top, right, bottom, left) evenly spaced sites."""
    guides = []
    sites = []
    for side, n in zip("trbl", counts):
        for k in range(1, n + 1):
            i = len(sites)
            if prst == "hexagon" and side in "tb":
                guides += [f"p{i}a */ fw {k} {n + 1}", f"p{i}x +- x1 p{i}a 0"]
                sites.append(f"3cd4 p{i}x y1" if side == "t" else f"cd4 p{i}x y2")
            elif prst == "hexagon":
                guides += [f"p{i}y */ h {k} {n + 1}", f"p{i}a */ x1 {abs(2 * k - n - 1)} {n + 1}",
                           f"p{i}x +- r 0 p{i}a" if side == "r" else f"p{i}x +- l p{i}a 0"]
                sites.append(f"{'0' if side == 'r' else 'cd2'} p{i}x p{i}y")
            elif side in "tb":
                guides.append(f"p{i}x */ w {k} {n + 1}")
                sites.append(f"{'3cd4' if side == 't' else 'cd4'} p{i}x {side}")
            else:
                guides.append(f"p{i}y */ h {k} {n + 1}")
                sites.append(f"{'0' if side == 'r' else 'cd2'} {side} p{i}y")
    return guides, sites


def _set_ports(sp, counts):
    """Swap the preset geometry of `sp` for an equivalent custGeom with per-side sites."""
    body = []
    sppr = sp._element.spPr
    prst_el = sppr.find(qn("a:prstGeom"))
    prst = None if prst_el is None else prst_el.get("prst")
    if prst not in _PORT_SHAPES:
        raise ValueError(f"ports() supports {sorted(_PORT_SHAPES)} shapes, not {prst!r}")
    defaults, guides, _, text = PRESETS[prst]
    adj = {**defaults, **_adjusts(prst_el)}
    extra, path = _PORT_SHAPES[prst]
    site_guides, sites = _port_sites(prst, counts)
    for cmd in path:
        if cmd[0] == "M":
            body.append(f'<a:moveTo><a:pt x="{cmd[1]}" y="{cmd[2]}"/></a:moveTo>')
        elif cmd[0] == "L":
            body.append(f'<a:lnTo><a:pt x="{cmd[1]}" y="{cmd[2]}"/></a:lnTo>')
        else:
            body.append(f'<a:arcTo wR="x1" hR="x1" stAng="{cmd[1]}" swAng="5400000"/>')
    av = "".join(f'<a:gd name="{k}" fmla="val {round(v)}"/>' for k, v in adj.items())
    gd = "".join(f'<a:gd name="{g.split()[0]}" fmla="{g.split(" ", 1)[1]}"/>'
                 for g in (*guides, *extra, *site_guides))
    cxn = "".join(f'<a:cxn ang="{a}"><a:pos x="{x}" y="{y}"/></a:cxn>'
                  for a, x, y in (s.split() for s in sites))
    rect = dict(zip("ltrb", text.split()))
    cust = etree.fromstring(
        f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:avLst>{av}</a:avLst><a:gdLst>{gd}</a:gdLst><a:ahLst/><a:cxnLst>{cxn}</a:cxnLst>'
        f'<a:rect l="{rect["l"]}" t="{rect["t"]}" r="{rect["r"]}" b="{rect["b"]}"/>'
        f'<a:pathLst><a:path>{"".join(body)}<a:close/></a:path></a:pathLst></a:custGeom>')
    prst_el.addprevious(cust)
    sppr.remove(prst_el)
    return sp


# --------------------------------------------------------------------------------------------
# text measuring in the deck font
# --------------------------------------------------------------------------------------------
_FONT_FILES = {}
_PIL_FONTS = {}


def _font_files():
    """{face name (lower case): (font file, index in a collection)} from the Windows registry."""
    if _FONT_FILES:
        return _FONT_FILES
    try:
        import winreg
    except ImportError:
        return _FONT_FILES
    fonts_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            key = winreg.OpenKey(root, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")
        except OSError:
            continue
        with key:
            for i in range(winreg.QueryInfoKey(key)[1]):
                name, value, _ = winreg.EnumValue(key, i)
                path = value if os.path.isabs(value) else os.path.join(fonts_dir, value)
                for index, face in enumerate(name.rsplit(" (", 1)[0].split(" & ")):
                    _FONT_FILES.setdefault(face.strip().lower(), (path, index))
    return _FONT_FILES


def _pil_font(family, bold=False, italic=False):
    """The Pillow font for a face. A family without a bold face (Calibri Light) is measured with
    its regular face: PowerPoint's simulated bold keeps the regular advance widths (measured)."""
    key = (family.lower(), bold, italic)
    if key not in _PIL_FONTS:
        files = _font_files()
        style = " ".join(word for word, on in (("bold", bold), ("italic", italic)) if on)
        names = [f"{family} {style}".strip(), f"{family} italic" if italic else "", family,
                 f"{family} regular"]
        found = next((files[n.lower()] for n in names if n and n.lower() in files), None)
        if found is None:
            raise RuntimeError(f"cannot measure text: font {family!r} is not installed (not in the "
                               "Windows font registry); choose an installed theme font")
        _PIL_FONTS[key] = ImageFont.truetype(found[0], MEASURE_PX, index=found[1],
                                             layout_engine=ImageFont.Layout.BASIC)
    return _PIL_FONTS[key]


def text_width(text, size, family, bold=False, italic=False):
    """Points PowerPoint needs to keep `text` on one line: the Pillow width plus the calibrated
    safety margin (TEXT_MARGIN, TEXT_PAD_PT)."""
    if not text:
        return 0.0
    return natural_width(text, size, family, bold, italic) * TEXT_MARGIN + TEXT_PAD_PT


def natural_width(text, size, family, bold=False, italic=False):
    """Pillow's width of `text` in points, without the safety margin."""
    return _pil_font(family, bold, italic).getlength(text) * size / MEASURE_PX


def check_size(size):
    if size < MIN_PT:
        raise ValueError(f"text at {size} pt is below the {MIN_PT} pt minimum; split the diagram "
                         "instead of shrinking its text")


def alt_text(shape, text, title=None):
    """Set a shape's alternative text (read by screen readers); returns the shape."""
    cnv = shape._element.find(f"*/{qn('p:cNvPr')}")
    cnv.set("descr", text)
    if title:
        cnv.set("title", title)
    return shape


def _short(name):
    return name.split(": ", 1)[-1]


def _facing(a, b):
    """Facing sides of boxes a and b, across the wider gap between them."""
    gap_x = max(b[0] - a[2], a[0] - b[2])
    gap_y = max(b[1] - a[3], a[1] - b[3])
    if gap_y >= gap_x:
        return ("b", "t") if b[1] + b[3] >= a[1] + a[3] else ("t", "b")
    return ("r", "l") if b[0] + b[2] >= a[0] + a[2] else ("l", "r")


def _toward(box, p):
    """The side of `box` that faces point p."""
    dx = p[0] - box[2] if p[0] > box[2] else (p[0] - box[0] if p[0] < box[0] else 0.0)
    dy = p[1] - box[3] if p[1] > box[3] else (p[1] - box[1] if p[1] < box[1] else 0.0)
    if abs(dx) > abs(dy):
        return "r" if dx > 0 else "l"
    return "b" if dy >= 0 else "t"


def _pick(sites, box, spec, toward):
    """Site index for `spec`: an index, a (side, k) pair or a side (its site nearest the middle)."""
    if isinstance(spec, int):
        if not 0 <= spec < len(sites):
            raise ValueError(f"site {spec} out of range: the shape has {len(sites)} sites")
        return spec
    if isinstance(spec, tuple):
        on_side = side_sites(sites, spec[0])
        if not on_side:
            raise ValueError(f"the shape has no site on side {spec[0]!r}")
        return on_side[spec[1]]
    if spec not in SIDES:
        raise ValueError(f"unknown site {spec!r}; use 't', 'r', 'b', 'l', (side, k) or an index")
    mid = {"b": ((box[0] + box[2]) / 2, box[3]), "l": (box[0], (box[1] + box[3]) / 2),
           "r": (box[2], (box[1] + box[3]) / 2), "t": ((box[0] + box[2]) / 2, box[1])}[spec]
    cands = side_sites(sites, spec) or list(range(len(sites)))
    return min(cands, key=lambda i: (round(math.dist((sites[i].x, sites[i].y), mid), 4),
                                     math.dist((sites[i].x, sites[i].y), toward)))


def _snap(pts, d0, d1):
    """Align the first and last bends of a caller's route with the two site points."""
    pts = [tuple(p) for p in pts]
    if len(pts) > 2:
        pts[1] = (pts[1][0], pts[0][1]) if d0[1] == 0 else (pts[0][0], pts[1][1])
        pts[-2] = (pts[-2][0], pts[-1][1]) if d1[1] == 0 else (pts[-1][0], pts[-2][1])
    return simplify(pts)


@dataclass
class Link:
    """What Deck.link() made: the connector, its label chip (or None) and its route in inches."""
    connector: object
    label: object
    points: list


class Shapes:
    """Diagram primitives, mixed into Deck (they use its theme, fonts and text helpers)."""

    # -- text measuring -----------------------------------------------------------------------
    def measure(self, text, size=13, bold=False, italic=False, mono=False, font=None):
        """Inches PowerPoint needs to keep each line of `text` on one line (the widest line), in
        the deck font; a safe estimate, calibrated against PowerPoint."""
        family = font or (self.MONO if mono else self.FONT)
        return max((text_width(line, size, family, bold, italic) for line in text.split("\n")),
                   default=0.0) / 72

    def wrap(self, text, max_w, size=13, bold=False, italic=False, mono=False, font=None):
        """Lines of `text` greedily filled up to `max_w` inches. Newlines are kept; a word wider
        than max_w stays whole on its own line (words are never broken)."""
        lines = []
        for para in text.split("\n"):
            line = ""
            for word in para.split():
                trial = f"{line} {word}" if line else word
                if line and self.measure(trial, size, bold, italic, mono, font) > max_w:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        return lines

    def text_size(self, text, size=13, bold=False, max_w=None, italic=False, mono=False, font=None):
        """(w, h, lines) of a text block in inches, wrapped to `max_w` when given."""
        lines = (self.wrap(text, max_w, size, bold, italic, mono, font) if max_w
                 else text.split("\n"))
        w = max((self.measure(line, size, bold, italic, mono, font) for line in lines), default=0.0)
        return w, len(lines) * size * LINE_SPACING / 72, lines

    def fit_shape(self, text, prst="roundRect", size=13, bold=False, max_w=2.4, pad=(0.08, 0.04),
                  min_w=0.0, min_h=0.0, adj=None):
        """(w, h, lines): the smallest `prst` shape whose text area holds `text` at `size` pt with
        `pad` insets, wrapped to at most `max_w` inches of text. Text never goes below 11 pt."""
        check_size(size)
        tw, th, lines = self.text_size(text, size, bold, max_w)
        w, h = size_for_text(prst, tw + 2 * pad[0], th + 2 * pad[1], adj, min_w, min_h)
        return math.ceil(w * 100 - 1e-6) / 100, math.ceil(h * 100 - 1e-6) / 100, lines

    # -- shapes -------------------------------------------------------------------------------
    def _prst_shape(self, s, prst, x, y, w, h):
        sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(round(x * EMU_IN)), Emu(round(y * EMU_IN)),
                                Emu(round(w * EMU_IN)), Emu(round(h * EMU_IN)))
        sp._element.spPr.find(qn("a:prstGeom")).set("prst", prst)
        return sp

    def _write(self, sp, text, size, color, bold=False, align="center", anchor="middle",
               margins=(0.08, 0.04), wrap=True):
        tf = sp.text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = {"bottom": MSO_ANCHOR.BOTTOM, "middle": MSO_ANCHOR.MIDDLE,
                              "top": MSO_ANCHOR.TOP}[anchor]
        tf.margin_left = tf.margin_right = Emu(round(margins[0] * EMU_IN))
        tf.margin_top = tf.margin_bottom = Emu(round(margins[1] * EMU_IN))
        for i, para in enumerate(text.split("\n")):
            self.run(self.para(tf, first=i == 0, align=align), para, size=size, color=color,
                     bold=bold)

    def styled_node(self, s, x, y, w, h, text, kind="process", size=None, prst=None, fill=None,
                    line=None, text_color=None, bold=False, align="center", anchor="middle",
                    name=None):
        """A node in the modern style of `kind` (fill/outline pair from the palette, thin outline,
        no shadow, rounded corners of a constant radius). Returns the shape."""
        size = size or TYPE_SCALE["node"]
        check_size(size)
        kprst, kfill, kline, ktext, dashed = kind_style(kind, self.palette_hex, self.t)
        sp = self._prst_shape(s, prst or kprst, x, y, w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill or _rgb(kfill)
        _style_line(sp, line or _rgb(kline), NODE_LINE_WIDTH, "dashed" if dashed else "solid")
        sp.shadow.inherit = False
        if (prst or kprst) == "roundRect":
            sp.adjustments[0] = min(0.5, NODE_RADIUS / max(min(w, h), 1e-6))
        if text:
            self._write(sp, text, size, text_color or _rgb(ktext), bold, align, anchor)
        sp.name = name or f"Node: {text.splitlines()[0] if text else kind}"
        return sp

    def _chip_size(self, text, size, bold=False):
        tw, th, _ = self.text_size(text, size, bold)
        return size_for_text("roundRect", tw + 2 * CHIP_PAD[0], th + 2 * CHIP_PAD[1],
                             {"adj": 50000})

    def label(self, s, cx, cy, text, size=None, color=None, fill=None, bold=False, name=None):
        """A text chip centred on (cx, cy) inches with the slide-background fill and no outline, so
        a line behind it never cuts the text. Returns the shape."""
        size = size or TYPE_SCALE["label"]
        check_size(size)
        w, h = self._chip_size(text, size, bold)
        sp = self._prst_shape(s, "roundRect", cx - w / 2, cy - h / 2, w, h)
        sp.adjustments[0] = 0.5
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill or self.LIGHT
        sp.line.fill.background()
        sp.shadow.inherit = False
        self._write(sp, text, size, color or self.TEXT, bold, margins=CHIP_PAD, wrap=False)
        sp.name = name or f"Label: {text.splitlines()[0]}"
        return sp

    def ports(self, sp, top=1, right=1, bottom=1, left=1):
        """Give a rect, roundRect or hexagon shape `top`/`right`/`bottom`/`left` evenly spaced
        connection sites per side (a custom geometry that looks the same), so several connectors
        can meet one side at distinct points; pick them in link() with (side, k). Returns sp."""
        return _set_ports(sp, (top, right, bottom, left))

    def alt_text(self, shape, text, title=None):
        """Set a shape's alternative text; returns the shape."""
        return alt_text(shape, text, title)

    def group(self, s, shapes, name=None, alt=None):
        """Group `shapes` (nodes, connectors, labels) into one group shape with alt text."""
        grp = s.shapes.add_group_shape(list(shapes))
        if name:
            grp.name = name
        if alt:
            alt_text(grp, alt)
        return grp

    # -- links --------------------------------------------------------------------------------
    def link(self, s, src, dst, label="", *, src_site=None, dst_site=None, via=None,
             route="ortho", arrow="triangle", head=None, style="solid", color=None, width=None,
             avoid=(), name=None, label_at=None):
        """A connector glued from shape `src` to shape `dst`; returns Link(connector, label, route).

        src_site / dst_site: None (the facing sides), 't' 'r' 'b' 'l' (the side's middle site), a
        (side, k) pair (the k-th site on that side) or a site index. route: 'ortho' (right angles,
        at most 5 segments; `via` lists the bend points), 'curved' (the same route, smoothed) or
        'straight'. arrow / head: line end at dst / src ('triangle', 'stealth', 'arrow', 'diamond',
        'oval' or None). style: 'solid', 'dashed', 'dotted' or 'dash_dot'. `avoid`: more shapes the
        automatic route keeps out of. `s` is the slide or group shape to draw in.
        """
        if route not in ("curved", "ortho", "straight"):
            raise ValueError(f"unknown route {route!r}; expected 'ortho', 'curved' or 'straight'")
        chip = None
        box_a, box_b = bbox(src), bbox(dst)
        sites_a, sites_b = connection_sites(src), connection_sites(dst)
        if via:
            side_a, side_b = _toward(box_a, via[0]), _toward(box_b, via[-1])
        else:
            side_a, side_b = _facing(box_a, box_b)
        centre_a = ((box_a[0] + box_a[2]) / 2, (box_a[1] + box_a[3]) / 2)
        centre_b = ((box_b[0] + box_b[2]) / 2, (box_b[1] + box_b[3]) / 2)
        ia = _pick(sites_a, box_a, side_a if src_site is None else src_site, centre_b)
        ib = _pick(sites_b, box_b, side_b if dst_site is None else dst_site, centre_a)
        pa, pb = sites_a[ia], sites_b[ib]
        d0, d1 = _DIRS[pa.side], _DIRS[pb.side]
        if route == "straight":
            pts = [(pa.x, pa.y), (pb.x, pb.y)]
        elif via:
            pts = _snap([(pa.x, pa.y), *via, (pb.x, pb.y)], d0, d1)
            if any(_direction(p, q) is None for p, q in zip(pts, pts[1:])):
                raise ValueError(f"link {src.name!r} -> {dst.name!r}: via points must make a "
                                 "route of horizontal and vertical segments")
        else:
            pts = orthogonal_route((pa.x, pa.y), d0, (pb.x, pb.y), d1,
                                   [box_a, box_b, *(bbox(o) for o in avoid)])
            if pts is None:
                raise ValueError(f"link {src.name!r} -> {dst.name!r}: no route with at most "
                                 f"{MAX_SEGMENTS} segments; move the shapes apart, choose other "
                                 "sides or pass via points (or split the diagram)")
        if len(pts) - 1 > MAX_SEGMENTS:
            raise ValueError(f"link {src.name!r} -> {dst.name!r}: the route has {len(pts) - 1} "
                             f"segments but PowerPoint connectors allow at most {MAX_SEGMENTS}; "
                             "split the diagram or add layout hints")
        container = s.shapes._spTree
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 1, 1)
        _shape_connector(cn, [_slide_to_container(container, x * EMU_IN, y * EMU_IN)
                              for x, y in pts], curved=route == "curved")
        nv = cn._element.find(f"{qn('p:nvCxnSpPr')}/{qn('p:cNvCxnSpPr')}")
        for tag, shape, idx in (("a:stCxn", src, ia), ("a:endCxn", dst, ib)):
            el = etree.SubElement(nv, qn(tag))
            el.set("id", str(shape.shape_id))
            el.set("idx", str(idx))
        _style_line(cn, color or self.LINE, width or EDGE_WIDTH, style, head, arrow)
        cn.shadow.inherit = False
        cn.name = name or f"Edge: {_short(src.name)} -> {_short(dst.name)}"
        if label:
            if label_at is None:
                label_at = _label_spot(pts, self._chip_size(label, TYPE_SCALE["label"]),
                                       [box_a, box_b, *(bbox(o) for o in avoid)])
            chip = self.label(s, label_at[0], label_at[1], label)
        return Link(cn, chip, pts)


def _label_spot(pts, size, boxes):
    """Centre for a label chip of `size` (w, h) on a route: the point nearest the middle of the
    longest segment where the chip stays clear of `boxes`, else the point with the least overlap."""
    spots = [((p[0] + q[0]) / 2 + (q[0] - p[0]) * f, (p[1] + q[1]) / 2 + (q[1] - p[1]) * f)
             for p, q in sorted(zip(pts, pts[1:]), key=lambda pq: -math.dist(*pq))
             for f in (0.0, -0.15, 0.15, -0.3, 0.3, -0.4, 0.4)]

    def overlap(spot):
        chip = (spot[0] - size[0] / 2, spot[1] - size[1] / 2, spot[0] + size[0] / 2,
                spot[1] + size[1] / 2)
        return sum(_overlap_area(chip, box, 0.02) for box in boxes)

    return next((spot for spot in spots if not overlap(spot)), min(spots, key=overlap))


def _overlap_area(a, b, clear):
    w = min(a[2], b[2] + clear) - max(a[0], b[0] - clear)
    h = min(a[3], b[3] + clear) - max(a[1], b[1] - clear)
    return w * h if w > 0 and h > 0 else 0.0


def _rgb(hexval):
    return RGBColor.from_string(hexval)

