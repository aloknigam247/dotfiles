"""Tree, grid, set and time builders, mixed into Deck: org chart, mind map, tree view, block
diagram, user journey, event modeling, timeline, Gantt chart, quadrant chart, Venn and Sankey
diagrams.

Like the graph builders, each draws into a slide region (`box`, default: the content area under
the header), keeps text at 11 pt or more (it raises LayoutError, "split the diagram", when the data
does not fit) and groups what it draws with alt text. Links between shapes are connectors glued
at both ends; Sankey ribbons are filled shapes and are not glued.
"""
import datetime as dt
import math

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import layout
from .diagrams import _ANGLE, _custom_sites
from .layout import LayoutError
from .shapes import EMU_IN, LINE_SPACING, _rgb, alt_text
from .style import CONTENT, MIN_PT, NODE_RADIUS, TYPE_SCALE, blend, tint

_BRANCH_SLOTS = (0, 3, 2, 5, 1, 4)  # org chart and mind map branches: sky, emerald, amber...
_EVENT = {  # event-modeling convention: kind -> (fill, outline)
    "cmd": ("BFDBFE", "3B82F6"), "evt": ("FED7AA", "F97316"), "pcr": ("F5D0FE", "C026D3"),
    "rmo": ("D9F99D", "65A30D"), "ui": ("FFFFFF", "CBD5E1"),
}
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _font(size):
    if size < MIN_PT:
        raise ValueError(f"text at {size} pt is below the {MIN_PT} pt minimum")
    return size


def _split(text):
    """('Name', 'Role') of an org chart text 'Name — Role' (the role may be empty)."""
    for sep in (" \u2014 ", " - ", "\n"):
        if sep in text:
            name, role = text.split(sep, 1)
            return name.strip(), role.strip()
    return text.strip(), ""


def _geom(sp, paths, sites=()):
    """Give `sp` a custom geometry. paths: [(points, closed, filled)] with points as (fx, fy)
    fractions of the box; sites: [(side, fx, fy)] connection sites, site k = sites[k]."""
    body, cxn, gd, names = [], [], [], {}

    def point(fx, fy):
        key = (round(fx, 5), round(fy, 5))
        if key not in names:
            k = len(names)
            names[key] = (f"x{k}", f"y{k}")
            gd.append(f'<a:gd name="x{k}" fmla="*/ w {round(fx * 100000)} 100000"/>'
                      f'<a:gd name="y{k}" fmla="*/ h {round(fy * 100000)} 100000"/>')
        return names[key]

    for pts, closed, filled in paths:
        cmds = []
        for i, (fx, fy) in enumerate(pts):
            tag = "a:moveTo" if i == 0 else "a:lnTo"
            x, y = point(fx, fy)
            cmds.append(f'<{tag}><a:pt x="{x}" y="{y}"/></{tag}>')
        if closed:
            cmds.append("<a:close/>")
        attr = "" if filled else ' fill="none"'
        body.append(f'<a:path{attr}>{"".join(cmds)}</a:path>')
    for side, fx, fy in sites:
        x, y = point(fx, fy)
        cxn.append(f'<a:cxn ang="{_ANGLE[side]}"><a:pos x="{x}" y="{y}"/></a:cxn>')
    cust = etree.fromstring(
        f'<a:custGeom xmlns:a="{_NS_A}"><a:avLst/><a:gdLst>{"".join(gd)}</a:gdLst><a:ahLst/>'
        f'<a:cxnLst>{"".join(cxn)}</a:cxnLst><a:rect l="l" t="t" r="r" b="b"/>'
        f'<a:pathLst>{"".join(body)}</a:pathLst></a:custGeom>')
    old = sp._element.spPr.find(qn("a:prstGeom"))
    old.addprevious(cust)
    sp._element.spPr.remove(old)
    return sp


def _alpha(sp, share):
    """Make the solid fill of `sp` `share` opaque (0..1)."""
    clr = sp._element.spPr.find(qn("a:solidFill"))[0]
    etree.SubElement(clr, qn("a:alpha")).set("val", str(round(share * 100000)))


def _cell_lines(cell, color, width=9525):
    """Thin borders of colour `color` on every side of a table cell."""
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcpr.findall(qn(tag)):
            tcpr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.fromstring(
            f'<{tag} xmlns:a="{_NS_A}" w="{width}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:prstDash val="solid"/>'
            f'</{tag}>')
        tcpr.insert(i, ln)


def _fmt_day(d):
    return f"{d:%b} {d.day}"


def _num(v):
    return f"{v:g}" if isinstance(v, float) else str(v)


class Structures:
    """Tree, grid, set and time builders, mixed into Deck (they use its theme and links)."""

    # -- helpers --------------------------------------------------------------------------------
    def _shape(self, s, prst, x, y, w, h, fill=None, line=None, width=1.0, dash=None,
               alpha=None, name="Deco: shape"):
        """A preset shape with a solid fill (or none), an outline (or none) and no shadow."""
        sp = self._prst_shape(s, prst, x, y, w, h)
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = _rgb(fill)
            if alpha is not None:
                _alpha(sp, alpha)
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = _rgb(line)
            sp.line.width = Pt(width)
            if dash:
                etree.SubElement(sp.line._get_or_add_ln(), qn("a:prstDash")).set("val", dash)
        sp.shadow.inherit = False
        sp.name = name
        return sp

    def _card(self, s, x, y, w, h, lines, fill, line, *, prst="roundRect", name, width=1.0,
              radius=None, align="center", anchor="middle", margins=(0.08, 0.04)):
        """A shape holding text lines [(text, size, bold, colour hex or None)]."""
        sp = self._shape(s, prst, x, y, w, h, fill, line, width, name=name)
        if prst == "roundRect":
            sp.adjustments[0] = min(0.5, (radius or NODE_RADIUS) / max(min(w, h), 1e-6))
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"bottom": 4, "middle": 3, "top": 1}[anchor]
        tf.margin_left = tf.margin_right = Emu(round(margins[0] * EMU_IN))
        tf.margin_top = tf.margin_bottom = Emu(round(margins[1] * EMU_IN))
        for i, (text, size, bold, color) in enumerate(lines):
            self.run(self.para(tf, first=i == 0, align=align), text, size=_font(size), bold=bold,
                     color=_rgb(color) if color else self.TEXT)
        return sp

    def _text(self, s, x, y, w, h, text, size, *, color=None, bold=False, italic=False,
              align="left", anchor="middle", name):
        """A text box without insets holding `text` (one paragraph per line)."""
        sp, tf = self._text_shape(s, x, y, w, h, anchor, name)
        for i, line in enumerate(text.split("\n")):
            self.run(self.para(tf, first=i == 0, align=align), line, size=_font(size),
                     bold=bold, italic=italic, color=_rgb(color) if color else self.TEXT)
        return sp

    def _line(self, s, x1, y1, x2, y2, color, width=1.0, dash=None, name="Deco: line",
              arrow=None):
        """A plain (unglued) line, for axes, grids and dividers."""
        cn = self.connector(s, x1, y1, x2, y2, color=_rgb(color), width=width)
        ln = cn.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
        if arrow:
            end = etree.SubElement(ln, qn("a:tailEnd"))
            end.set("type", arrow)
            end.set("w", "med")
            end.set("len", "med")
        cn.name = name
        return cn

    def _place(self, size, box, what):
        """Offset that centres a drawing of `size` (w, h) in `box`; raises when it is larger."""
        if size[0] > box[2] + 1e-6 or size[1] > box[3] + 1e-6:
            raise LayoutError(f"{what} needs {size[0]:.1f} x {size[1]:.1f} in but its region is "
                              f"{box[2]:.1f} x {box[3]:.1f} in at 11 pt or more: split the "
                              "diagram or give it a larger box")
        return box[0] + (box[2] - size[0]) / 2, box[1] + (box[3] - size[1]) / 2

    def _finish(self, s, made, kind, title, alt):
        return self.group(s, made, name=f"Diagram: {kind} {title or ''}".rstrip(), alt=alt)

    # -- trees ----------------------------------------------------------------------------------
    def org_chart(self, s, root, children, *, box=None, direction="auto", font=None,
                  title=None, alt=None):
        """An org chart laid out as a tidy tree. root: 'Name — Role'; children: [(text,
        [children])] to any depth. ' — ' splits a person into a bold name and a role line; each
        first-level branch takes a palette colour; links are elbow connectors glued at both
        ends. direction: 'auto' (top-down, else left-to-right), 'TB' or 'LR'."""
        kids = {}
        made = []
        shapes = {}
        slot = {}
        texts = {}
        tries = [("TB", 0.3, 0.55), ("TB", 0.18, 0.45), ("TB", 0.1, 0.35), ("LR", 0.18, 0.55),
                 ("LR", 0.1, 0.4)]
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        role_size = max(MIN_PT, font - 2)
        if direction not in ("auto", "TB", "LR"):
            raise ValueError(f"org_chart() direction must be 'auto', 'TB' or 'LR', not "
                             f"{direction!r}")

        def walk(nid, text, subs, colour):
            texts[nid] = _split(text)
            slot[nid] = colour
            kids[nid] = []
            for k, (sub_text, sub) in enumerate(subs):
                cid = f"{nid}.{k}"
                kids[nid].append(cid)
                walk(cid, sub_text, sub, _BRANCH_SLOTS[k % 6] if nid == "0" else colour)

        walk("0", root, children, None)
        h = (font + (role_size if any(r for _, r in texts.values()) else 0)) * LINE_SPACING / 72 \
            + 0.24
        w = max(max(self.measure(n, font, bold=True), self.measure(r, role_size) if r else 0)
                for n, r in texts.values()) + 0.36
        sizes = {n: (max(w, 1.4), h) for n in texts}
        tries = [t for t in tries if direction in ("auto", t[0])]
        for d, sib, gap in tries:
            boxes, size = layout.tidy_tree("0", kids, sizes, direction=d, sib_gap=sib,
                                           level_gap=gap)
            if size[0] <= box[2] + 1e-6 and size[1] <= box[3] + 1e-6:
                break
        ox, oy = self._place(size, box, "org chart")
        for nid, (x, y, bw, bh) in boxes.items():
            accent = self._hex(slot[nid] if slot[nid] is not None else 5)
            name, role = texts[nid]
            lines = [(name, font, True, None)] + ([(role, role_size, False, self.t["line"])]
                                                   if role else [])
            shapes[nid] = self._card(s, ox + x, oy + y, bw, bh, lines,
                                     tint(accent, 0.2 if nid == "0" else 0.12), accent,
                                     name=f"Node: {name}")
            made.append(shapes[nid])
        for parent, subs in kids.items():
            a = shapes[parent]
            ax, ay, aw, ah = boxes[parent]
            for c in subs:
                b = shapes[c]
                bx, by, bw, bh = boxes[c]
                if d == "TB":
                    sites = ("b", "t")
                    mid = oy + (ay + ah + by) / 2
                    p, q = ox + ax + aw / 2, ox + bx + bw / 2
                    via = [(p, mid), (q, mid)] if abs(p - q) > 1e-3 else None
                else:
                    sites = ("r", "l")
                    mid = ox + (ax + aw + bx) / 2
                    p, q = oy + ay + ah / 2, oy + by + bh / 2
                    via = [(mid, p), (mid, q)] if abs(p - q) > 1e-3 else None
                made.append(self.link(s, a, b, src_site=sites[0], dst_site=sites[1], via=via,
                                      route="ortho" if via else "straight").connector)
        outline = "; ".join(f"{texts[p][0]} leads {', '.join(texts[c][0] for c in cs)}"
                            for p, cs in kids.items() if cs)
        return self._finish(s, made, "org_chart", title, alt or (
            f"Org chart{': ' + title if title else ''}. {texts['0'][0]} at the top. "
            f"{outline}."))

    def mind_map(self, s, root, branches, *, box=None, font=None, title=None, alt=None):
        """A mind map: `root` in the middle and branches [(text, [leaf texts])] split left and
        right to balance their leaves. Branches are pills in palette colours, leaves sit on an
        underline in their branch's colour, links are curved connectors glued at both ends."""
        children = {"root": []}
        made = []
        shapes = {}
        sizes = {}
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        texts = {"root": root}
        rw, rh, rlines = self.fit_shape(root, "ellipse", font + 5, bold=True, max_w=2.6,
                                        min_w=1.8, min_h=0.85)
        sizes["root"] = (rw, rh)
        for k, (branch, leaves) in enumerate(branches):
            bid = f"b{k}"
            children["root"].append(bid)
            texts[bid] = branch
            bw, bh, _ = self.fit_shape(branch, "roundRect", font + 1, bold=True, max_w=2.4,
                                       min_w=1.3, min_h=0.48, adj={"adj": 50000})
            sizes[bid] = (bw + 0.1, bh)
            children[bid] = []
            for j, leaf in enumerate(leaves):
                lid = f"{bid}.{j}"
                children[bid].append(lid)
                texts[lid] = leaf
                sizes[lid] = (self.measure(leaf, font) + 0.3, font * LINE_SPACING / 72 + 0.14)
        for level, sib, gap in ((0.6, 0.12, 0.3), (0.45, 0.08, 0.2), (0.35, 0.04, 0.12)):
            boxes, side, size = layout.side_tree("root", children, sizes, level_gap=level,
                                                 sib_gap=sib, branch_gap=gap)
            if size[0] <= box[2] + 1e-6 and size[1] <= box[3] + 1e-6:
                break
        ox, oy = self._place(size, box, "mind map")
        root_accent = self._hex(2)
        x, y, w, h = boxes["root"]
        shapes["root"] = self._card(s, ox + x, oy + y, w, h, [("\n".join(rlines), font + 5,
                                                               True, None)],
                                    tint(root_accent, 0.25), root_accent, prst="ellipse",
                                    width=2.0, name=f"Node: {root}")
        made.append(shapes["root"])
        branch_slots = [slot for slot in _BRANCH_SLOTS if slot != 2]  # amber is the root's
        for k, bid in enumerate(children["root"]):
            accent = self._hex(branch_slots[k % len(branch_slots)])
            x, y, w, h = boxes[bid]
            br = self._card(s, ox + x, oy + y, w, h, [(texts[bid], font + 1, True, None)],
                            tint(accent, 0.18), accent, width=1.75, radius=h / 2,
                            name=f"Node: {texts[bid]}")
            made.append(br)
            out, into = ("r", "l") if side[bid] > 0 else ("l", "r")
            made.append(self.link(s, shapes["root"], br, src_site=out, dst_site=into,
                                  route="curved", arrow=None, color=_rgb(accent),
                                  width=1.75).connector)
            for lid in children[bid]:
                x, y, w, h = boxes[lid]
                leaf = self._card(s, ox + x, oy + y, w, h, [(texts[lid], font, False, None)],
                                  None, accent, prst="rect", width=1.5, anchor="bottom",
                                  margins=(0.05, 0.04), name=f"Node: {texts[lid]}")
                _geom(leaf, [([(0, 1), (1, 1)], False, False)], [("l", 0, 1), ("r", 1, 1)])
                made.append(leaf)
                made.append(self.link(s, br, leaf, src_site=out, dst_site=0 if out == "r" else 1,
                                      route="curved", arrow=None, color=_rgb(accent),
                                      width=1.25).connector)
        outline = "; ".join(f"{texts[b]}: {', '.join(texts[c] for c in children[b])}"
                            for b in children["root"])
        return self._finish(s, made, "mind_map", title, alt or (
            f"Mind map{': ' + title if title else ''}. {root}. {outline}."))

    def tree_view(self, s, rows, *, box=None, font=None, title=None, alt=None):
        """A file tree. rows: [(depth, name, note, highlight)] in display order; names ending
        in '/' are folders. Guides are elbow connectors glued from the parent's icon to the
        child's; notes line up in a column on the right; a highlighted row gets a band."""
        icon = 0.2
        indent = 0.32
        made = []
        parents = []
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        pitch = font * LINE_SPACING / 72 + 0.12
        if len(rows) * pitch > box[3] + 1e-6:
            raise LayoutError(f"tree view: {len(rows)} rows need {len(rows) * pitch:.1f} in of "
                              f"height but the region has {box[3]:.1f} in: split the diagram")
        green = blend(self._hex(3), "0F172A", 0.25)
        line = self.t["muted"]
        names_w = max(d * indent + icon + 0.12 + self.measure(n.rstrip("/"), font,
                                                               bold=n.endswith("/"))
                      for d, n, *_ in rows)
        notes_w = max((self.measure(r[2], font, italic=True) for r in rows if len(r) > 2 and
                       r[2]), default=0.0)
        w = names_w + (0.6 + notes_w if notes_w else 0.0) + 0.3
        ox, oy = self._place((w, len(rows) * pitch), box, "tree view")
        note_x = ox + names_w + 0.6
        for i, (depth, name, *rest) in enumerate(rows):
            folder = name.endswith("/")
            highlight = len(rest) > 1 and rest[1]
            note = rest[0] if rest else ""
            x = ox + depth * indent
            y = oy + i * pitch
            if highlight:
                made.append(self._shape(s, "roundRect", x - 0.08, y + 0.02,
                                        ox + names_w + 0.1 - x, pitch - 0.04,
                                        tint(self._hex(2), 0.18), self._hex(2),
                                        name=f"Deco: highlight {name}"))
            ih = icon * (0.8 if folder else 1.0)
            ic = self._shape(s, "rect", x, y + (pitch - ih) / 2, icon * (1.0 if folder else 0.8),
                             ih, self.t["line"] if folder else self.t["muted"], None,
                             name=f"Node: icon {name}")
            if folder:
                _geom(ic, [([(0, 0), (0.42, 0), (0.54, 0.18), (1, 0.18), (1, 1), (0, 1)], True,
                            True)], [("b", 0.32, 1), ("l", 0, 0.6)])
            else:
                _geom(ic, [([(0, 0), (0.62, 0), (1, 0.32), (1, 1), (0, 1)], True, True)],
                      [("b", 0.4, 1), ("l", 0, 0.6)])
            made.append(ic)
            label = name.rstrip("/") or name
            tw = self.measure(label, font, bold=folder) + 0.05
            made.append(self._text(s, x + icon + 0.1, y, tw, pitch, label, font, bold=folder,
                                   name=f"Label: {label}"))
            if note:
                made.append(self._text(s, note_x, y, self.measure(note, font, italic=True) + 0.05,
                                       pitch, note, font, italic=True, color=green,
                                       name=f"Label: note {note}"))
            while parents and parents[-1][0] >= depth:
                parents.pop()
            if parents:
                made.append(self.link(s, parents[-1][1], ic, src_site=0, dst_site=1, arrow=None,
                                      color=_rgb(line), width=1.0).connector)
            parents.append((depth, ic))
        listing = "; ".join(f"{'  ' * d}{n}{' (' + r[0] + ')' if r and r[0] else ''}"
                            for d, n, *r in rows)
        return self._finish(s, made, "tree_view", title, alt or (
            f"File tree{': ' + title if title else ''}: {listing}."))

    # -- grids and lanes ------------------------------------------------------------------------
    def block(self, s, columns, arrows=(), *, box=None, font=None, title=None, alt=None):
        """A block diagram: columns [(id, label, [(id, label, span)])] drawn as panels side by
        side; a block with span 'full' takes a whole row of its panel and the other blocks of
        that panel share the rows below it, otherwise blocks stack. arrows: [(src, dst)] between
        panels or blocks, drawn as straight glued connectors."""
        arrow_gap = 0.75
        ends = {}
        gap = 0.2
        made = []
        pad = 0.22
        plans = []
        shapes = {}
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        title_h = (font + 2) * LINE_SPACING / 72 + 0.3
        for cid, label, blocks in columns:
            full = [b for b in blocks if len(b) > 2 and b[2] == "full"]
            need = max([self.measure(b[1], font) + 0.4 for b in blocks] + [1.0])
            rest = [b for b in blocks if not (len(b) > 2 and b[2] == "full")]
            rows = [[b] for b in full] + ([rest] if full and rest else [[b] for b in rest])
            across = max(len(r) for r in rows) if rows else 1
            plans.append((cid, label, rows, across, need))
        fixed = sum(2 * pad + gap * (p[3] - 1) for p in plans) + arrow_gap * (len(plans) - 1)
        units = sum(p[3] for p in plans)
        unit = min(max(p[4] for p in plans), (box[2] - fixed) / units, 2.4)
        if unit < max(p[4] for p in plans) - 1e-6:
            unit = (box[2] - fixed) / units
            if any(unit < p[4] - 1e-6 for p in plans):
                raise LayoutError("block diagram: the blocks do not fit side by side at 11 pt or "
                                  "more: split the diagram")
        h = min(box[3], max(len(p[2]) for p in plans) * 1.0 + title_h + pad)
        total_w = units * unit + fixed
        ox, oy = self._place((total_w, h), box, "block diagram")
        x = ox
        for k, (cid, label, rows, across, _) in enumerate(plans):
            pw = across * unit + gap * (across - 1) + 2 * pad
            panel = self._card(s, x, oy, pw, h, [(label, font + 2, True, None)],
                               self.t["panel"], self.t["line"], anchor="top", radius=0.12,
                               margins=(0.1, 0.14), name=f"Container: {label}")
            made.append(panel)
            shapes[cid] = panel
            accent = self._hex(_BRANCH_SLOTS[k % 6])
            avail = h - title_h - pad
            full_h = (font * LINE_SPACING / 72 + 0.5) if len(rows) > 1 and len(rows[0]) == 1 \
                and any(len(r) > 1 for r in rows) else None
            heights = [avail / len(rows) - gap * (len(rows) - 1) / len(rows)] * len(rows)
            top = oy + title_h
            if full_h:
                heights = [full_h] + [(avail - full_h - gap * (len(rows) - 1)) / (len(rows) - 1)] \
                    * (len(rows) - 1)
            for row, rh in zip(rows, heights):
                bw = (pw - 2 * pad - gap * (len(row) - 1)) / len(row)
                for j, (bid, blabel, *_) in enumerate(row):
                    sp = self._card(s, x + pad + j * (bw + gap), top, bw, rh,
                                    [(blabel, font, False, None)], tint(accent, 0.16), accent,
                                    radius=0.1, name=f"Node: {blabel}")
                    made.append(sp)
                    shapes[bid] = sp
                top += rh + gap
            x += pw + arrow_gap
        for src, dst in arrows:
            a, b = shapes[src], shapes[dst]
            ay0, ay1 = a.top, a.top + a.height
            by0, by1 = b.top, b.top + b.height
            lo, hi = max(ay0, by0), min(ay1, by1)
            y = (lo + hi) / 2 if hi > lo else (by0 + by1) / 2
            ends.setdefault(src, []).append(("r", 1.0, (y - ay0) / (ay1 - ay0)))
            ends.setdefault(dst, []).append(("l", 0.0, (y - by0) / (by1 - by0)))
        for sid, sites in ends.items():
            _custom_sites(shapes[sid], "roundRect", sites)
        used = {sid: 0 for sid in ends}
        for src, dst in arrows:
            ia, ib = used[src], used[dst]
            used[src] += 1
            used[dst] += 1
            made.append(self.link(s, shapes[src], shapes[dst], src_site=ia, dst_site=ib,
                                  route="straight", width=2.0).connector)
        names = {cid: label for cid, label, _ in columns}
        names.update({b[0]: b[1] for _, _, blocks in columns for b in blocks})
        flows = "; ".join(f"{names[a]} to {names[b]}" for a, b in arrows)
        parts = "; ".join(f"{label}: {', '.join(b[1] for b in blocks)}"
                          for _, label, blocks in columns)
        return self._finish(s, made, "block", title, alt or (
            f"Block diagram{': ' + title if title else ''}. {parts}. Arrows: {flows}."))

    def user_journey(self, s, stages, *, box=None, font=None, title=None, alt=None):
        """A user journey: stages [(stage, [(task, score 1-5, [actors])])] drawn as stage bands
        over task cards, a satisfaction plot whose score dots are linked by glued connectors, and
        actor chips under each task."""
        col_gap = 0.12
        dot_d = 0.4
        dots = []
        heads = ("Stage", "Task", "Satisfaction", "Actors")
        made = []
        scale = "1 (poor) \u2013 5 (great)"
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        small = TYPE_SCALE["label"]
        tasks = [(k, *t) for k, (_, ts) in enumerate(stages) for t in ts]
        label_w = max(self.measure(t, small, bold=True) for t in heads) + 0.25
        label_w = max(label_w, self.measure(scale, small) + 0.1)
        tick_w = self.measure("5", small) + 0.12
        col_w = (box[2] - label_w - tick_w - col_gap * (len(tasks) - 1)) / len(tasks)
        task_lines = [self.wrap(t[1], col_w - 0.2, font) for t in tasks]
        if any(self.measure(line, font) > col_w - 0.16 for lines in task_lines for line in lines):
            raise LayoutError("user journey: a task name is too long for its column at 11 pt "
                              "or more: split the diagram")
        chip_h = small * LINE_SPACING / 72 + 0.08
        actors_h = max(len(t[3]) for t in tasks) * (chip_h + 0.06)
        line_h = font * LINE_SPACING / 72
        stage_h = (font + 1) * LINE_SPACING / 72 + 0.16
        task_h = max(len(lines) for lines in task_lines) * line_h + 0.3
        plot_h = box[3] - stage_h - task_h - actors_h - 0.12 * 3
        if plot_h < 1.6:
            raise LayoutError("user journey: no room for the satisfaction plot: split the "
                              "diagram")
        x0, y0 = box[0] + label_w + tick_w, box[1]
        ys = (y0, y0 + stage_h + 0.12, y0 + stage_h + task_h + 0.24,
              y0 + stage_h + task_h + plot_h + 0.36)
        for text, y, hh in zip(heads, ys, (stage_h, task_h, plot_h, chip_h)):
            made.append(self._text(s, box[0], y + hh / 2 - small * LINE_SPACING / 72 / 2 - 0.01,
                                   self.measure(text, small, bold=True) + 0.05,
                                   small * LINE_SPACING / 72 + 0.02, text, small, bold=True,
                                   color=self.t["line"], name=f"Label: row {text}"))
        made.append(self._text(s, box[0], ys[2] + plot_h / 2 + 0.12, self.measure(scale, small)
                               + 0.05, small * LINE_SPACING / 72 + 0.02, scale, small,
                               color=self.t["muted"], name="Label: scale"))
        cx = [x0 + i * (col_w + col_gap) + col_w / 2 for i in range(len(tasks))]
        for k, (stage, _) in enumerate(stages):
            accent = self._hex(_BRANCH_SLOTS[k % 6])
            idx = [i for i, t in enumerate(tasks) if t[0] == k]
            sx = cx[idx[0]] - col_w / 2
            sw = cx[idx[-1]] + col_w / 2 - sx
            made.append(self._card(s, sx, ys[0], sw, stage_h, [(stage, font + 1, True, None)],
                                   tint(accent, 0.18), accent, name=f"Node: stage {stage}"))
        plot_top, plot_bottom = ys[2] + 0.15, ys[2] + plot_h - 0.15

        def yv(v):
            return plot_bottom - (v - 1) / 4 * (plot_bottom - plot_top)

        for v in range(1, 6):
            made.append(self._line(s, x0 - 0.05, yv(v), cx[-1] + col_w / 2, yv(v),
                                   self.t["subtle"], 0.75, "dash", name=f"Deco: grid {v}"))
            made.append(self._text(s, x0 - tick_w, yv(v) - small * LINE_SPACING / 72 / 2,
                                   tick_w - 0.08, small * LINE_SPACING / 72, str(v), small,
                                   align="right", color=self.t["muted"], name=f"Deco: tick {v}"))
        for i, (k, task, score, actors) in enumerate(tasks):
            accent = self._hex(_BRANCH_SLOTS[k % 6])
            made.append(self._card(s, cx[i] - col_w / 2, ys[1], col_w, task_h,
                                   [("\n".join(task_lines[i]), font, False, None)],
                                   "FFFFFF", accent, name=f"Node: task {task}"))
            made.append(self._line(s, cx[i], ys[2], cx[i], ys[2] + plot_h, self.t["subtle"],
                                   0.75, "dash", name="Deco: column"))
            colour = self._hex({1: 4, 2: 4, 3: 2, 4: 3, 5: 3}[max(1, min(5, int(score)))])
            dot = self._card(s, cx[i] - dot_d / 2, yv(score) - dot_d / 2, dot_d, dot_d,
                             [(str(score), small + 1, True, None)], tint(colour, 0.25), colour,
                             prst="ellipse", width=1.75, margins=(0.0, 0.0),
                             name=f"Node: score {task}")
            dots.append(dot)
            for j, actor in enumerate(actors):
                cw = self.measure(actor, small) + 0.3
                tone = self._hex(5) if j else self.t["line"]
                made.append(self._card(s, cx[i] - cw / 2, ys[3] + j * (chip_h + 0.06), cw, chip_h,
                                       [(actor, small, False, None)], tint(tone, 0.12), tone,
                                       radius=chip_h / 2, margins=(0.06, 0.0),
                                       name=f"Label: actor {task} {actor}"))
        for a, b in zip(dots, dots[1:]):
            made.append(self.link(s, a, b, route="straight", arrow=None,
                                  color=_rgb(self.t["muted"]), width=2.0).connector)
        made += dots
        steps = "; ".join(f"{stages[k][0]}: {t} (score {sc}, {', '.join(ac)})"
                          for k, t, sc, ac in tasks)
        return self._finish(s, made, "user_journey", title, alt or (
            f"User journey{': ' + title if title else ''}. {steps}."))

    def event_model(self, s, lanes, items, *, box=None, font=None, title=None, alt=None):
        """An event model: lanes [(name, [item kinds])] stacked as swimlanes, items [(order,
        kind, name, fields)] with kind ui, pcr (processor), cmd (command), rmo (read model) or
        evt (event), placed left to right in order and linked in order by glued connectors."""
        ends = {}
        lane_gap = 0.08
        made = []
        shapes = []
        sizes = []
        xs = []
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"] - 1)
        items = sorted(items, key=lambda it: it[0])
        lane_h = (box[3] - lane_gap * (len(lanes) - 1)) / len(lanes)
        lane_of = {kind: i for i, (_, kinds) in enumerate(lanes) for kind in kinds}
        small = TYPE_SCALE["label"]
        title_w = max(self.measure(name, small, bold=True) for name, _ in lanes) + 0.3
        for _, _, name, fields in items:
            w = max(self.measure(name, font, bold=True),
                    self.measure(fields, small, mono=True) if fields else 0.0) + 0.3
            sizes.append(max(w, 1.2))
        item_h = font * LINE_SPACING / 72 + (small * LINE_SPACING / 72 + 0.06 if any(
            it[3] for it in items) else 0.0) + 0.3
        item_h = min(max(item_h, 0.75), lane_h - 0.2)
        for i, (_, kind, _, _) in enumerate(items):
            x = 0.0
            lane = lane_of[kind]
            if i:
                prev_lane = lane_of[items[i - 1][1]]
                x = xs[-1] + (sizes[i - 1] + 0.3 if prev_lane == lane else sizes[i - 1] * 0.55)
            x = max(x, ends.get(lane, -1.0) + 0.3)
            xs.append(x)
            ends[lane] = x + sizes[i]
        room = box[2] - title_w - 0.2
        span = max(x + w for x, w in zip(xs, sizes))
        if span > room + 1e-6:
            raise LayoutError(f"event model needs {span + title_w:.1f} in of width but its region "
                              f"has {box[2]:.1f} in: split the diagram")
        ox = box[0] + title_w + (room - span) / 2
        for i, (name, _) in enumerate(lanes):
            made.append(self._card(s, box[0], box[1] + i * (lane_h + lane_gap), box[2], lane_h,
                                   [(name, small, True, self.t["line"])], self.t["panel"], None,
                                   prst="rect", align="left", anchor="top", margins=(0.12, 0.08),
                                   name=f"Container: {name}"))
        for (_, kind, name, fields), x, w in zip(items, xs, sizes):
            fill, line = _EVENT.get(kind, (self.t["panel"], self.t["line"]))
            lane = lane_of[kind]
            y = box[1] + lane * (lane_h + lane_gap) + (lane_h - item_h) / 2 + 0.06
            sp = self._card(s, ox + x, y, w, item_h, [(name, font, True, None)], fill, line,
                            prst="rect", name=f"Node: {name}")
            if fields:
                self.run(self.para(sp.text_frame, align="center"), fields, size=small,
                         mono=True, color=self.TEXT)
            made.append(sp)
            shapes.append((sp, lane))
        for (a, la), (b, lb) in zip(shapes, shapes[1:]):
            if la == lb:
                sites = ("r", "l")
            else:
                sites = ("b", "t") if lb > la else ("t", "b")
            made.append(self.link(s, a, b, src_site=sites[0], dst_site=sites[1],
                                  route="straight" if la != lb else "ortho").connector)
        steps = "; ".join(f"{kind} {name}" for _, kind, name, _ in items)
        return self._finish(s, made, "event_model", title, alt or (
            f"Event model{': ' + title if title else ''}. Lanes: "
            f"{', '.join(n for n, _ in lanes)}. In order: {steps}."))

    # -- time -----------------------------------------------------------------------------------
    def timeline(self, s, events, *, box=None, font=None, title=None, alt=None):
        """A timeline: events [(when, text)] on an axis with a circle per event, cards
        alternating above and below linked by glued dashed connectors; numeric years leave room
        for the years in between, which get a small tick."""
        cards = []
        stem = 0.35
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        keys = [str(e[0]) for e in events]
        numeric = all(k.isdigit() for k in keys)
        small = TYPE_SCALE["label"]
        if numeric:
            lo, hi = int(keys[0]), int(keys[-1])
            slots = list(range(lo, hi + 1)) if hi - lo < 15 else [int(k) for k in keys]
        else:
            slots = keys
        dw, dh, _ = self.fit_shape(max(keys, key=len), "ellipse", font, bold=True, max_w=1.2,
                                   min_w=0.7, min_h=0.7)
        d = max(dw, dh)
        card_w = min(2.6, (box[2] - d - 0.4) / max(len(slots) - 1, 1) * 2 - 0.2)
        slot_of = {str(v): i for i, v in enumerate(slots)}
        for _, text in events:
            tw, th, lines = self.text_size(text, font, max_w=card_w - 0.3)
            if tw > card_w - 0.2 + 1e-6:
                raise LayoutError("timeline: an event text is too long for its card: split the "
                                  "diagram")
            cards.append((max(tw + 0.4, 1.4), th + 0.36, lines))
        card_h = max(c[1] for c in cards)
        height = 2 * (card_h + stem) + d
        ox, oy = self._place((box[2], height), box, "timeline")
        axis_y = oy + card_h + stem + d / 2
        made = [self._line(s, box[0], axis_y, box[0] + box[2], axis_y, self.t["line"], 3.0,
                           name="Deco: axis", arrow="triangle")]
        x_first = box[0] + max(0.2 + d / 2, cards[0][0] / 2)
        x_last = box[0] + box[2] - max(0.4 + d / 2, cards[-1][0] / 2)
        step = (x_last - x_first) / max(len(slots) - 1, 1)
        x_of = {k: x_first + slot_of[k] * step for k in keys}
        for v in slots:
            if str(v) in x_of:
                continue
            x = x_first + slot_of[str(v)] * step
            made.append(self._shape(s, "ellipse", x - 0.08, axis_y - 0.08, 0.16, 0.16,
                                    self.t["light"], self.t["muted"], name=f"Deco: tick {v}"))
            tw = self.measure(str(v), small) + 0.05
            made.append(self._text(s, x - tw / 2, axis_y + 0.14, tw,
                                   small * LINE_SPACING / 72, str(v), small,
                                   color=self.t["muted"], name=f"Deco: year {v}"))
        for i, ((when, text), (cw, ch, lines)) in enumerate(zip(events, cards)):
            accent = self._hex(i)
            x = x_of[str(when)]
            dot = self._card(s, x - d / 2, axis_y - d / 2, d, d, [(str(when), font, True, None)],
                             tint(accent, 0.22), accent, prst="ellipse", width=2.0,
                             margins=(0.0, 0.0), name=f"Node: {when}")
            cx = min(max(x - cw / 2, box[0]), box[0] + box[2] - cw)
            up = i % 2 == 0
            cy = oy + (0 if up else card_h + 2 * stem + d)
            card = self._card(s, cx, cy if up else cy + card_h - ch, cw, ch,
                              [("\n".join(lines), font, False, None)], tint(accent, 0.14),
                              accent, name=f"Node: {text}")
            made += [card, dot]
            made.append(self.link(s, card, dot, src_site="b" if up else "t",
                                  dst_site="t" if up else "b", route="straight", arrow=None,
                                  style="dashed", color=_rgb(accent), width=1.5).connector)
        steps = "; ".join(f"{w}: {t}" for w, t in events)
        return self._finish(s, made, "timeline", title, alt or (
            f"Timeline{': ' + title if title else ''}. {steps}."))

    def gantt(self, s, sections, milestones=(), *, box=None, font=None, title=None, alt=None):
        """A Gantt chart as a native table (Task, Start, End and one column per week under a
        month header) with a bar shape per task and a diamond per milestone. sections: [(name,
        [(task, start, end)])] with ISO dates (end inclusive); milestones: [(name, date)]."""
        k = -1
        made = []
        months = []
        rows = []  # (kind, label, start, end)
        box = box or CONTENT
        day = dt.date.fromisoformat
        dates = [day(d) for _, ts in sections for _, a, b in ts for d in (a, b)]
        dates += [day(d) for _, d in milestones]
        font = _font(font or TYPE_SCALE["node"] - 1)
        small = TYPE_SCALE["label"]
        week0 = min(dates) - dt.timedelta(days=min(dates).weekday())
        weeks = math.ceil(((max(dates) - week0).days + 1) / 7)
        week_days = [week0 + dt.timedelta(days=7 * i) for i in range(weeks)]
        for name, tasks in sections:
            rows.append(("section", name, None, None))
            rows += [("task", t, day(a), day(b)) for t, a, b in tasks]
        rows += [("milestone", m, day(d), day(d)) for m, d in milestones]
        n_rows = len(rows) + 2
        row_h = min(0.4, box[3] / n_rows)
        if row_h < font * LINE_SPACING / 72 + 0.1:
            raise LayoutError(f"Gantt chart: {n_rows} rows do not fit the region's height at 11 pt "
                              "or more: split the diagram")
        date_w = max(self.measure(_fmt_day(d), font) for d in dates) + 0.3
        task_w = max(self.measure(r[1], font, bold=r[0] == "section") for r in rows) + 0.3
        task_w = max(task_w, self.measure("Task", font, bold=True) + 0.3)
        week_w = (box[2] - task_w - 2 * date_w) / weeks
        for i, d in enumerate(week_days):
            if months and months[-1][0] == (d.year, d.month):
                months[-1][2] = i
            else:
                months.append([(d.year, d.month), i, i])
        need = max(self.measure(f"{dt.date(y, m, 1):%B %Y}", font, bold=True) /
                   (b - a + 1) for (y, m), a, b in months)
        if week_w < max(self.measure("22", small) + 0.12, need + 0.04):
            raise LayoutError("Gantt chart: too many weeks for the region's width: split the "
                              "diagram")
        x0, y0 = box[0], box[1] + (box[3] - n_rows * row_h) / 2
        frame = s.shapes.add_table(n_rows, weeks + 3, Inches(x0), Inches(y0), Inches(box[2]),
                                   Inches(n_rows * row_h))
        tbl = frame.table
        tbl._tbl.tblPr.set("firstRow", "0")
        tbl._tbl.tblPr.set("bandRow", "0")
        for i, wcol in enumerate([task_w, date_w, date_w] + [week_w] * weeks):
            tbl.columns[i].width = Emu(round(wcol * EMU_IN))
        for r in range(n_rows):
            tbl.rows[r].height = Emu(round(row_h * EMU_IN))
        head, band, grid = self.t["panel"], self.t["light"], self.t["card_border"]

        def put(r, c, text, size=font, bold=False, align="center", fill="FFFFFF", color=None):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = 3
            tf = cell.text_frame
            tf.word_wrap = False
            p = self.para(tf, first=True, align=align)
            if text:
                self.run(p, text, size=size, bold=bold, color=_rgb(color) if color else self.TEXT)
            _cell_lines(cell, grid)

        for c, text in enumerate(("Task", "Start", "End")):
            put(0, c, text, bold=True, align="left" if c == 0 else "center", fill=head)
            put(1, c, "", fill=head)
            tbl.cell(0, c).merge(tbl.cell(1, c))
        for (y, m), a, b in months:
            for c in range(a, b + 1):
                put(0, 3 + c, f"{dt.date(y, m, 1):%B %Y}" if c == a else "", bold=True,
                    fill=head)
            if b > a:
                tbl.cell(0, 3 + a).merge(tbl.cell(0, 3 + b))
        for i, d in enumerate(week_days):
            put(1, 3 + i, str(d.day), size=small, fill=head, color=self.t["line"])
        for r, (kind, label, a, b) in enumerate(rows, 2):
            fill = band if kind == "section" else "FFFFFF"
            put(r, 0, label, bold=kind == "section", align="left", fill=fill)
            put(r, 1, _fmt_day(a) if a else "", fill=fill)
            put(r, 2, _fmt_day(b) if b else "", fill=fill)
            for c in range(weeks):
                put(r, 3 + c, "", fill=fill)
        alt_text(frame, alt or (
            f"Gantt chart{': ' + title if title else ''}. " + "; ".join(
                f"{label}: {_fmt_day(a)} to {_fmt_day(b)}" for kind, label, a, b in rows
                if kind != "section")))
        grid_x = x0 + task_w + 2 * date_w

        def x_at(d):
            return grid_x + (d - week0).days / 7 * week_w

        for r, (kind, label, a, b) in enumerate(rows, 2):
            y = y0 + r * row_h
            if kind == "section":
                k += 1
                continue
            if kind == "task":
                accent = self._hex(_BRANCH_SLOTS[k % 6])
                made.append(self._shape(s, "roundRect", x_at(a), y + 0.08,
                                        x_at(b + dt.timedelta(days=1)) - x_at(a), row_h - 0.16,
                                        tint(accent, 0.25), accent, name=f"Node: bar {label}"))
                made[-1].adjustments[0] = 0.3
            else:
                cx = x_at(a) + week_w / 14
                m = min(0.26, row_h - 0.08)
                made.append(self._shape(s, "diamond", cx - m / 2, y + (row_h - m) / 2, m, m,
                                        tint(self._hex(4), 0.3), self._hex(4),
                                        name=f"Node: milestone {label}"))
        return self._finish(s, made, "gantt", title, f"Bars of the Gantt chart "
                                                     f"{title or ''}".rstrip())

    # -- sets and flows -------------------------------------------------------------------------
    def quadrant(self, s, points, quadrants, axes, *, box=None, font=None, title=None,
                 alt=None):
        """A quadrant chart. points: [(label, x, y)] with x and y from 0 to 1; quadrants: names
        in reading order (top left, top right, bottom left, bottom right); axes: (x, y) names,
        or ((low, high), (low, high)) end labels. Quadrant titles sit at the top of their
        quadrant, centred, or against the frame when a point would crowd them."""
        dot_d = 0.13
        made = []
        shades = (0.08, 0.15, 0.05, 0.03)
        taken = []
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        py = box[1]
        small = TYPE_SCALE["label"] + 1
        line_h = small * LINE_SPACING / 72
        ph = box[3] - line_h - 0.2
        pw = min(box[2] - line_h - 0.3, ph * 1.7)
        xs, ys = ((f"Low {a.lower()}", f"High {a.lower()}") if isinstance(a, str) else tuple(a)
                  for a in axes)
        side = max(self.measure(t, small) for t in ys) + 0.1
        px = max(box[0] + (box[2] - pw - line_h - 0.2) / 2 + line_h + 0.2,
                 box[0] + side / 2 + line_h / 2 + 0.1)
        if pw < 4:
            raise LayoutError("quadrant chart: the region is too small: split the diagram")
        accent = self._hex(1)
        for i, name in enumerate(quadrants):
            qx, qy = px + (i % 2) * pw / 2, py + (i // 2) * ph / 2
            made.append(self._shape(s, "rect", qx, qy, pw / 2, ph / 2, tint(accent, shades[i]),
                                    None, name=f"Deco: quadrant {name}"))
        made.append(self._shape(s, "rect", px, py, pw, ph, None, tint(accent, 0.35),
                                name="Deco: plot frame"))
        made.append(self._line(s, px + pw / 2, py, px + pw / 2, py + ph, tint(accent, 0.35),
                               name="Deco: divider"))
        made.append(self._line(s, px, py + ph / 2, px + pw, py + ph / 2, tint(accent, 0.35),
                               name="Deco: divider"))
        dots = [(px + x * pw, py + (1 - y) * ph) for _, x, y in points]
        h = font * LINE_SPACING / 72 + 0.04
        marks = [(x - dot_d / 2, y - dot_d / 2, x + dot_d / 2, y + dot_d / 2) for x, y in dots]
        tops = [py + (i // 2) * ph / 2 + 0.08 for i in range(len(quadrants))]
        widths = [self.measure(name, font) + 0.1 for name in quadrants]

        def title_xs(i):
            """Where the title of quadrant i may go: centred, at the frame, at the divider."""
            left = px + (i % 2) * pw / 2
            ends = (left + 0.12, left + pw / 2 - widths[i] - 0.12)
            return (left + pw / 4 - widths[i] / 2,) + (ends if i % 2 == 0 else ends[::-1])

        def clear(i, x):
            return not any(_hit((x, tops[i], x + widths[i], tops[i] + h), m, 0.15) for m in marks)

        options = [title_xs(i) for i in range(len(quadrants))]
        for k in (0, 1):  # every title centred, else every title against the frame
            title_x = [opt[k] for opt in options]
            if all(clear(i, x) for i, x in enumerate(title_x)):
                break
        else:  # each title at its first spot clear of the points
            title_x = [next((x for x in opt if clear(i, x)), opt[0])
                       for i, opt in enumerate(options)]
        for i, name in enumerate(quadrants):
            made.append(self._text(s, title_x[i], tops[i], widths[i], h, name, font,
                                   align="center", name=f"Label: {name}"))
            taken.append((title_x[i], tops[i], title_x[i] + widths[i], tops[i] + h))
        for i, text in enumerate(xs):
            w = self.measure(text, small) + 0.1
            made.append(self._text(s, px + (i + 0.5) * pw / 2 - w / 2, py + ph + 0.1, w, line_h,
                                   text, small, align="center", color=self.t["line"],
                                   name=f"Label: {text}"))
        for i, text in enumerate(reversed(ys)):
            cx, cy = px - 0.12 - line_h / 2, py + (i + 0.5) * ph / 2
            w = self.measure(text, small) + 0.1
            lab = self._text(s, cx - w / 2, cy - line_h / 2, w, line_h, text, small,
                             align="center", color=self.t["line"], name=f"Label: {text}")
            lab.rotation = 270
            made.append(lab)
        taken += marks
        for (label, _, _), (x, y) in zip(points, dots):
            made.append(self._shape(s, "ellipse", x - dot_d / 2, y - dot_d / 2, dot_d, dot_d,
                                    blend(accent, "0F172A", 0.45), None, name=f"Node: {label}"))
            gap = dot_d / 2 + 0.05
            w = self.measure(label, small) + 0.1
            spots = [(x - w / 2, y + gap), (x - w / 2, y - gap - line_h),
                     (x + gap, y - line_h / 2), (x - gap - w, y - line_h / 2)]
            for lx, ly in spots:
                rect = (lx, ly, lx + w, ly + line_h)
                inside = lx >= px and lx + w <= px + pw and ly >= py and ly + line_h <= py + ph
                if inside and not any(_hit(rect, t, 0.0) for t in taken):
                    break
            else:
                raise LayoutError(f"quadrant chart: no room for the label {label!r}: move the "
                                  "point or split the chart")
            taken.append(rect)
            made.append(self._text(s, lx, ly, w, line_h, label, small, align="center",
                                   name=f"Label: point {label}"))
        listing = "; ".join(f"{label} ({x:g}, {y:g})" for label, x, y in points)
        return self._finish(s, made, "quadrant", title, alt or (
            f"Quadrant chart{': ' + title if title else ''}. Quadrants: "
            f"{', '.join(quadrants)}; x: {xs[0]} to {xs[1]}; y: {ys[0]} to {ys[1]}. "
            f"Points: {listing}."))

    def venn(self, s, sets, overlaps=(), *, title=None, box=None, font=None, alt=None):
        """A Venn diagram of two or three sets: sets [(id, label)] drawn as translucent ovals
        with coloured outlines; overlaps [((ids), label)] name the shared regions."""
        made = []
        slots = (5, 3, 2)
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        n = len(sets)
        if n not in (2, 3):
            raise ValueError("venn() draws two or three sets")
        angles = (210, 330, 90) if n == 3 else (180, 0)
        r = min(box[3] / 2.95, box[2] / 3.1) if n == 3 else min(box[3] / 2.0, box[2] / 3.3)
        d = 0.62 * r
        cx0, cy0 = box[0] + box[2] / 2, box[1] + box[3] / 2 + (0.25 * d if n == 3 else 0.0)
        centres = [(cx0 + d * math.cos(math.radians(a)), cy0 - d * math.sin(math.radians(a)))
                   for a in angles]

        def clearance(rect, members):
            """How far the label rect stays inside its member sets and outside the others
            (negative: it crosses an outline)."""
            out = math.inf
            corners = [(rect[0], rect[1]), (rect[2], rect[1]), (rect[0], rect[3]),
                       (rect[2], rect[3])]
            for i in range(n):
                if i in members:
                    out = min(out, min(r - math.dist(c, centres[i]) for c in corners))
                else:
                    out = min(out, _gap(rect, centres[i]) - r)
            return out

        ids = [sid for sid, _ in sets]
        for i, (_, label) in enumerate(sets):
            accent = self._hex(slots[i])
            x, y = centres[i]
            made.append(self._shape(s, "ellipse", x - r, y - r, 2 * r, 2 * r, accent, accent,
                                    2.0, alpha=0.12, name=f"Node: {label}"))
        regions = [((i,), label, slots[i], font + 2) for i, (_, label) in enumerate(sets)]
        regions += [(tuple(ids.index(m) for m in members), label, None, font)
                    for members, label in overlaps]
        steps = int(2 * r / 0.05) + int(2 * d / 0.05) + 2
        x_lo = min(x for x, _ in centres) - r
        y_lo = min(y for _, y in centres) - r
        for members, label, slot, size in regions:
            best = None
            h = size * LINE_SPACING / 72 + 0.02
            w = self.measure(label, size) + 0.1
            for gx in range(steps):
                for gy in range(steps):
                    x, y = x_lo + gx * 0.05, y_lo + gy * 0.05
                    rect = (x - w / 2, y - h / 2, x + w / 2, y + h / 2)
                    c = clearance(rect, set(members))
                    if best is None or c > best[0]:
                        best = (c, rect)
            if best[0] < 0.06:
                raise LayoutError(f"Venn diagram: the label {label!r} does not fit its region: "
                                  "shorten it or split the diagram")
            colour = blend(self._hex(slot), "0F172A", 0.3) if slot is not None else None
            rect = best[1]
            made.append(self._text(s, rect[0], rect[1], w, h, label, size, align="center",
                                   color=colour, name=f"Label: {label}"))
        listing = "; ".join(f"{label} = {' and '.join(sets[ids.index(m)][1] for m in members)}"
                            for members, label in overlaps)
        return self._finish(s, made, "venn", title, alt or (
            f"Venn diagram{': ' + title if title else ''} of "
            f"{', '.join(label for _, label in sets)}. {listing}."))

    def sankey(self, s, flows, *, box=None, font=None, title=None, alt=None):
        """A Sankey diagram: flows [(src, dst, value)]; nodes stand in columns from the sources
        to the sinks with heights by throughput, and every flow is a filled ribbon with a
        gradient from its source's colour to its target's."""
        made = []
        box = box or CONTENT
        font = _font(font or TYPE_SCALE["node"])
        label_h = font * LINE_SPACING / 72 + 0.04
        nodes, links = layout.sankey(flows, box[2], box[3], node_w=0.16, node_gap=0.35)
        order = list(nodes)
        colour = {n: self._hex(i) for i, n in enumerate(order)}
        for src, dst, value, ya, yb, thick in links:
            sx, _, sw, _, _ = nodes[src]
            tx = nodes[dst][0]
            made.append(self._ribbon(s, box[0] + sx + sw, box[1] + ya, box[0] + tx, box[1] + yb,
                                     thick, colour[src], colour[dst],
                                     f"{src} -> {dst} {_num(value)}"))
        last = max(x for x, *_ in nodes.values())
        for n, (x, y, w, h, value) in nodes.items():
            made.append(self._shape(s, "rect", box[0] + x, box[1] + y, w, max(h, 0.02),
                                    colour[n], None, name=f"Node: {n}"))
            ly = min(max(box[1] + y + h / 2 - label_h / 2, box[1]), box[1] + box[3] - label_h)
            text = f"{n} {_num(round(value, 2))}"
            tw = self.measure(text, font) + 0.1
            lx = box[0] + x + w + 0.08 if x < last - 1e-6 else box[0] + x - 0.08 - tw
            made.append(self._text(s, lx, ly, tw, label_h, text, font, name=f"Label: {n}"))
        listing = "; ".join(f"{a} to {b}: {_num(v)}" for a, b, v in flows)
        return self._finish(s, made, "sankey", title, alt or (
            f"Sankey diagram{': ' + title if title else ''}. Flows: {listing}."))

    def _ribbon(self, s, x0, y0, x1, y1, thick, c0, c1, name):
        """A flow ribbon from (x0, y0) to (x1, y1), `thick` high: two cubic Bezier edges, filled
        with a translucent gradient from c0 to c1."""
        bottom = max(y0, y1) + thick
        top = min(y0, y1)
        h = bottom - top
        w = x1 - x0
        sp = self._prst_shape(s, "rect", x0, top, w, h)
        ew, eh = round(w * EMU_IN), round(h * EMU_IN)

        def p(x, y):
            return f'<a:pt x="{round((x - x0) * EMU_IN)}" y="{round((y - top) * EMU_IN)}"/>'

        old = sp._element.spPr.find(qn("a:prstGeom"))
        xm = (x0 + x1) / 2
        path = (f'<a:path w="{ew}" h="{eh}"><a:moveTo>{p(x0, y0)}</a:moveTo>'
                f'<a:cubicBezTo>{p(xm, y0)}{p(xm, y1)}{p(x1, y1)}</a:cubicBezTo>'
                f'<a:lnTo>{p(x1, y1 + thick)}</a:lnTo>'
                f'<a:cubicBezTo>{p(xm, y1 + thick)}{p(xm, y0 + thick)}{p(x0, y0 + thick)}'
                f'</a:cubicBezTo><a:close/></a:path>')
        cust = etree.fromstring(
            f'<a:custGeom xmlns:a="{_NS_A}"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
            f'<a:rect l="l" t="t" r="r" b="b"/><a:pathLst>{path}</a:pathLst></a:custGeom>')
        old.addprevious(cust)
        sp._element.spPr.remove(old)
        sp.fill.gradient()
        sp.fill.gradient_angle = 0
        for stop, c in zip(sp.fill.gradient_stops, (c0, c1)):
            stop.color.rgb = _rgb(c)
            etree.SubElement(stop._gs[0], qn("a:alpha")).set("val", "42000")
        sp.line.fill.background()
        sp.shadow.inherit = False
        sp.name = f"Deco: flow {name}"
        return sp


def _gap(rect, c):
    """Distance from point c to the nearest point of rect (x0, y0, x1, y1)."""
    dx = max(rect[0] - c[0], 0.0, c[0] - rect[2])
    dy = max(rect[1] - c[1], 0.0, c[1] - rect[3])
    return math.hypot(dx, dy)


def _hit(a, b, pad=0.02):
    return a[0] < b[2] + pad and b[0] < a[2] + pad and a[1] < b[3] + pad and b[1] < a[3] + pad
