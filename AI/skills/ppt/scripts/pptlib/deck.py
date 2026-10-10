"""
pptlib - reusable python-pptx helpers for themed 16:9 presentations.

Usage:
    from pptlib import Deck
    d = Deck()                      # default Modern Minimal (slate + sky blue)
    s = d.slide()
    d.header(s, "Architecture", "System Overview", 3)
    d.bullet(d.textbox(s, 0.9, 2.0, 5, 4), "First point", first=True)
    d.img_fit(s, "img/arch.png", 6.3, 2.0, 6.3, 4.5)
    d.save(r"C:\\path\\Deck.pptx")

Override the theme by passing a dict of hex colors / font names:
    Deck(theme={"accent": "E11D48", "navy": "1A1A2E", "font": "Calibri"})
"""
import os
from pptx import Presentation
from pptx.chart.data import CategoryChartData  # noqa: F401  (re-exported by the package)
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: F401  (as above)
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.shapes.autoshape import AutoShapeType
from PIL import Image

from .charts import CHART_KINDS, LEGEND_POS, Charts  # noqa: F401
from .diagrams import Diagrams
from .shapes import Shapes, alt_text
from .structures import Structures
from .style import DEFAULT_PALETTE, write_theme

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

DEFAULT_THEME = {
    "navy": "0F172A", "navy_dk": "0A0F1E", "accent": "0EA5E9", "accent2": "6366F1",
    "light": "F8FAFC", "white": "FFFFFF", "text": "0F172A", "muted": "94A3B8",
    "warn": "F59E0B", "panel": "F1F5F9", "card_border": "E2E8F0",
    "subtle": "CBD5E1", "line": "64748B", "font": "Calibri Light", "mono": "Cascadia Code",
}


def hex2rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Deck(Shapes, Diagrams, Structures, Charts):
    """A 16:9 presentation builder with a consistent theme."""

    def __init__(self, theme=None):
        self.t = {**DEFAULT_THEME, **(theme or {})}
        for k, v in self.t.items():
            if k not in ("font", "mono", "palette"):
                setattr(self, k.upper(), hex2rgb(v))
        self.FONT = self.t["font"]
        self.MONO = self.t["mono"]
        # The palette defaults to the deck's accents, so charts and diagrams follow a new accent.
        palette = self.t.get("palette") or (self.t["accent"], self.t["accent2"], self.t["warn"],
                                            *DEFAULT_PALETTE[3:])
        if len(palette) != 6:
            raise ValueError(f"theme palette needs 6 colours, got {len(palette)}")
        self.palette_hex = [c.lstrip("#").upper() for c in palette]
        self.t["palette"] = self.palette_hex
        self.PALETTE = [hex2rgb(c) for c in self.palette_hex]
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self._blank = self.prs.slide_layouts[6]
        write_theme(self.prs, self.palette_hex, self.t)

    # -- slide / save -----------------------------------------------------
    def slide(self):
        return self.prs.slides.add_slide(self._blank)

    def save(self, path):
        self.prs.save(path)
        return path

    # -- shapes -----------------------------------------------------------
    def rect(self, s, x, y, w, h, fill=None, line=None, line_w=1.0,
             shape=MSO_SHAPE.RECTANGLE, shadow=False):
        sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line
            sp.line.width = Pt(line_w)
        # IMPORTANT: shadow.inherit=False inserts an empty <a:effectLst/>. Adding a
        # SECOND effectLst produces schema-invalid XML that PowerPoint refuses to
        # open (python-pptx is more lenient). Reuse the existing element instead.
        sp.shadow.inherit = False
        if shadow:
            spPr = sp._element.spPr
            ef = spPr.find(qn("a:effectLst"))
            if ef is None:
                ef = spPr.makeelement(qn("a:effectLst"), {})
                spPr.append(ef)
            sh = ef.makeelement(qn("a:outerShdw"),
                                {"blurRad": "90000", "dist": "38100",
                                 "dir": "5400000", "rotWithShape": "0"})
            clr = sh.makeelement(qn("a:srgbClr"), {"val": self.t["navy"]})
            clr.append(clr.makeelement(qn("a:alpha"), {"val": "28000"}))
            sh.append(clr)
            ef.append(sh)
        return sp

    def connector(self, s, x1, y1, x2, y2, color=None, width=1.5,
                  kind=MSO_CONNECTOR.STRAIGHT):
        cn = s.shapes.add_connector(kind, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        cn.line.color.rgb = color or self.ACCENT2
        cn.line.width = Pt(width)
        cn.shadow.inherit = False
        return cn

    # -- text -------------------------------------------------------------
    def textbox(self, s, x, y, w, h, anchor="top"):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = ANCHOR[anchor]
        for m in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{m}", 0)
        return tf

    def _para(self, tf, first):
        if first and not tf.paragraphs[0].runs:
            return tf.paragraphs[0]
        return tf.add_paragraph()

    def run(self, p, text, size=18, color=None, bold=False, italic=False,
            mono=False, font=None):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color if color is not None else self.TEXT
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font or (self.MONO if mono else self.FONT)
        return r

    def para(self, tf, first=False, align=None, space_before=0, space_after=0):
        p = self._para(tf, first)
        if align:
            p.alignment = ALIGN[align]
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        return p

    def bullet(self, tf, text, size=18, color=None, bold=False, level=0,
               space_after=10, first=False):
        p = self.para(tf, first, space_after=space_after)
        p.level = level
        self.run(p, text, size=size, color=color, bold=bold)
        self._native_bullet(p, level)
        return p

    def _native_bullet(self, p, level):
        """Give paragraph p real DrawingML bullet formatting (not a literal glyph run)."""
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.25 + 0.25 * level))))
        pPr.set("indent", str(int(-Inches(0.25))))
        for tag in ("a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts",
                    "a:buFontTx", "a:buFont", "a:buNone", "a:buAutoNum", "a:buChar"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        for xml in ('<a:buClr><a:srgbClr val="%s"/></a:buClr>' % self.t["accent"],
                    '<a:buFont typeface="Arial"/>', '<a:buChar char="\u2022"/>'):
            el = parse_xml(
                '<a:x xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '%s</a:x>' % xml)[0]
            pPr.insert_element_before(el, "a:tabLst", "a:defRPr", "a:extLst")

    # -- images -----------------------------------------------------------
    def img_fit(self, s, path, bx, by, bw, bh, align="center", valign="middle", alt=None):
        """Insert an image scaled to fit inside the box, preserving aspect ratio."""
        iw, ih = Image.open(path).size
        ar, bar = iw / ih, bw / bh
        if ar > bar:
            w, h = bw, bw / ar
        else:
            w, h = bh * ar, bh
        x = {"left": bx, "center": bx + (bw - w) / 2, "right": bx + bw - w}[align]
        y = {"top": by, "middle": by + (bh - h) / 2, "bottom": by + bh - h}[valign]
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        if alt:
            alt_text(pic, alt)
        return x, y, w, h

    # -- composite layout -------------------------------------------------
    def header(self, s, kicker, title, num, footer_label="Presentation"):
        """Standard content-slide chrome: bg, top accent line, kicker, title, footer."""
        self.rect(s, 0, 0, 13.333, 7.5, fill=self.LIGHT)
        self.rect(s, 0, 0, 13.333, 0.05, fill=self.ACCENT)
        tf = self.textbox(s, 1.2, 0.5, 10, 0.35)
        self.run(self.para(tf, True), kicker.upper(), size=11, color=self.ACCENT, bold=True)
        tf2 = self.textbox(s, 1.2, 0.85, 10, 0.75)
        self.run(self.para(tf2, True), title, size=28, color=self.NAVY, bold=True)
        self.rect(s, 1.2, 1.55, 1.5, 0.04, fill=self.ACCENT2)
        self.footer(s, num, footer_label)

    def footer(self, s, num, label="Presentation"):
        tf = self.textbox(s, 0.7, 7.02, 8, 0.35)
        self.run(self.para(tf, True), label, size=10, color=self.MUTED)
        tf2 = self.textbox(s, 11.8, 7.02, 0.8, 0.35)
        self.run(self.para(tf2, True, align="right"), f"{num:02d}", size=10,
                 color=self.MUTED, bold=True)

    def node(self, s, x, y, w, h, text, fill=None, text_color=None, size=14,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True, kind=None, name=None):
        """A labelled box for hand-built (native) diagrams. Returns the shape.

        With `kind` (e.g. "process", "decision") the box takes the modern diagram style of that
        node kind (see styled_node()): `shape` then only overrides the kind's geometry, and
        `shadow` is ignored. `name` names the shape (default "Node: <text>" for kinds)."""
        if kind is not None:
            prst = None if shape == MSO_SHAPE.ROUNDED_RECTANGLE else AutoShapeType(shape).prst
            return self.styled_node(s, x, y, w, h, text, kind, size=size, prst=prst, fill=fill,
                                    text_color=text_color, name=name)
        sp = self.rect(s, x, y, w, h, fill=fill or self.PANEL,
                       line=self.ACCENT, line_w=1.5, shape=shape, shadow=shadow)
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        self.run(self.para(tf, True, align="center"), text, size=size,
                 color=text_color or self.NAVY, bold=True)
        if name:
            sp.name = name
        return sp

    # -- tables / charts --------------------------------------------------
    def table(self, s, x, y, w, h, headers, rows, col_widths=None, col_align=None,
              font_size=12, header_fill=None, banded=True, alt=None):
        """A native, editable PowerPoint table with themed navy header and banded body.
        `alt` is its alternative text (default: the headers and row count)."""
        if not headers:
            raise ValueError("headers must be non-empty")
        ncols = len(headers)
        for i, row in enumerate(rows):
            if len(row) != ncols:
                raise ValueError(f"row {i} has {len(row)} cells, expected {ncols}")
        if col_widths is not None and len(col_widths) != ncols:
            raise ValueError("col_widths length must match column count")
        if col_align is not None and len(col_align) != ncols:
            raise ValueError("col_align length must match column count")
        gf = s.shapes.add_table(len(rows) + 1, ncols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
        tbl = gf.table
        # Suppress the built-in style banding (template-themed, clashes with our
        # palette) and colour every cell explicitly from the deck theme instead.
        tblPr = tbl._tbl.tblPr
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
        if col_widths is not None:
            for i, cw in enumerate(col_widths):
                tbl.columns[i].width = Inches(cw)
        grid = [headers] + list(rows)
        for r, row in enumerate(grid):
            fill = header_fill or self.NAVY if r == 0 else (
                self.PANEL if banded and r % 2 == 0 else self.WHITE)
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
                p = cell.text_frame.paragraphs[0]
                if col_align is not None:
                    p.alignment = ALIGN[col_align[c]]
                run = p.add_run()
                run.text = str(val)
                run.font.size = Pt(font_size)
                run.font.name = self.FONT
                run.font.bold = r == 0
                run.font.color.rgb = self.WHITE if r == 0 else self.TEXT
        alt_text(gf, alt or f"Table with columns {', '.join(map(str, headers))}; "
                            f"{len(rows)} rows")
        return tbl
